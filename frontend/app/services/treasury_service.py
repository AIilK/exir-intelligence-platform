from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
import re
from typing import Any
import unicodedata

from sqlalchemy import text
from sqlalchemy.engine import Engine

from app.core.config import settings
from app.database.sqlserver import get_sqlserver_engine
from app.services.received_cheque_current_status import (
    current_received_holding_label,
    current_received_state_expr,
    current_received_status_apply,
    received_cheque_open_expression,
    received_cheque_open_predicate,
    received_cheque_is_approved_open_holding,
)


_PERSIAN_SEARCH_TRANSLATION = str.maketrans(
    {
        "آ": "ا",
        "أ": "ا",
        "إ": "ا",
        "ٱ": "ا",
        "ي": "ی",
        "ى": "ی",
        "ك": "ک",
        "ة": "ه",
        "ۀ": "ه",
        "\u200c": " ",
    }
)


def _normalize_account_search(value: str) -> str:
    normalized = unicodedata.normalize("NFKC", value)
    normalized = "".join(
        character
        for character in normalized
        if unicodedata.category(character) != "Mn"
    )
    normalized = normalized.translate(_PERSIAN_SEARCH_TRANSLATION)
    normalized = re.sub(r"[^\w\u0600-\u06ff]+", " ", normalized)
    return " ".join(normalized.split())


def _as_number(value: Any) -> int | float:
    if value is None:
        return 0.0
    number = value if isinstance(value, Decimal) else Decimal(str(value))
    if number == number.to_integral_value():
        return int(number)
    return float(number)


def _as_datetime(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def _as_jalali_date(value: Any) -> str | None:
    if value is None:
        return None
    try:
        import jdatetime
    except ImportError:
        return None

    if isinstance(value, datetime):
        gregorian_date = value.date()
    elif isinstance(value, date):
        gregorian_date = value
    else:
        try:
            gregorian_date = datetime.fromisoformat(str(value)).date()
        except ValueError:
            return None

    return jdatetime.date.fromgregorian(date=gregorian_date).isoformat().replace("-", "/")


def _currency_name(currency_ref: Any) -> str | None:
    if currency_ref == settings.treasury_operational_currency_ref:
        return settings.treasury_operational_currency_name
    return None


def _cheque_state_label(cheque_type: str, state: int) -> str:
    """برچسب‌های تأییدشده یا استنباط‌شده از تراکنش‌های واقعی راهکاران."""

    labels = {
        "received": {
            1: "registered",
            2: "in_collection",
            3: "collected",
            4: "protested",
            10: "delivered_to_counterparty",
        },
        "issued": {
            11: "issued_open_unposted",
            15: "restored_or_cancelled",
            # Business rule confirmed by Finance: when the payable cheque is
            # withdrawn from the bank, Rahkaran moves it out of the long-term
            # bucket into the daily/cleared state. In this installation that
            # finalized state is 28 and it means the cheque has been posted.
            28: "posted_withdrawn",
            -27: "clearing_reversed",
            -12: "issue_reversed",
        },
    }
    return labels.get(cheque_type, {}).get(state, "unknown")



def _issued_posting_status(state: Any, status_document_description: Any = None) -> tuple[bool, str]:
    """Return the Finance-approved accounting meaning of an issued cheque state.

    State 11 is still an open long-term obligation. A move to a configured
    cleared state (currently State=28) means the bank withdrawal happened and
    the cheque has been posted/accounted for; it must not remain in future
    cash-flow obligations.
    """

    status_text = str(status_document_description or "").strip()
    normalized_status_text = status_text.replace("ي", "ی").replace("ك", "ک")
    if any(marker in normalized_status_text for marker in ("تعیین وضعیت چک", "تعین وضعیت چک", "وصول چ", "پرداخت چک", "برداشت وجه چک")):
        return True, "سند خورده / برداشت‌شده"

    try:
        code = int(state)
    except (TypeError, ValueError):
        return False, "وضعیت نامشخص"
    paid_states = {
        int(value.strip())
        for value in settings.treasury_issued_paid_states.split(",")
        if value.strip().lstrip("-").isdigit()
    }
    if code in paid_states:
        return True, "سند خورده / برداشت‌شده"
    if code == 11:
        return False, "باز / سند نخورده"
    if code == 15:
        return False, "ابطال / اعاده"
    if code in {-27, -12}:
        return False, "برگشت عملیات"
    return False, "وضعیت نامشخص"

def _master_cheque_payload(row: Any, cheque_type: str) -> dict[str, Any]:
    state = row["ChequeState"]
    current_status_description = row.get("CurrentStatusDescription")
    is_posted, posting_status = _issued_posting_status(state) if cheque_type == "issued" else (False, None)
    received_status_label = (
        current_received_holding_label(state, current_status_description)
        if cheque_type == "received"
        else None
    )
    return {
        "cheque_id": row["ChequeID"],
        "document_id": row.get("DocumentID"),
        "document_number": row.get("DocumentNumber"),
        "document_date": _as_datetime(row.get("DocumentDate")),
        "document_date_jalali": _as_jalali_date(row.get("DocumentDate")),
        "cheque_type": cheque_type,
        "state_code": state,
        "master_state_code": row.get("MasterChequeState", state),
        "state": _cheque_state_label(cheque_type, state),
        "state_label": posting_status if cheque_type == "issued" else received_status_label,
        "current_status_description": current_status_description,
        "current_status_date": _as_datetime(row.get("CurrentStatusDate")),
        "current_status_date_jalali": _as_jalali_date(row.get("CurrentStatusDate")),
        "current_status_document_date": _as_datetime(row.get("CurrentStatusDocumentDate")),
        "current_status_document_date_jalali": _as_jalali_date(row.get("CurrentStatusDocumentDate")),
        "current_status_transaction_id": row.get("CurrentStatusTransactionID"),
        "status_source": (
            "latest_receivable_note_transaction"
            if cheque_type == "received" and row.get("CurrentStatusTransactionID") is not None
            else "receivable_note_master" if cheque_type == "received" else None
        ),
        "is_posted": is_posted if cheque_type == "issued" else None,
        "cashflow_open": (not is_posted and state == 11) if cheque_type == "issued" else None,
        "status_document_id": row.get("StatusDocumentID"),
        "status_document_number": row.get("StatusDocumentNumber"),
        "status_document_date": _as_datetime(row.get("StatusDocumentDate")),
        "status_document_date_jalali": _as_jalali_date(row.get("StatusDocumentDate")),
        "status_document_description": row.get("StatusDocumentDescription"),
        "amount": _as_number(row["Amount"]),
        "due_date": _as_datetime(row["DueDate"]),
        "due_date_jalali": _as_jalali_date(row["DueDate"]),
        "days_until_due": row["DaysUntilDue"],
        "serial_number": row["SerialNumber"],
        "series": row["Series"],
        "sayad_number": row["SayadNumber"],
        "account_number": row["AccountNumber"],
        "bank_ref": row["BankRef"],
        "bank_name": row["BankName"],
        "bank_branch_name": row["BankBranchName"],
        "bank_branch_code": row["BankBranchCode"],
        "counterpart_ref": row["CounterPartRef"],
        "counterpart_code": row["CounterPartCode"],
        "counterpart_name": row["CounterPartName"],
        "account_ref": row["AccountRef"],
        "currency_ref": row["CurrencyRef"],
        "currency_name": _currency_name(row["CurrencyRef"]),
        "normal_or_guarantee": row["NormalORGuarantee"],
        "description": row["Description"],
        "source_system": "rahkaran",
        "source_label": "راهکاران",
    }


def _karamad_cheques(source_kind: str) -> list[dict[str, Any]]:
    from app.services.karamad_manual_import_service import KaramadManualImportService
    return KaramadManualImportService().cheque_rows(source_kind)


def _karamad_received_period(rows: list[dict[str, Any]], period: str) -> list[dict[str, Any]]:
    horizon = {"1m": 30, "3m": 90, "6m": 180, "12m": 365}
    if period == "all":
        return rows
    if period == "overdue":
        return [row for row in rows if row.get("days_until_due") is not None and row["days_until_due"] < 0]
    days = horizon[period]
    return [row for row in rows if row.get("days_until_due") is not None and 0 <= row["days_until_due"] < days]


def _engine(engine: Engine | None) -> Engine:
    return engine or get_sqlserver_engine()


def _account_payload(row: Any) -> dict[str, Any]:
    return {
        "account_id": row["AccountID"],
        "name": row["Name"],
        "number": row["Number"],
        "currency_ref": row["CurrencyRef"],
        "currency_name": _currency_name(row["CurrencyRef"]),
        "account_type_ref": row["AccountTypeRef"],
        "party_ref": row["PartyRef"],
        "status": row["Status"],
        "debit": _as_number(row["DebitBalance"]),
        "credit": _as_number(row["CreditBalance"]),
        "balance": _as_number(row["Balance"]),
    }


def search_accounts(
    account_name: str,
    limit: int = 20,
    engine: Engine | None = None,
) -> list[dict[str, Any]]:
    """حساب‌ها را بدون انتخاب خودکار یک نتیجه مبهم جست‌وجو می‌کند."""

    search_term = _normalize_account_search(account_name)
    if not search_term:
        return []

    safe_limit = max(1, min(int(limit), 50))
    tokens = search_term.split()[:8]

    normalized_name_sql = """
        REPLACE(
            REPLACE(
                REPLACE(
                    REPLACE(
                        REPLACE(
                            REPLACE(Name, N'آ', N'ا'),
                            N'أ', N'ا'
                        ),
                        N'إ', N'ا'
                    ),
                    N'ي', N'ی'
                ),
                N'ى', N'ی'
            ),
            N'ك', N'ک'
        )
    """

    token_filters = " AND ".join(
        f"{normalized_name_sql} LIKE :token_{index}"
        for index in range(len(tokens))
    )

    query = text(
        f"""
        SELECT TOP ({safe_limit})
            AccountID,
            Name,
            Number,
            CurrencyRef,
            AccountTypeRef,
            PartyRef,
            Status,
            DebitBalance,
            CreditBalance,
            Balance
        FROM FIN3.Account
        WHERE {token_filters}
        ORDER BY
            CASE
                WHEN LTRIM(RTRIM({normalized_name_sql})) = :exact_name
                THEN 0 ELSE 1
            END,
            Name,
            CurrencyRef
        """
    )

    parameters = {
        f"token_{index}": f"%{token}%"
        for index, token in enumerate(tokens)
    }
    parameters["exact_name"] = search_term

    with _engine(engine).connect() as connection:
        rows = connection.execute(
            query,
            parameters,
        ).mappings().all()

    return [_account_payload(row) for row in rows]


def _get_account_by_id(
    account_id: int,
    engine: Engine | None = None,
) -> dict[str, Any] | None:
    query = text(
        """
        SELECT
            AccountID,
            Name,
            Number,
            CurrencyRef,
            AccountTypeRef,
            PartyRef,
            Status,
            DebitBalance,
            CreditBalance,
            Balance
        FROM FIN3.Account
        WHERE AccountID = :account_id
        """
    )

    with _engine(engine).connect() as connection:
        row = connection.execute(
            query,
            {"account_id": int(account_id)},
        ).mappings().first()

    return _account_payload(row) if row else None


def _resolve_account(
    account_name: str | None,
    account_id: int | None,
    engine: Engine | None = None,
) -> dict[str, Any]:
    if account_id is not None:
        account = _get_account_by_id(account_id, engine=engine)
        if account is None:
            return {
                "status": "not_found",
                "message": "حساب با شناسه داده‌شده پیدا نشد.",
            }
        return {"status": "success", "account": account}

    matches = search_accounts(account_name or "", engine=engine)

    if not matches:
        return {
            "status": "not_found",
            "message": "حساب پیدا نشد.",
        }

    if len(matches) > 1:
        return {
            "status": "ambiguous",
            "message": "چند حساب پیدا شد؛ حساب و ارز موردنظر را مشخص کنید.",
            "matches": matches,
        }

    return {"status": "success", "account": matches[0]}


def get_account_balance(
    account_name: str | None = None,
    account_id: int | None = None,
    engine: Engine | None = None,
) -> dict[str, Any]:
    """مانده ذخیره‌شده یک حساب را از FIN3.Account برمی‌گرداند."""

    resolved = _resolve_account(account_name, account_id, engine=engine)
    if resolved["status"] != "success":
        return resolved

    return {
        "status": "success",
        **resolved["account"],
    }


def get_account_transactions(
    account_name: str | None = None,
    limit: int = 10,
    account_id: int | None = None,
    engine: Engine | None = None,
) -> dict[str, Any]:
    """گردش عمومی حساب را از هسته FIN3.Transaction دریافت می‌کند."""

    resolved = _resolve_account(account_name, account_id, engine=engine)
    if resolved["status"] != "success":
        return resolved

    account = resolved["account"]
    safe_limit = max(1, min(int(limit), 1000))

    transactions_query = text(
        f"""
        SELECT TOP ({safe_limit})
            TransactionID,
            BookingDate,
            SettlementDate,
            Description,
            EntryAmount,
            GLAmount,
            Effect,
            Status,
            ReferenceTypeComponentName,
            ReferenceTypeEntityName,
            ReferenceRef
        FROM FIN3.[Transaction]
        WHERE AccountRef = :account_id
          AND Status = 1
        ORDER BY BookingDate DESC, TransactionID DESC
        """
    )

    summary_query = text(
        """
        SELECT
            COALESCE(SUM(CASE WHEN Effect = 1 THEN EntryAmount ELSE 0 END), 0)
                AS DebitTotal,
            COALESCE(SUM(CASE WHEN Effect = 2 THEN EntryAmount ELSE 0 END), 0)
                AS CreditTotal
        FROM FIN3.[Transaction]
        WHERE AccountRef = :account_id
          AND Status = 1
        """
    )

    with _engine(engine).connect() as connection:
        transaction_rows = connection.execute(
            transactions_query,
            {"account_id": account["account_id"]},
        ).mappings().all()
        summary = connection.execute(
            summary_query,
            {"account_id": account["account_id"]},
        ).mappings().one()

    transactions = []
    for row in transaction_rows:
        effect = row["Effect"]
        transactions.append(
            {
                "transaction_id": row["TransactionID"],
                "booking_date": _as_datetime(row["BookingDate"]),
                "booking_date_jalali": _as_jalali_date(row["BookingDate"]),
                "settlement_date": _as_datetime(row["SettlementDate"]),
                "settlement_date_jalali": _as_jalali_date(row["SettlementDate"]),
                "description": row["Description"],
                "amount": _as_number(row["EntryAmount"]),
                "gl_amount": _as_number(row["GLAmount"]),
                "gl_currency_name": settings.treasury_operational_currency_name,
                "effect": effect,
                "side": (
                    "debit"
                    if effect == 1
                    else "credit"
                    if effect == 2
                    else "unknown"
                ),
                "status": row["Status"],
                "reference_component": row["ReferenceTypeComponentName"],
                "reference_entity": row["ReferenceTypeEntityName"],
                "reference_ref": row["ReferenceRef"],
            }
        )

    debit_total = _as_number(summary["DebitTotal"])
    credit_total = _as_number(summary["CreditTotal"])

    return {
        "status": "success",
        "account": account,
        "summary": {
            "debit_total": debit_total,
            "credit_total": credit_total,
            "calculated_balance": debit_total - credit_total,
            "stored_balance": account["balance"],
            "gl_currency_name": settings.treasury_operational_currency_name,
        },
        "transactions": transactions,
        "returned_count": len(transactions),
    }


def _get_latest_treasury_documents(
    *,
    table_name: str,
    id_column: str,
    document_type: str,
    limit: int,
    engine: Engine | None = None,
) -> dict[str, Any]:
    """آخرین هدرهای تأییدشده دریافت یا پرداخت را برمی‌گرداند."""

    allowed_sources = {
        ("Receipt", "ReceiptID", "receipt"),
        ("Payment", "PaymentID", "payment"),
    }
    if (table_name, id_column, document_type) not in allowed_sources:
        raise ValueError("Unsupported treasury document source")

    safe_limit = max(1, min(int(limit), 1000))
    query = text(
        f"""
        SELECT TOP ({safe_limit})
            [{id_column}] AS [DocumentID],
            [Number] AS [DocumentNumber],
            [Date] AS [DocumentDate],
            [ApproveState],
            [CounterPartRef],
            [TotalOperationalCurrencyAmount] AS [TotalAmount],
            [ItemType],
            [IsDeployment],
            [FiscalYearRef],
            [Description]
        FROM RPA3.[{table_name}]
        WHERE [ApproveState] = 3
        ORDER BY [Date] DESC, [{id_column}] DESC
        """
    )

    with _engine(engine).connect() as connection:
        rows = connection.execute(query).mappings().all()

    documents = [
        {
            "document_id": row["DocumentID"],
            "number": row["DocumentNumber"],
            "date": _as_datetime(row["DocumentDate"]),
            "date_jalali": _as_jalali_date(row["DocumentDate"]),
            "approve_state": row["ApproveState"],
            "counterpart_ref": row["CounterPartRef"],
            "total_amount": _as_number(row["TotalAmount"]),
            "item_type": row["ItemType"],
            "is_deployment": bool(row["IsDeployment"]),
            "fiscal_year_ref": row["FiscalYearRef"],
            "description": row["Description"],
        }
        for row in rows
    ]

    return {
        "status": "success",
        "document_type": document_type,
        "approve_state_filter": 3,
        "currency_name": settings.treasury_operational_currency_name,
        "documents": documents,
        "returned_count": len(documents),
        "returned_total_amount": sum(
            document["total_amount"]
            for document in documents
        ),
    }


def get_latest_receipts(
    limit: int = 10,
    engine: Engine | None = None,
) -> dict[str, Any]:
    """آخرین اسناد دریافت تأییدشده خزانه را می‌گیرد."""

    return _get_latest_treasury_documents(
        table_name="Receipt",
        id_column="ReceiptID",
        document_type="receipt",
        limit=limit,
        engine=engine,
    )


def get_latest_payments(
    limit: int = 10,
    engine: Engine | None = None,
) -> dict[str, Any]:
    """آخرین اسناد پرداخت تأییدشده خزانه را می‌گیرد."""

    return _get_latest_treasury_documents(
        table_name="Payment",
        id_column="PaymentID",
        document_type="payment",
        limit=limit,
        engine=engine,
    )


def get_latest_received_cheques(
    limit: int = 10,
    engine: Engine | None = None,
) -> dict[str, Any]:
    """آخرین چک‌های ثبت‌شده در اسناد دریافت تأییدشده را برمی‌گرداند."""

    safe_limit = max(1, min(int(limit), 100))
    query = text(
        f"""
        SELECT TOP ({safe_limit})
            note.[ReceiptReceivableNoteID] AS [ChequeItemID],
            receipt.[ReceiptID] AS [DocumentID],
            receipt.[Number] AS [DocumentNumber],
            receipt.[Date] AS [DocumentDate],
            note.[Amount],
            note.[DueDate],
            note.[SerialNumber],
            note.[Series],
            note.[SayadNumber],
            note.[AccountNumber],
            note.[BankRef],
            note.[BankBranchRef],
            note.[BankBranchName],
            note.[BankBranchCode],
            note.[CounterPartRef],
            note.[AccountRef],
            note.[NoteType],
            note.[NormalORGuarantee],
            note.[CurrencyRef],
            master_note.[State] AS [ChequeState],
            bank.[Name] AS [BankName],
            counterpart.[Code] AS [CounterPartCode],
            counterpart.[Title] AS [CounterPartName],
            note.[Description] AS [ChequeDescription],
            receipt.[Description] AS [DocumentDescription]
        FROM RPA3.[ReceiptReceivableNote] AS note
        INNER JOIN RPA3.[Receipt] AS receipt
            ON receipt.[ReceiptID] = note.[ReceiptRef]
        LEFT JOIN RPA3.[ReceivableNote] AS master_note
            ON master_note.[ReceivableNoteID] = note.[ReceivableNoteRef]
        LEFT JOIN RPA3.[Bank] AS bank
            ON bank.[BankID] = note.[BankRef]
        LEFT JOIN FIN3.[DL] AS counterpart
            ON counterpart.[DLID] = note.[CounterPartRef]
        WHERE receipt.[ApproveState] = 3
          AND receipt.[ItemType] = 1
          AND COALESCE(
                master_note.[NormalORGuarantee],
                note.[NormalORGuarantee]
              ) = 1
          AND ISNULL(master_note.[Description], N'') NOT LIKE N'%ضمانت%'
          AND ISNULL(master_note.[Description], N'') NOT LIKE N'%تضمین%'
          AND ISNULL(master_note.[Description], N'') NOT LIKE N'%حسن انجام%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%ضمانت%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%تضمین%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%حسن انجام%'
        ORDER BY
            receipt.[Date] DESC,
            note.[ReceiptReceivableNoteID] DESC
        """
    )

    with _engine(engine).connect() as connection:
        rows = connection.execute(query).mappings().all()

    cheques = [
        {
            "cheque_item_id": row["ChequeItemID"],
            "document_id": row["DocumentID"],
            "document_number": row["DocumentNumber"],
            "document_date": _as_datetime(row["DocumentDate"]),
            "document_date_jalali": _as_jalali_date(row["DocumentDate"]),
            "amount": _as_number(row["Amount"]),
            "due_date": _as_datetime(row["DueDate"]),
            "due_date_jalali": _as_jalali_date(row["DueDate"]),
            "serial_number": row["SerialNumber"],
            "series": row["Series"],
            "sayad_number": row["SayadNumber"],
            "account_number": row["AccountNumber"],
            "bank_ref": row["BankRef"],
            "bank_branch_ref": row["BankBranchRef"],
            "bank_branch_name": row["BankBranchName"],
            "bank_branch_code": row["BankBranchCode"],
            "counterpart_ref": row["CounterPartRef"],
            "account_ref": row["AccountRef"],
            "note_type": row["NoteType"],
            "normal_or_guarantee": row["NormalORGuarantee"],
            "currency_ref": row["CurrencyRef"],
            "currency_name": _currency_name(row["CurrencyRef"]),
            "state_code": row["ChequeState"],
            "state": _cheque_state_label(
                "received",
                row["ChequeState"],
            ),
            "bank_name": row["BankName"],
            "counterpart_code": row["CounterPartCode"],
            "counterpart_name": row["CounterPartName"],
            "description": (
                row["ChequeDescription"]
                or row["DocumentDescription"]
            ),
        }
        for row in rows
    ]

    return {
        "status": "success",
        "cheque_type": "received",
        "approve_state_filter": 3,
        "item_type_filter": 1,
        "normal_or_guarantee_filter": 1,
        "guarantee_cheques_excluded": True,
        "guarantee_detection_rules": [
            "NormalORGuarantee must equal 1",
            "Description must not contain ضمانت, تضمین or حسن انجام",
        ],
        "cheques": cheques,
        "returned_count": len(cheques),
        "returned_total_amount": sum(
            cheque["amount"] for cheque in cheques
        ),
    }


def get_open_received_cheques(
    period: str = "all",
    engine: Engine | None = None,
) -> dict[str, Any]:
    """All non-guarantee open received cheques, optionally filtered by due horizon."""

    filters = {
        "all": "",
        "overdue": "AND note.[DueDate] < CAST(GETDATE() AS date)",
        "1m": "AND note.[DueDate] >= CAST(GETDATE() AS date) AND note.[DueDate] < DATEADD(month, 1, CAST(GETDATE() AS date))",
        "3m": "AND note.[DueDate] >= CAST(GETDATE() AS date) AND note.[DueDate] < DATEADD(month, 3, CAST(GETDATE() AS date))",
        "6m": "AND note.[DueDate] >= CAST(GETDATE() AS date) AND note.[DueDate] < DATEADD(month, 6, CAST(GETDATE() AS date))",
        "12m": "AND note.[DueDate] >= CAST(GETDATE() AS date) AND note.[DueDate] < DATEADD(month, 12, CAST(GETDATE() AS date))",
    }
    normalized_period = period.strip().lower()
    if normalized_period not in filters:
        raise ValueError("period must be all, overdue, 1m, 3m, 6m or 12m")

    current_status_apply = current_received_status_apply("note", "current_status")
    open_predicate = received_cheque_open_predicate("note", "current_status")
    effective_state = current_received_state_expr("note", "current_status")

    query = text(
        f"""
        SELECT
            note.[ReceivableNoteID] AS [ChequeID],
            note.[State] AS [MasterChequeState],
            {effective_state} AS [ChequeState],
            current_status.[CurrentStatusDescription],
            current_status.[CurrentStatusDate],
            current_status.[CurrentStatusDocumentDate],
            current_status.[CurrentStatusTransactionID],
            note.[Amount], note.[DueDate],
            DATEDIFF(day, CAST(GETDATE() AS date), CAST(note.[DueDate] AS date)) AS [DaysUntilDue],
            note.[SerialNumber], note.[Series], note.[SayadNumber], note.[AccountNumber],
            note.[BankRef], bank.[Name] AS [BankName], note.[BankBranchName], note.[BankBranchCode],
            note.[CounterPartRef], counterpart.[Code] AS [CounterPartCode],
            counterpart.[Title] AS [CounterPartName], note.[AccountRef], note.[CurrencyRef],
            note.[NormalORGuarantee], note.[Description],
            document_link.[DocumentID], document_link.[DocumentNumber], document_link.[DocumentDate]
        FROM RPA3.[ReceivableNote] AS note
        LEFT JOIN RPA3.[Bank] AS bank ON bank.[BankID] = note.[BankRef]
        LEFT JOIN FIN3.[DL] AS counterpart ON counterpart.[DLID] = note.[CounterPartRef]
        {current_status_apply}
        OUTER APPLY (
            SELECT TOP (1)
                receipt.[ReceiptID] AS [DocumentID],
                receipt.[Number] AS [DocumentNumber],
                receipt.[Date] AS [DocumentDate]
            FROM RPA3.[ReceiptReceivableNote] AS receipt_note
            INNER JOIN RPA3.[Receipt] AS receipt ON receipt.[ReceiptID] = receipt_note.[ReceiptRef]
            WHERE receipt_note.[ReceivableNoteRef] = note.[ReceivableNoteID]
              AND receipt.[ApproveState] = 3
            ORDER BY receipt.[Date] DESC, receipt.[ReceiptID] DESC
        ) AS document_link
        WHERE note.[NoteType] = 1
          {open_predicate}
          AND note.[NormalORGuarantee] = 1
          AND ISNULL(note.[Description], N'') NOT LIKE N'%ضمانت%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%تضمین%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%حسن انجام%'
          {filters[normalized_period]}
        ORDER BY CASE WHEN note.[DueDate] IS NULL THEN 1 ELSE 0 END,
                 note.[DueDate] ASC, note.[ReceivableNoteID] DESC
        """
    )
    with _engine(engine).connect() as connection:
        rows = connection.execute(query).mappings().all()
    rahkaran_cheques = [_master_cheque_payload(row, "received") for row in rows]
    karamad_cheques = _karamad_received_period(
        [row for row in _karamad_cheques("received_cheques") if received_cheque_is_approved_open_holding(row)],
        normalized_period,
    )
    cheques = sorted(
        rahkaran_cheques + karamad_cheques,
        key=lambda item: (item.get("due_date") or "9999-12-31", str(item.get("cheque_id") or "")),
    )
    return {
        "status": "success",
        "report_type": "open_received_cheques",
        "state_filter": "approved_open_holding_only",
        "allowed_open_holdings": ["نزد مأمور وصول", "نزد بانک", "نزد صندوق"],
        "legacy_master_state_filter": [1, 2],
        "current_status_basis": "فقط سه محل باز مصوب: نزد مأمور وصول، نزد بانک، نزد صندوق",
        "period_filter": normalized_period,
        "sql_due_date_horizon": "unbounded" if normalized_period == "all" else normalized_period,
        "row_limit": None,
        "guarantee_cheques_excluded": True,
        "data_sources": ["راهکاران", "کارآمد"],
        "guarantee_detection_rules": [
            "NormalORGuarantee must equal 1",
            "Description must not contain ضمانت, تضمین or حسن انجام",
        ],
        "cheques": cheques,
        "returned_count": len(cheques),
        "returned_total_amount": sum(cheque["amount"] for cheque in cheques),
    }


def get_open_issued_cheques(
    engine: Engine | None = None,
) -> dict[str, Any]:
    """Authoritative active issued cheques used by the dashboard.

    State 11 is the finance-confirmed open/active state. Cleared, cancelled,
    reversed, guarantee and shareholder cheques are deliberately excluded.
    """

    query = text(
        """
        SELECT
            note.[PayableNoteID] AS [ChequeID],
            note.[State] AS [ChequeState],
            note.[Amount], note.[DueDate],
            DATEDIFF(day, CAST(GETDATE() AS date), CAST(note.[DueDate] AS date)) AS [DaysUntilDue],
            note.[SerialNumber], note.[Series],
            CAST(NULL AS nvarchar(100)) AS [SayadNumber],
            note.[AccountNumber], note.[BankRef],
            bank.[Name] AS [BankName], note.[BankBranchName], note.[BankBranchCode],
            note.[CounterPartRef], counterpart.[Code] AS [CounterPartCode],
            counterpart.[Title] AS [CounterPartName], note.[AccountRef], note.[CurrencyRef],
            note.[NormalORGuarantee], note.[Description],
            document_link.[DocumentID], document_link.[DocumentNumber],
            COALESCE(document_link.[DocumentDate], note.[AgreementDate]) AS [DocumentDate],
            status_link.[StatusDocumentID], status_link.[StatusDocumentNumber],
            status_link.[StatusDocumentDate], status_link.[StatusDocumentDescription]
        FROM RPA3.[PayableNote] AS note
        LEFT JOIN RPA3.[Bank] AS bank ON bank.[BankID] = note.[BankRef]
        LEFT JOIN FIN3.[DL] AS counterpart ON counterpart.[DLID] = note.[CounterPartRef]
        OUTER APPLY (
            SELECT TOP (1)
                payment.[PaymentID] AS [DocumentID],
                payment.[Number] AS [DocumentNumber],
                payment.[Date] AS [DocumentDate]
            FROM RPA3.[PaymentPayableNote] AS payment_note
            INNER JOIN RPA3.[Payment] AS payment
                ON payment.[PaymentID] = payment_note.[PaymentRef]
            WHERE payment_note.[PayableNoteRef] = note.[PayableNoteID]
              AND payment.[ApproveState] = 3
            ORDER BY payment.[Date] DESC, payment.[PaymentID] DESC
        ) AS document_link
        OUTER APPLY (
            SELECT TOP (1)
                pnt.[DocumentRef] AS [StatusDocumentID],
                pnt.[DocumentNumber] AS [StatusDocumentNumber],
                COALESCE(pnt.[DocumentDate], pnt.[Date]) AS [StatusDocumentDate],
                pnt.[Description] AS [StatusDocumentDescription]
            FROM RPA3.[PayableNoteTransaction] AS pnt
            WHERE pnt.[PayableNoteRef] = note.[PayableNoteID]
              AND pnt.[State] = 11
              AND pnt.[DocumentItemType] IN (24, 26)
              AND pnt.[DocumentState] = 3
              AND (
                    ISNULL(pnt.[Description], N'') LIKE N'%تعیین وضعیت چک%'
                 OR ISNULL(pnt.[Description], N'') LIKE N'%تعین وضعیت چک%'
                 OR ISNULL(pnt.[Description], N'') LIKE N'%وصول چ%'
                 OR ISNULL(pnt.[Description], N'') LIKE N'%پرداخت چک%'
                 OR ISNULL(pnt.[Description], N'') LIKE N'%برداشت وجه چک%'
              )
            ORDER BY COALESCE(pnt.[DocumentDate], pnt.[Date]) DESC, pnt.[PayableNoteTransactionID] DESC
        ) AS status_link
        WHERE note.[NoteType] = 1
          AND note.[State] = 11
          AND status_link.[StatusDocumentID] IS NULL
          AND note.[NormalORGuarantee] = 1
          AND note.[DueDate] >= DATEADD(day, -20, CAST(GETDATE() AS date))
          AND ISNULL(note.[Description], N'') NOT LIKE N'%ضمانت%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%تضمین%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%حسن انجام%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%سهامدار%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%سود سهام%'
        ORDER BY CASE WHEN note.[DueDate] IS NULL THEN 1 ELSE 0 END,
                 note.[DueDate] ASC, note.[PayableNoteID] DESC
        """
    )
    with _engine(engine).connect() as connection:
        rows = connection.execute(query).mappings().all()
    rahkaran_cheques = [_master_cheque_payload(row, "issued") for row in rows]
    # KarAmand issued-cheque file is a snapshot of all registered cheques, not a complete
    # open-commitment feed. Keep the full snapshot visible in the KarAmand source page;
    # Rahkaran retains the 20-day overdue policy for open commitments.
    karamad_cheques = _karamad_cheques("issued_cheques")
    cheques = sorted(
        rahkaran_cheques + karamad_cheques,
        key=lambda item: (item.get("due_date") or "9999-12-31", str(item.get("cheque_id") or "")),
    )
    return {
        "status": "success",
        "report_type": "open_issued_cheques",
        "data_source": "RPA3.PayableNote + KarAmand Excel",
        "data_sources": ["راهکاران", "کارآمد"],
        "state_filter": [11],
        "paid_state_excluded": 28,
        "posting_business_rule": "وجود سند تأییدشده تعیین وضعیت/وصول چک در PayableNoteTransaction یا State پرداخت‌شده = سند خورده و برداشت‌شده؛ از تعهدات باز و Cash Flow آینده حذف می‌شود",
        "due_policy": "Rahkaran: from 20 days overdue through all future due dates; KarAmand: all registered cheques from latest snapshot (future coverage incomplete)",
        "row_limit": None,
        "guarantee_cheques_excluded": True,
        "shareholder_cheques_excluded": True,
        "cheques": cheques,
        "returned_count": len(cheques),
        "returned_total_amount": sum(cheque["amount"] for cheque in cheques),
    }


def get_latest_issued_cheques(
    limit: int = 10,
    engine: Engine | None = None,
    upcoming_only: bool = False,
    sort_by_due: bool = False,
) -> dict[str, Any]:
    """آخرین چک‌های ثبت‌شده در اسناد پرداخت تأییدشده را برمی‌گرداند."""

    safe_limit = max(1, min(int(limit), 1000))
    primary_future_filter = (
        "AND note.[DueDate] >= CAST(GETDATE() AS date) AND master_note.[State] = 11"
        if upcoming_only else ""
    )
    primary_order = (
        "CASE WHEN note.[DueDate] IS NULL THEN 1 ELSE 0 END, note.[DueDate] ASC, note.[PaymentPayableNoteID] DESC"
        if sort_by_due else "payment.[Date] DESC, note.[PaymentPayableNoteID] DESC"
    )
    query = text(
        f"""
        SELECT TOP ({safe_limit})
            note.[PaymentPayableNoteID] AS [ChequeItemID],
            payment.[PaymentID] AS [DocumentID],
            payment.[Number] AS [DocumentNumber],
            payment.[Date] AS [DocumentDate],
            note.[Amount],
            note.[DueDate],
            note.[SerialNumber],
            note.[Series],
            note.[AccountNumber],
            note.[BankRef],
            note.[BankAccountRef],
            note.[BankBranchName],
            note.[BankBranchCode],
            note.[CounterPartRef],
            note.[AccountRef],
            note.[ApproveState] AS [ChequeApproveState],
            note.[NoteType],
            note.[NormalORGuarantee],
            note.[CurrencyRef],
            master_note.[State] AS [ChequeState],
            bank.[Name] AS [BankName],
            counterpart.[Code] AS [CounterPartCode],
            counterpart.[Title] AS [CounterPartName],
            note.[Description] AS [ChequeDescription],
            payment.[Description] AS [DocumentDescription],
            status_link.[StatusDocumentID], status_link.[StatusDocumentNumber],
            status_link.[StatusDocumentDate], status_link.[StatusDocumentDescription]
        FROM RPA3.[PaymentPayableNote] AS note
        INNER JOIN RPA3.[Payment] AS payment
            ON payment.[PaymentID] = note.[PaymentRef]
        LEFT JOIN RPA3.[PayableNote] AS master_note
            ON master_note.[PayableNoteID] = note.[PayableNoteRef]
        LEFT JOIN RPA3.[Bank] AS bank
            ON bank.[BankID] = note.[BankRef]
        LEFT JOIN FIN3.[DL] AS counterpart
            ON counterpart.[DLID] = note.[CounterPartRef]
        OUTER APPLY (
            SELECT TOP (1)
                pnt.[DocumentRef] AS [StatusDocumentID],
                pnt.[DocumentNumber] AS [StatusDocumentNumber],
                COALESCE(pnt.[DocumentDate], pnt.[Date]) AS [StatusDocumentDate],
                pnt.[Description] AS [StatusDocumentDescription]
            FROM RPA3.[PayableNoteTransaction] AS pnt
            WHERE pnt.[PayableNoteRef] = master_note.[PayableNoteID]
              AND pnt.[State] = 11
              AND pnt.[DocumentItemType] IN (24, 26)
              AND pnt.[DocumentState] = 3
              AND (
                    ISNULL(pnt.[Description], N'') LIKE N'%تعیین وضعیت چک%'
                 OR ISNULL(pnt.[Description], N'') LIKE N'%تعین وضعیت چک%'
                 OR ISNULL(pnt.[Description], N'') LIKE N'%وصول چ%'
                 OR ISNULL(pnt.[Description], N'') LIKE N'%پرداخت چک%'
                 OR ISNULL(pnt.[Description], N'') LIKE N'%برداشت وجه چک%'
              )
            ORDER BY COALESCE(pnt.[DocumentDate], pnt.[Date]) DESC, pnt.[PayableNoteTransactionID] DESC
        ) AS status_link
        WHERE payment.[ApproveState] = 3
          AND note.[NoteType] = 1
          AND COALESCE(
                master_note.[NormalORGuarantee],
                note.[NormalORGuarantee]
              ) = 1
          AND ISNULL(master_note.[Description], N'') NOT LIKE N'%ضمانت%'
          AND ISNULL(master_note.[Description], N'') NOT LIKE N'%تضمین%'
          AND ISNULL(master_note.[Description], N'') NOT LIKE N'%حسن انجام%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%ضمانت%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%تضمین%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%حسن انجام%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%سهامدار%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%سود سهام%'
          AND ISNULL(master_note.[Description], N'') NOT LIKE N'%سهامدار%'
          AND ISNULL(master_note.[Description], N'') NOT LIKE N'%سود سهام%'
          AND note.[DueDate] >= DATEADD(day, -20, CAST(GETDATE() AS date))
          {primary_future_filter}
        ORDER BY {primary_order}
        """
    )

    primary_error: str | None = None
    with _engine(engine).connect() as connection:
        try:
            rows = connection.execute(query).mappings().all()
        except Exception as exc:
            # Some Rahkaran installations do not retain every issued cheque in
            # PaymentPayableNote, while the authoritative master PayableNote
            # still contains it. Keep the dashboard available through the
            # master-table fallback and expose the source in the response.
            primary_error = str(exc)
            rows = []

        source = "payment_document"
        if not rows:
            source = "payable_note_master_fallback"
            fallback_future_filter = (
                "AND master_note.[DueDate] >= CAST(GETDATE() AS date)"
                if upcoming_only else ""
            )
            fallback_query = text(
                f"""
                SELECT TOP ({safe_limit})
                    master_note.[PayableNoteID] AS [ChequeItemID],
                    CAST(NULL AS bigint) AS [DocumentID],
                    CAST(NULL AS bigint) AS [DocumentNumber],
                    master_note.[AgreementDate] AS [DocumentDate],
                    master_note.[Amount],
                    master_note.[DueDate],
                    master_note.[SerialNumber],
                    master_note.[Series],
                    master_note.[AccountNumber],
                    master_note.[BankRef],
                    master_note.[BankAccountRef],
                    master_note.[BankBranchName],
                    master_note.[BankBranchCode],
                    master_note.[CounterPartRef],
                    master_note.[AccountRef],
                    CAST(NULL AS int) AS [ChequeApproveState],
                    master_note.[NoteType],
                    master_note.[NormalORGuarantee],
                    master_note.[CurrencyRef],
                    master_note.[State] AS [ChequeState],
                    bank.[Name] AS [BankName],
                    counterpart.[Code] AS [CounterPartCode],
                    counterpart.[Title] AS [CounterPartName],
                    master_note.[Description] AS [ChequeDescription],
                    CAST(NULL AS nvarchar(1024)) AS [DocumentDescription],
                    status_link.[StatusDocumentID], status_link.[StatusDocumentNumber],
                    status_link.[StatusDocumentDate], status_link.[StatusDocumentDescription]
                FROM RPA3.[PayableNote] AS master_note
                LEFT JOIN RPA3.[Bank] AS bank
                    ON bank.[BankID] = master_note.[BankRef]
                LEFT JOIN FIN3.[DL] AS counterpart
                    ON counterpart.[DLID] = master_note.[CounterPartRef]
                OUTER APPLY (
                    SELECT TOP (1)
                        pnt.[DocumentRef] AS [StatusDocumentID],
                        pnt.[DocumentNumber] AS [StatusDocumentNumber],
                        COALESCE(pnt.[DocumentDate], pnt.[Date]) AS [StatusDocumentDate],
                        pnt.[Description] AS [StatusDocumentDescription]
                    FROM RPA3.[PayableNoteTransaction] AS pnt
                    WHERE pnt.[PayableNoteRef] = master_note.[PayableNoteID]
                      AND pnt.[State] = 11
                      AND pnt.[DocumentItemType] IN (24, 26)
                      AND pnt.[DocumentState] = 3
                      AND (
                            ISNULL(pnt.[Description], N'') LIKE N'%تعیین وضعیت چک%'
                         OR ISNULL(pnt.[Description], N'') LIKE N'%تعین وضعیت چک%'
                         OR ISNULL(pnt.[Description], N'') LIKE N'%وصول چ%'
                         OR ISNULL(pnt.[Description], N'') LIKE N'%پرداخت چک%'
                         OR ISNULL(pnt.[Description], N'') LIKE N'%برداشت وجه چک%'
                      )
                    ORDER BY COALESCE(pnt.[DocumentDate], pnt.[Date]) DESC, pnt.[PayableNoteTransactionID] DESC
                ) AS status_link
                WHERE master_note.[NoteType] = 1
                  AND master_note.[NormalORGuarantee] = 1
                  AND ISNULL(master_note.[Description], N'') NOT LIKE N'%ضمانت%'
                  AND ISNULL(master_note.[Description], N'') NOT LIKE N'%تضمین%'
                  AND ISNULL(master_note.[Description], N'') NOT LIKE N'%حسن انجام%'
                  AND ISNULL(master_note.[Description], N'') NOT LIKE N'%سهامدار%'
                  AND ISNULL(master_note.[Description], N'') NOT LIKE N'%سود سهام%'
                  AND master_note.[State] = 11
                  AND master_note.[DueDate] >= DATEADD(day, -20, CAST(GETDATE() AS date))
                  {fallback_future_filter}
                ORDER BY
                    CASE WHEN master_note.[DueDate] IS NULL THEN 1 ELSE 0 END,
                    master_note.[DueDate] ASC,
                    master_note.[PayableNoteID] DESC
                """
            )
            rows = connection.execute(fallback_query).mappings().all()

    cheques = [
        {
            "cheque_item_id": row["ChequeItemID"],
            "document_id": row["DocumentID"],
            "document_number": row["DocumentNumber"],
            "document_date": _as_datetime(row["DocumentDate"]),
            "document_date_jalali": _as_jalali_date(row["DocumentDate"]),
            "amount": _as_number(row["Amount"]),
            "due_date": _as_datetime(row["DueDate"]),
            "due_date_jalali": _as_jalali_date(row["DueDate"]),
            "serial_number": row["SerialNumber"],
            "series": row["Series"],
            "account_number": row["AccountNumber"],
            "bank_ref": row["BankRef"],
            "bank_account_ref": row["BankAccountRef"],
            "bank_branch_name": row["BankBranchName"],
            "bank_branch_code": row["BankBranchCode"],
            "counterpart_ref": row["CounterPartRef"],
            "account_ref": row["AccountRef"],
            "cheque_approve_state": row["ChequeApproveState"],
            "note_type": row["NoteType"],
            "normal_or_guarantee": row["NormalORGuarantee"],
            "currency_ref": row["CurrencyRef"],
            "currency_name": _currency_name(row["CurrencyRef"]),
            "state_code": row["ChequeState"],
            "state": _cheque_state_label(
                "issued",
                row["ChequeState"],
            ),
            "state_label": _issued_posting_status(row["ChequeState"], row.get("StatusDocumentDescription"))[1],
            "is_posted": _issued_posting_status(row["ChequeState"], row.get("StatusDocumentDescription"))[0],
            "cashflow_open": (not _issued_posting_status(row["ChequeState"], row.get("StatusDocumentDescription"))[0] and row["ChequeState"] == 11),
            "status_document_id": row.get("StatusDocumentID"),
            "status_document_number": row.get("StatusDocumentNumber"),
            "status_document_date": _as_datetime(row.get("StatusDocumentDate")),
            "status_document_date_jalali": _as_jalali_date(row.get("StatusDocumentDate")),
            "status_document_description": row.get("StatusDocumentDescription"),
            "bank_name": row["BankName"],
            "counterpart_code": row["CounterPartCode"],
            "counterpart_name": row["CounterPartName"],
            "description": (
                row["ChequeDescription"]
                or row["DocumentDescription"]
            ),
        }
        for row in rows
    ]

    return {
        "status": "success",
        "cheque_type": "issued",
        "approve_state_filter": 3,
        "note_type_filter": 1,
        "normal_or_guarantee_filter": 1,
        "guarantee_cheques_excluded": True,
        "shareholder_paid_cheques_excluded": True,
        "shareholder_exclusion_rules": ["Description does not contain سهامدار or سود سهام"],
        "data_source": source,
        "fallback_used": source != "payment_document",
        "primary_query_error": primary_error,
        "cheques": cheques,
        "returned_count": len(cheques),
        "returned_total_amount": sum(
            cheque["amount"] for cheque in cheques
        ),
    }




def get_received_cheque_sql_scope(engine: Engine | None = None) -> dict[str, Any]:
    """Read-only SQL coverage check for Rahkaran received cheques.

    This endpoint intentionally applies NO due-date horizon so Finance can see
    whether SQL itself contains cheques beyond the dates currently visible in UI.
    """
    current_status_apply = current_received_status_apply("note", "current_status")
    open_expression = received_cheque_open_expression("note", "current_status")
    query = text(
        f"""
        SELECT
            MIN(CAST(note.[DueDate] AS date)) AS [MinDueDate],
            MAX(CAST(note.[DueDate] AS date)) AS [MaxDueDate],
            COUNT_BIG(*) AS [AllNormalChequeCount],
            COALESCE(SUM(note.[Amount]), 0) AS [AllNormalChequeAmount],
            SUM(CASE WHEN {open_expression} THEN 1 ELSE 0 END) AS [OpenChequeCount],
            COALESCE(SUM(CASE WHEN {open_expression} THEN note.[Amount] ELSE 0 END), 0) AS [OpenChequeAmount]
        FROM RPA3.[ReceivableNote] AS note
        {current_status_apply}
        WHERE note.[NoteType] = 1
          AND note.[NormalORGuarantee] = 1
          AND ISNULL(note.[Description], N'') NOT LIKE N'%ضمانت%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%تضمین%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%حسن انجام%'
        """
    )
    monthly_query = text(
        f"""
        SELECT
            YEAR(note.[DueDate]) AS [DueYear],
            MONTH(note.[DueDate]) AS [DueMonth],
            COUNT_BIG(*) AS [OpenChequeCount],
            COALESCE(SUM(note.[Amount]), 0) AS [OpenChequeAmount]
        FROM RPA3.[ReceivableNote] AS note
        {current_status_apply}
        WHERE note.[NoteType] = 1
          AND note.[NormalORGuarantee] = 1
          AND note.[DueDate] IS NOT NULL
          AND ISNULL(note.[Description], N'') NOT LIKE N'%ضمانت%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%تضمین%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%حسن انجام%'
          AND {open_expression}
        GROUP BY YEAR(note.[DueDate]), MONTH(note.[DueDate])
        ORDER BY YEAR(note.[DueDate]), MONTH(note.[DueDate])
        """
    )
    with _engine(engine).connect() as connection:
        summary = connection.execute(query).mappings().one()
        monthly = connection.execute(monthly_query).mappings().all()
    return {
        "status": "success",
        "source_system": "rahkaran_sql",
        "date_horizon": "unbounded",
        "min_due_date": _as_datetime(summary["MinDueDate"]),
        "max_due_date": _as_datetime(summary["MaxDueDate"]),
        "all_normal_cheque_count": int(summary["AllNormalChequeCount"] or 0),
        "all_normal_cheque_amount": _as_number(summary["AllNormalChequeAmount"]),
        "open_cheque_count": int(summary["OpenChequeCount"] or 0),
        "open_cheque_amount": _as_number(summary["OpenChequeAmount"]),
        "open_due_months": [
            {
                "year": int(row["DueYear"]),
                "month": int(row["DueMonth"]),
                "count": int(row["OpenChequeCount"] or 0),
                "amount": _as_number(row["OpenChequeAmount"]),
            }
            for row in monthly
        ],
    }

def get_received_cheque_current_state_quality(engine: Engine | None = None) -> dict[str, Any]:
    """Compare legacy master-state filtering with Rahkaran's latest operation state.

    This is intentionally read-only and exists so Finance can audit exactly why
    a cheque entered or left the open portfolio after V115.
    """

    current_status_apply = current_received_status_apply("note", "current_status")
    effective_state = current_received_state_expr("note", "current_status")
    open_expression = received_cheque_open_expression("note", "current_status")
    query = text(
        f"""
        WITH CurrentRows AS (
            SELECT
                note.[ReceivableNoteID] AS [ChequeID],
                note.[SerialNumber],
                note.[SayadNumber],
                note.[Amount],
                note.[DueDate],
                note.[State] AS [MasterState],
                {effective_state} AS [EffectiveState],
                current_status.[CurrentStatusDescription],
                current_status.[CurrentStatusDate],
                current_status.[CurrentStatusDocumentDate],
                current_status.[CurrentStatusTransactionID],
                CASE WHEN note.[State] IN (1, 2) THEN 1 ELSE 0 END AS [LegacyOpen],
                CASE WHEN {open_expression} THEN 1 ELSE 0 END AS [BusinessOpen]
            FROM RPA3.[ReceivableNote] AS note
            {current_status_apply}
            WHERE note.[NoteType] = 1
              AND note.[NormalORGuarantee] = 1
              AND ISNULL(note.[Description], N'') NOT LIKE N'%ضمانت%'
              AND ISNULL(note.[Description], N'') NOT LIKE N'%تضمین%'
              AND ISNULL(note.[Description], N'') NOT LIKE N'%حسن انجام%'
        )
        SELECT
            COUNT_BIG(*) AS [TotalNormalChequeCount],
            SUM(CASE WHEN [LegacyOpen] = 1 THEN 1 ELSE 0 END) AS [LegacyOpenCount],
            COALESCE(SUM(CASE WHEN [LegacyOpen] = 1 THEN [Amount] ELSE 0 END), 0) AS [LegacyOpenAmount],
            SUM(CASE WHEN [BusinessOpen] = 1 THEN 1 ELSE 0 END) AS [BusinessOpenCount],
            COALESCE(SUM(CASE WHEN [BusinessOpen] = 1 THEN [Amount] ELSE 0 END), 0) AS [BusinessOpenAmount],
            SUM(CASE WHEN [LegacyOpen] = 1 AND [BusinessOpen] = 0 THEN 1 ELSE 0 END) AS [LegacyOnlyCount],
            COALESCE(SUM(CASE WHEN [LegacyOpen] = 1 AND [BusinessOpen] = 0 THEN [Amount] ELSE 0 END), 0) AS [LegacyOnlyAmount],
            SUM(CASE WHEN [LegacyOpen] = 0 AND [BusinessOpen] = 1 THEN 1 ELSE 0 END) AS [BusinessOnlyCount],
            COALESCE(SUM(CASE WHEN [LegacyOpen] = 0 AND [BusinessOpen] = 1 THEN [Amount] ELSE 0 END), 0) AS [BusinessOnlyAmount]
        FROM CurrentRows
        """
    )
    mismatch_query = text(
        f"""
        SELECT TOP (200)
            note.[ReceivableNoteID] AS [ChequeID],
            note.[SerialNumber],
            note.[SayadNumber],
            note.[Amount],
            note.[DueDate],
            note.[State] AS [MasterState],
            {effective_state} AS [EffectiveState],
            current_status.[CurrentStatusDescription],
            current_status.[CurrentStatusDate],
            current_status.[CurrentStatusDocumentDate],
            current_status.[CurrentStatusTransactionID],
            CASE WHEN note.[State] IN (1, 2) THEN 1 ELSE 0 END AS [LegacyOpen],
            CASE WHEN {open_expression} THEN 1 ELSE 0 END AS [BusinessOpen]
        FROM RPA3.[ReceivableNote] AS note
        {current_status_apply}
        WHERE note.[NoteType] = 1
          AND note.[NormalORGuarantee] = 1
          AND ISNULL(note.[Description], N'') NOT LIKE N'%ضمانت%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%تضمین%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%حسن انجام%'
          AND (
                (note.[State] IN (1, 2) AND NOT {open_expression})
                OR
                (note.[State] NOT IN (1, 2) AND {open_expression})
              )
        ORDER BY
            COALESCE(current_status.[CurrentStatusDate], current_status.[CurrentStatusDocumentDate], note.[DueDate]) DESC,
            note.[ReceivableNoteID] DESC
        """
    )
    with _engine(engine).connect() as connection:
        summary = connection.execute(query).mappings().one()
        mismatches = connection.execute(mismatch_query).mappings().all()

    def money(value: Any) -> float:
        return _as_number(value)

    return {
        "status": "success",
        "definition": "سبد باز = State اصلی ۱/۲؛ آخرین عملیات تأییدشده فقط برای حذف وصول/مسترد/واخواست/واگذاری به غیر",
        "legacy_definition": "ReceivableNote.State IN (1,2)",
        "total_normal_cheque_count": int(summary["TotalNormalChequeCount"] or 0),
        "legacy_open": {
            "count": int(summary["LegacyOpenCount"] or 0),
            "amount": money(summary["LegacyOpenAmount"]),
        },
        "business_open": {
            "count": int(summary["BusinessOpenCount"] or 0),
            "amount": money(summary["BusinessOpenAmount"]),
        },
        "legacy_only": {
            "count": int(summary["LegacyOnlyCount"] or 0),
            "amount": money(summary["LegacyOnlyAmount"]),
        },
        "business_only": {
            "count": int(summary["BusinessOnlyCount"] or 0),
            "amount": money(summary["BusinessOnlyAmount"]),
        },
        "mismatches": [
            {
                "cheque_id": row["ChequeID"],
                "serial_number": row["SerialNumber"],
                "sayad_number": row["SayadNumber"],
                "amount": money(row["Amount"]),
                "due_date": _as_datetime(row["DueDate"]),
                "due_date_jalali": _as_jalali_date(row["DueDate"]),
                "master_state": row["MasterState"],
                "effective_state": row["EffectiveState"],
                "current_status_description": row["CurrentStatusDescription"],
                "current_status_date": _as_datetime(row["CurrentStatusDate"]),
                "current_status_date_jalali": _as_jalali_date(row["CurrentStatusDate"]),
                "legacy_open": bool(row["LegacyOpen"]),
                "business_open": bool(row["BusinessOpen"]),
            }
            for row in mismatches
        ],
        "mismatch_count_returned": len(mismatches),
    }

def get_cheque_state_quality(engine: Engine | None = None) -> dict[str, Any]:
    """Return auditable cheque-state counts and the verified received pass rate.

    Issued-state business meanings vary by Rahkaran configuration. We therefore
    expose their distribution without silently treating unknown states as paid.
    """
    received_query = text(
        """
        SELECT
            COUNT_BIG(*) AS [TotalCount],
            SUM(CASE WHEN [State] = 3 THEN 1 ELSE 0 END) AS [CollectedCount],
            SUM(CASE WHEN [State] = 3 THEN [Amount] ELSE 0 END) AS [CollectedAmount],
            SUM([Amount]) AS [TotalAmount]
        FROM RPA3.[ReceivableNote]
        WHERE [NoteType] = 1
          AND [NormalORGuarantee] = 1
          AND ISNULL([Description], N'') NOT LIKE N'%ضمانت%'
          AND ISNULL([Description], N'') NOT LIKE N'%تضمین%'
          AND ISNULL([Description], N'') NOT LIKE N'%حسن انجام%'
        """
    )
    issued_query = text(
        """
        SELECT
            [State] AS [StateCode],
            COUNT_BIG(*) AS [ChequeCount],
            SUM([Amount]) AS [ChequeAmount],
            SUM(CASE WHEN [DueDate] < CAST(GETDATE() AS date) THEN 1 ELSE 0 END) AS [PastDueCount],
            SUM(CASE WHEN [DueDate] < CAST(GETDATE() AS date) THEN [Amount] ELSE 0 END) AS [PastDueAmount]
        FROM RPA3.[PayableNote]
        WHERE [NoteType] = 1
          AND [NormalORGuarantee] = 1
          AND ISNULL([Description], N'') NOT LIKE N'%ضمانت%'
          AND ISNULL([Description], N'') NOT LIKE N'%تضمین%'
          AND ISNULL([Description], N'') NOT LIKE N'%حسن انجام%'
        GROUP BY [State]
        ORDER BY [State]
        """
    )
    with _engine(engine).connect() as connection:
        received = connection.execute(received_query).mappings().one()
        issued = connection.execute(issued_query).mappings().all()
    total_count = int(received["TotalCount"] or 0)
    total_amount = _as_number(received["TotalAmount"])
    collected_count = int(received["CollectedCount"] or 0)
    collected_amount = _as_number(received["CollectedAmount"])
    issued_rows = [
        {
            "state_code": row["StateCode"],
            "cheque_count": int(row["ChequeCount"] or 0),
            "cheque_amount": _as_number(row["ChequeAmount"]),
            "past_due_count": int(row["PastDueCount"] or 0),
            "past_due_amount": _as_number(row["PastDueAmount"]),
            "verified_business_label": _cheque_state_label("issued", row["StateCode"]),
            "posting_status_fa": _issued_posting_status(row["StateCode"])[1],
            "means_accounting_posted": _issued_posting_status(row["StateCode"])[0],
        }
        for row in issued
    ]
    paid_states = {
        int(value.strip())
        for value in settings.treasury_issued_paid_states.split(",")
        if value.strip().lstrip("-").isdigit()
    }
    # Denominator contains only valid matured cheques: still active (11) plus
    # finance-confirmed cleared states. Restored/cancelled and reversal states
    # must not dilute the real payment rate.
    eligible_states = {11} | paid_states
    due_count = sum(item["past_due_count"] for item in issued_rows if item["state_code"] in eligible_states)
    due_amount = sum(item["past_due_amount"] for item in issued_rows if item["state_code"] in eligible_states)
    paid_count = sum(item["past_due_count"] for item in issued_rows if item["state_code"] in paid_states)
    paid_amount = sum(item["past_due_amount"] for item in issued_rows if item["state_code"] in paid_states)
    issued_actual_rate = None
    if paid_states:
        issued_actual_rate = {
            "definition": "سندخورده/برداشت‌شده (بلندمدت → روز؛ configured cleared states) / چک‌های سررسیدشده معتبر",
            "configured_paid_states": sorted(paid_states),
            "eligible_denominator_states": sorted(eligible_states),
            "due_count": due_count,
            "paid_count": paid_count,
            "count_rate_percent": paid_count / due_count * 100 if due_count else 0,
            "due_amount": due_amount,
            "paid_amount": paid_amount,
            "amount_rate_percent": paid_amount / due_amount * 100 if due_amount else 0,
        }
    return {
        "status": "success",
        "received_verified": {
            "definition": "RPA3.ReceivableNote State=3 / normal received cheque notes; guarantee cheques excluded",
            "total_count": total_count,
            "collected_count": collected_count,
            "count_rate_percent": collected_count / total_count * 100 if total_count else 0,
            "total_amount": total_amount,
            "collected_amount": collected_amount,
            "amount_rate_percent": collected_amount / total_amount * 100 if total_amount else 0,
        },
        "issued_state_distribution": issued_rows,
        "issued_actual_payment_rate": issued_actual_rate,
        "issued_rate_limitation": None if paid_states else "کد وضعیت تسویه‌شده را پس از تأیید مالی در TREASURY_ISSUED_PAID_STATES فایل .env وارد کنید.",
    }


def get_cheque_due_report(
    cheque_type: str,
    days: int = 30,
    overdue: bool = False,
    limit: int = 100,
    engine: Engine | None = None,
) -> dict[str, Any]:
    """گزارش چک‌های فعال سررسیدشونده یا سررسیدگذشته."""

    normalized_type = cheque_type.strip().lower()
    configurations = {
        "received": {
            "table": "ReceivableNote",
            "id_column": "ReceivableNoteID",
            "sayad_sql": "note.[SayadNumber]",
            "active_states": (1, 2),
        },
        "issued": {
            "table": "PayableNote",
            "id_column": "PayableNoteID",
            "sayad_sql": "CAST(NULL AS nvarchar(32))",
            "active_states": (11,),
        },
    }
    configuration = configurations.get(normalized_type)
    if configuration is None:
        raise ValueError("cheque_type must be received or issued")

    safe_days = max(1, min(int(days), 365))
    safe_limit = max(1, min(int(limit), 500))
    state_sql = ", ".join(
        str(state) for state in configuration["active_states"]
    )
    if normalized_type == "received":
        current_status_apply = current_received_status_apply("note", "current_status")
        effective_state = current_received_state_expr("note", "current_status")
        state_select_sql = f"""
            note.[State] AS [MasterChequeState],
            {effective_state} AS [ChequeState],
            current_status.[CurrentStatusDescription],
            current_status.[CurrentStatusDate],
            current_status.[CurrentStatusDocumentDate],
            current_status.[CurrentStatusTransactionID],
        """
        active_state_predicate = received_cheque_open_predicate("note", "current_status")
    else:
        current_status_apply = ""
        state_select_sql = "note.[State] AS [ChequeState],"
        active_state_predicate = f"AND note.[State] IN ({state_sql})"
    if overdue:
        due_filter = """
            CAST(note.[DueDate] AS date) >=
                DATEADD(day, -:days, CAST(GETDATE() AS date))
            AND CAST(note.[DueDate] AS date) < CAST(GETDATE() AS date)
        """
    else:
        due_filter = """
            CAST(note.[DueDate] AS date) >= CAST(GETDATE() AS date)
            AND CAST(note.[DueDate] AS date) <=
                DATEADD(day, :days, CAST(GETDATE() AS date))
        """

    query = text(
        f"""
        SELECT TOP ({safe_limit})
            note.[{configuration['id_column']}] AS [ChequeID],
            {state_select_sql}
            note.[Amount],
            note.[DueDate],
            DATEDIFF(
                day,
                CAST(GETDATE() AS date),
                CAST(note.[DueDate] AS date)
            ) AS [DaysUntilDue],
            note.[SerialNumber],
            note.[Series],
            {configuration['sayad_sql']} AS [SayadNumber],
            note.[AccountNumber],
            note.[BankRef],
            bank.[Name] AS [BankName],
            note.[BankBranchName],
            note.[BankBranchCode],
            note.[CounterPartRef],
            counterpart.[Code] AS [CounterPartCode],
            counterpart.[Title] AS [CounterPartName],
            note.[AccountRef],
            note.[CurrencyRef],
            note.[NormalORGuarantee],
            note.[Description]
        FROM RPA3.[{configuration['table']}] AS note
        LEFT JOIN RPA3.[Bank] AS bank
            ON bank.[BankID] = note.[BankRef]
        LEFT JOIN FIN3.[DL] AS counterpart
            ON counterpart.[DLID] = note.[CounterPartRef]
        {current_status_apply}
        WHERE note.[NoteType] = 1
          AND note.[NormalORGuarantee] = 1
          AND ISNULL(note.[Description], N'') NOT LIKE N'%ضمانت%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%تضمین%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%حسن انجام%'
          {active_state_predicate}
          AND note.[DueDate] IS NOT NULL
          AND {due_filter}
        ORDER BY note.[DueDate], note.[{configuration['id_column']}]
        """
    )

    aggregate_query = text(
        f"""
        SELECT
            COUNT_BIG(*) AS [TotalCount],
            COALESCE(SUM(note.[Amount]), 0) AS [TotalAmount]
        FROM RPA3.[{configuration['table']}] AS note
        {current_status_apply}
        WHERE note.[NoteType] = 1
          AND note.[NormalORGuarantee] = 1
          AND ISNULL(note.[Description], N'') NOT LIKE N'%ضمانت%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%تضمین%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%حسن انجام%'
          {active_state_predicate}
          AND note.[DueDate] IS NOT NULL
          AND {due_filter}
        """
    )

    parameters = {"days": safe_days}
    with _engine(engine).connect() as connection:
        rows = connection.execute(query, parameters).mappings().all()
        aggregate = connection.execute(
            aggregate_query,
            parameters,
        ).mappings().one()

    cheques = [
        _master_cheque_payload(row, normalized_type)
        for row in rows
    ]
    total_count = int(aggregate["TotalCount"] or 0)
    total_amount = _as_number(aggregate["TotalAmount"])
    return {
        "status": "success",
        "report_type": "overdue" if overdue else "upcoming",
        "cheque_type": normalized_type,
        "as_of_date": date.today().isoformat(),
        "as_of_date_jalali": _as_jalali_date(date.today()),
        "days": safe_days,
        "date_window": (
            f"previous_{safe_days}_days"
            if overdue
            else f"next_{safe_days}_days"
        ),
        "active_state_filter": (
            "latest_positive_receivable_note_transaction_business_open"
            if normalized_type == "received"
            else list(configuration["active_states"])
        ),
        "normal_or_guarantee_filter": 1,
        "cheques": cheques,
        "returned_count": len(cheques),
        "returned_total_amount": sum(
            cheque["amount"] for cheque in cheques
        ),
        "total_count": total_count,
        "total_amount": total_amount,
        "is_truncated": total_count > len(cheques),
    }


def get_received_cheques_by_status(
    cheque_status: str,
    limit: int = 100,
    engine: Engine | None = None,
) -> dict[str, Any]:
    """چک‌های دریافتی با وضعیت قطعی وصول‌شده یا واخواست‌شده."""

    normalized_status = cheque_status.strip().lower()
    state_by_status = {"collected": 3, "protested": 4}
    state_code = state_by_status.get(normalized_status)
    if state_code is None:
        raise ValueError("cheque_status must be collected or protested")

    safe_limit = max(1, min(int(limit), 500))
    query = text(
        f"""
        SELECT TOP ({safe_limit})
            note.[ReceivableNoteID] AS [ChequeID],
            note.[State] AS [ChequeState],
            note.[Amount],
            note.[DueDate],
            DATEDIFF(
                day,
                CAST(GETDATE() AS date),
                CAST(note.[DueDate] AS date)
            ) AS [DaysUntilDue],
            note.[SerialNumber],
            note.[Series],
            note.[SayadNumber],
            note.[AccountNumber],
            note.[BankRef],
            bank.[Name] AS [BankName],
            note.[BankBranchName],
            note.[BankBranchCode],
            note.[CounterPartRef],
            counterpart.[Code] AS [CounterPartCode],
            counterpart.[Title] AS [CounterPartName],
            note.[AccountRef],
            note.[CurrencyRef],
            note.[NormalORGuarantee],
            note.[Description]
        FROM RPA3.[ReceivableNote] AS note
        LEFT JOIN RPA3.[Bank] AS bank
            ON bank.[BankID] = note.[BankRef]
        LEFT JOIN FIN3.[DL] AS counterpart
            ON counterpart.[DLID] = note.[CounterPartRef]
        WHERE note.[NoteType] = 1
          AND note.[NormalORGuarantee] = 1
          AND ISNULL(note.[Description], N'') NOT LIKE N'%ضمانت%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%تضمین%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%حسن انجام%'
          AND note.[State] = :state_code
        ORDER BY note.[DueDate] DESC, note.[ReceivableNoteID] DESC
        """
    )

    with _engine(engine).connect() as connection:
        rows = connection.execute(
            query,
            {"state_code": state_code},
        ).mappings().all()

    cheques = [
        _master_cheque_payload(row, "received")
        for row in rows
    ]
    return {
        "status": "success",
        "report_type": "received_cheque_status",
        "cheque_status": normalized_status,
        "state_filter": state_code,
        "normal_or_guarantee_filter": 1,
        "guarantee_cheques_excluded": True,
        "cheques": cheques,
        "returned_count": len(cheques),
        "returned_total_amount": sum(
            cheque["amount"] for cheque in cheques
        ),
    }


def _owner_name_matches_debtor(
    debtor_name: str | None,
    cheque_account_text: str | None,
) -> bool | None:
    """Informational only; allocation never depends on this fuzzy comparison."""

    account_text = _normalize_account_search(cheque_account_text or "")
    debtor_text = _normalize_account_search(debtor_name or "")
    if not account_text or not debtor_text:
        return None
    ignored = {
        "شرکت",
        "بازرگانی",
        "دریافتنی",
        "ریال",
        "اشخاص",
        "تجاری",
    }
    debtor_tokens = {
        token
        for token in debtor_text.split()
        if len(token) >= 3 and token not in ignored
    }
    if not debtor_tokens:
        return None
    return any(token in account_text for token in debtor_tokens)


def _settlement_remaining(
    opening_debt: Decimal,
    settlement_amount: Decimal,
) -> tuple[Decimal, Decimal]:
    difference = opening_debt - settlement_amount
    return max(difference, Decimal("0")), max(-difference, Decimal("0"))


def _build_customer_cheque_settlement_payload(
    *,
    cheque_rows: list[Any],
    account_rows: list[Any],
    counterpart_ref: int,
    counterpart_code: str | None,
    counterpart_name: str | None,
    opening_debt: Decimal | None,
) -> dict[str, Any]:
    """Build safe many-cheques-to-one-debtor totals from explicit ERP refs."""

    cheques: list[dict[str, Any]] = []
    state_amounts: dict[str, Decimal] = {
        "registered": Decimal("0"),
        "in_collection": Decimal("0"),
        "collected": Decimal("0"),
        "protested": Decimal("0"),
        "delivered_to_counterparty": Decimal("0"),
        "unknown": Decimal("0"),
    }
    state_counts = {state: 0 for state in state_amounts}

    for row in cheque_rows:
        state_code = row["ChequeState"]
        state = _cheque_state_label("received", state_code)
        amount = Decimal(str(row["Amount"] or 0))
        state_amounts.setdefault(state, Decimal("0"))
        state_counts.setdefault(state, 0)
        state_amounts[state] += amount
        state_counts[state] += 1
        allocation_source = (
            "receipt_note_counterpart"
            if row["NoteCounterPartRef"] is not None
            else "receipt_document_counterpart"
        )
        owner_matches = _owner_name_matches_debtor(
            counterpart_name,
            row["AccountNumber"],
        )
        cheques.append(
            {
                "cheque_id": row["ChequeID"],
                "cheque_item_id": row["ChequeItemID"],
                "document_id": row["DocumentID"],
                "document_number": row["DocumentNumber"],
                "document_date": _as_datetime(row["DocumentDate"]),
                "document_date_jalali": _as_jalali_date(row["DocumentDate"]),
                "amount": _as_number(amount),
                "due_date": _as_datetime(row["DueDate"]),
                "due_date_jalali": _as_jalali_date(row["DueDate"]),
                "serial_number": row["SerialNumber"],
                "series": row["Series"],
                "sayad_number": row["SayadNumber"],
                "cheque_account_text": row["AccountNumber"],
                "bank_name": row["BankName"],
                "state_code": state_code,
                "state": state,
                "allocated_debtor_ref": counterpart_ref,
                "allocated_debtor_name": counterpart_name,
                "allocation_source": allocation_source,
                "allocation_is_explicit": True,
                "owner_name_matches_debtor": owner_matches,
                "third_party_indicator": owner_matches is False,
                "description": (
                    row["ChequeDescription"]
                    or row["DocumentDescription"]
                ),
            }
        )

    received_total = sum(state_amounts.values(), Decimal("0"))
    collected_total = state_amounts.get("collected", Decimal("0"))
    pending_total = (
        state_amounts.get("registered", Decimal("0"))
        + state_amounts.get("in_collection", Decimal("0"))
    )
    provisional_settlement = collected_total + pending_total
    third_party_count = sum(
        1 for cheque in cheques if cheque["third_party_indicator"]
    )
    current_book_balance = sum(
        (Decimal(str(row["Balance"] or 0)) for row in account_rows),
        Decimal("0"),
    )

    debt_calculation: dict[str, Any] | None = None
    if opening_debt is not None:
        confirmed_remaining, confirmed_overpayment = _settlement_remaining(
            opening_debt,
            collected_total,
        )
        provisional_remaining, provisional_overpayment = _settlement_remaining(
            opening_debt,
            provisional_settlement,
        )
        debt_calculation = {
            "opening_debt": _as_number(opening_debt),
            "confirmed_settlement_collected_only": _as_number(collected_total),
            "confirmed_remaining_debt": _as_number(confirmed_remaining),
            "confirmed_overpayment": _as_number(confirmed_overpayment),
            "provisional_settlement_collected_and_pending": _as_number(
                provisional_settlement
            ),
            "provisional_remaining_debt": _as_number(provisional_remaining),
            "provisional_overpayment": _as_number(provisional_overpayment),
            "protested_cheques_are_excluded": True,
        }

    return {
        "status": "success",
        "report_type": "customer_cheque_settlement",
        "debtor": {
            "counterpart_ref": counterpart_ref,
            "code": counterpart_code,
            "name": counterpart_name,
        },
        "allocation_rule": (
            "CounterPartRef ردیف چک؛ در نبود آن CounterPartRef سند دریافت"
        ),
        "owner_name_is_not_used_for_allocation": True,
        "current_receivable_accounts": [
            {
                "account_id": row["AccountID"],
                "name": row["AccountName"],
                "number": row["AccountNumber"],
                "balance": _as_number(row["Balance"]),
                "currency_ref": row["CurrencyRef"],
                "currency_name": _currency_name(row["CurrencyRef"]),
            }
            for row in account_rows
        ],
        "current_book_balance_total": _as_number(current_book_balance),
        "summary": {
            "cheque_count": len(cheques),
            "received_cheque_total": _as_number(received_total),
            "collected_total": _as_number(collected_total),
            "pending_total": _as_number(pending_total),
            "protested_total": _as_number(
                state_amounts.get("protested", Decimal("0"))
            ),
            "third_party_indicator_count": third_party_count,
            "state_counts": state_counts,
            "state_amounts": {
                key: _as_number(value) for key, value in state_amounts.items()
            },
        },
        "debt_calculation": debt_calculation,
        "cheques": cheques,
        "returned_count": len(cheques),
    }


def search_received_cheque_debtors(
    name: str,
    limit: int = 20,
    engine: Engine | None = None,
) -> dict[str, Any]:
    """Search explicit cheque debtors/deliverers, not cheque account owners."""

    search_term = _normalize_account_search(name)
    if not search_term:
        return {"status": "success", "debtors": []}
    safe_limit = max(1, min(int(limit), 50))
    tokens = search_term.split()[:8]
    conditions = []
    parameters: dict[str, Any] = {}
    normalized_title_sql = """
        REPLACE(REPLACE(REPLACE(REPLACE(dl.[Title], N'آ', N'ا'),
        N'ي', N'ی'), N'ك', N'ک'), NCHAR(8204), N' ')
    """
    for index, token in enumerate(tokens):
        key = f"token_{index}"
        conditions.append(f"{normalized_title_sql} LIKE :{key}")
        parameters[key] = f"%{token}%"

    query = text(
        f"""
        SELECT TOP ({safe_limit})
            dl.[DLID] AS [CounterPartRef],
            dl.[Code] AS [CounterPartCode],
            dl.[Title] AS [CounterPartName],
            COUNT_BIG(*) AS [ChequeCount],
            SUM(note.[Amount]) AS [ChequeTotalAmount]
        FROM RPA3.[ReceiptReceivableNote] AS note
        INNER JOIN RPA3.[Receipt] AS receipt
            ON receipt.[ReceiptID] = note.[ReceiptRef]
        LEFT JOIN RPA3.[ReceivableNote] AS master_note
            ON master_note.[ReceivableNoteID] = note.[ReceivableNoteRef]
        INNER JOIN FIN3.[DL] AS dl
            ON dl.[DLID] = COALESCE(
                note.[CounterPartRef], receipt.[CounterPartRef]
            )
        WHERE receipt.[ApproveState] = 3
          AND receipt.[ItemType] = 1
          AND COALESCE(
                master_note.[NormalORGuarantee],
                note.[NormalORGuarantee]
              ) = 1
          AND ISNULL(master_note.[Description], N'') NOT LIKE N'%ضمانت%'
          AND ISNULL(master_note.[Description], N'') NOT LIKE N'%تضمین%'
          AND ISNULL(master_note.[Description], N'') NOT LIKE N'%حسن انجام%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%ضمانت%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%تضمین%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%حسن انجام%'
          AND {' AND '.join(conditions)}
        GROUP BY dl.[DLID], dl.[Code], dl.[Title]
        ORDER BY COUNT_BIG(*) DESC, dl.[Title]
        """
    )
    with _engine(engine).connect() as connection:
        rows = connection.execute(query, parameters).mappings().all()
    return {
        "status": "success",
        "debtors": [
            {
                "counterpart_ref": row["CounterPartRef"],
                "code": row["CounterPartCode"],
                "name": row["CounterPartName"],
                "cheque_count": row["ChequeCount"],
                "cheque_total_amount": _as_number(row["ChequeTotalAmount"]),
            }
            for row in rows
        ],
        "returned_count": len(rows),
    }


def get_customer_cheque_settlement(
    *,
    counterpart_ref: int,
    opening_debt: str | int | float | Decimal | None = None,
    limit: int = 500,
    engine: Engine | None = None,
) -> dict[str, Any]:
    """Report multiple own/third-party cheques allocated to one debtor."""

    safe_limit = max(1, min(int(limit), 1000))
    parsed_opening_debt: Decimal | None = None
    if opening_debt not in (None, ""):
        try:
            parsed_opening_debt = Decimal(str(opening_debt))
        except Exception as exc:
            raise ValueError("opening_debt must be a valid number") from exc
        if parsed_opening_debt < 0:
            raise ValueError("opening_debt cannot be negative")

    cheque_query = text(
        f"""
        SELECT TOP ({safe_limit})
            note.[ReceiptReceivableNoteID] AS [ChequeItemID],
            note.[ReceivableNoteRef] AS [ChequeID],
            receipt.[ReceiptID] AS [DocumentID],
            receipt.[Number] AS [DocumentNumber],
            receipt.[Date] AS [DocumentDate],
            note.[Amount], note.[DueDate], note.[SerialNumber], note.[Series],
            note.[SayadNumber], note.[AccountNumber],
            note.[CounterPartRef] AS [NoteCounterPartRef],
            receipt.[CounterPartRef] AS [DocumentCounterPartRef],
            master_note.[State] AS [ChequeState],
            bank.[Name] AS [BankName],
            dl.[Code] AS [CounterPartCode],
            dl.[Title] AS [CounterPartName],
            note.[Description] AS [ChequeDescription],
            receipt.[Description] AS [DocumentDescription]
        FROM RPA3.[ReceiptReceivableNote] AS note
        INNER JOIN RPA3.[Receipt] AS receipt
            ON receipt.[ReceiptID] = note.[ReceiptRef]
        LEFT JOIN RPA3.[ReceivableNote] AS master_note
            ON master_note.[ReceivableNoteID] = note.[ReceivableNoteRef]
        LEFT JOIN RPA3.[Bank] AS bank ON bank.[BankID] = note.[BankRef]
        INNER JOIN FIN3.[DL] AS dl
            ON dl.[DLID] = COALESCE(
                note.[CounterPartRef], receipt.[CounterPartRef]
            )
        WHERE receipt.[ApproveState] = 3
          AND receipt.[ItemType] = 1
          AND COALESCE(
                master_note.[NormalORGuarantee],
                note.[NormalORGuarantee]
              ) = 1
          AND ISNULL(master_note.[Description], N'') NOT LIKE N'%ضمانت%'
          AND ISNULL(master_note.[Description], N'') NOT LIKE N'%تضمین%'
          AND ISNULL(master_note.[Description], N'') NOT LIKE N'%حسن انجام%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%ضمانت%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%تضمین%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%حسن انجام%'
          AND COALESCE(note.[CounterPartRef], receipt.[CounterPartRef])
              = :counterpart_ref
          AND note.[CurrencyRef] = :currency_ref
        ORDER BY receipt.[Date] DESC, note.[ReceiptReceivableNoteID] DESC
        """
    )
    account_query = text(
        """
        SELECT
            account.[AccountID], account.[Name] AS [AccountName],
            account.[Number] AS [AccountNumber], account.[Balance],
            account.[CurrencyRef]
        FROM FIN3.[Account] AS account
        INNER JOIN FIN3.[DL] AS dl
            ON dl.[ReferenceID] = account.[PartyRef]
        WHERE dl.[DLID] = :counterpart_ref
          AND account.[CurrencyRef] = :currency_ref
          AND account.[Status] = 2
        ORDER BY account.[AccountID]
        """
    )
    parameters = {
        "counterpart_ref": int(counterpart_ref),
        "currency_ref": settings.treasury_operational_currency_ref,
    }
    with _engine(engine).connect() as connection:
        cheque_rows = connection.execute(
            cheque_query, parameters
        ).mappings().all()
        account_rows = connection.execute(
            account_query, parameters
        ).mappings().all()

    if not cheque_rows:
        return {
            "status": "not_found",
            "message": "برای این طرف حساب چک دریافتی تخصیص‌یافته پیدا نشد.",
            "counterpart_ref": int(counterpart_ref),
        }
    first = cheque_rows[0]
    return _build_customer_cheque_settlement_payload(
        cheque_rows=cheque_rows,
        account_rows=account_rows,
        counterpart_ref=int(counterpart_ref),
        counterpart_code=first["CounterPartCode"],
        counterpart_name=first["CounterPartName"],
        opening_debt=parsed_opening_debt,
    )


def get_receivable_report() -> dict[str, Any]:
    """تا زمان تأیید تعریف تجاری دریافتنی، عدد احتمالی تولید نمی‌کند."""

    return {
        "status": "not_ready",
        "message": (
            "تعریف و Query گزارش دریافتنی هنوز توسط واحد خزانه تأیید نشده است."
        ),
    }


def get_received_cheque_status_signatures(engine: Engine | None = None) -> dict[str, Any]:
    """SQL-only audit of Rahkaran current-status signatures.

    This does not use Excel as a data source. It exposes the fields that Rahkaran
    uses around the latest approved ReceivableNoteTransaction so the official
    report's «وضعیت فعلی» can be calibrated without guessing from master State.
    """
    current_status_apply = current_received_status_apply("note", "current_status")
    query = text(
        f"""
        SELECT
            note.[State] AS [MasterState],
            current_status.[CurrentChequeState],
            current_status.[CurrentDocumentItemType],
            CASE WHEN current_status.[CurrentBankAccountRef] IS NULL THEN 0 ELSE 1 END AS [HasBankAccount],
            CASE
                WHEN ISNULL(current_status.[CurrentStatusDescription], N'') LIKE N'%مأمور وصول%'
                  OR ISNULL(current_status.[CurrentStatusDescription], N'') LIKE N'%مامور وصول%' THEN N'collector_text'
                WHEN ISNULL(current_status.[CurrentStatusDescription], N'') LIKE N'%واخواست%' THEN N'protested_text'
                WHEN ISNULL(current_status.[CurrentStatusDescription], N'') LIKE N'%مسترد%'
                  OR ISNULL(current_status.[CurrentStatusDescription], N'') LIKE N'%استرداد%' THEN N'returned_text'
                WHEN ISNULL(current_status.[CurrentStatusDescription], N'') LIKE N'%وصول%'
                  OR ISNULL(current_status.[CurrentStatusDescription], N'') LIKE N'%نقد%' THEN N'collected_text'
                WHEN ISNULL(current_status.[CurrentStatusDescription], N'') LIKE N'%واگذار%'
                  AND ISNULL(current_status.[CurrentStatusDescription], N'') LIKE N'%غیر%' THEN N'transferred_to_third_party_text'
                WHEN ISNULL(current_status.[CurrentStatusDescription], N'') LIKE N'%برگ دریافت چک%'
                  OR ISNULL(current_status.[CurrentStatusDescription], N'') LIKE N'%برگه دریافت چک%' THEN N'receipt_sheet_text'
                WHEN NULLIF(LTRIM(RTRIM(ISNULL(current_status.[CurrentStatusDescription], N''))), N'') IS NULL THEN N'empty_text'
                ELSE N'other_text'
            END AS [DescriptionClass],
            COUNT_BIG(*) AS [ChequeCount],
            COALESCE(SUM(note.[Amount]), 0) AS [ChequeAmount],
            MIN(CAST(note.[DueDate] AS date)) AS [MinDueDate],
            MAX(CAST(note.[DueDate] AS date)) AS [MaxDueDate]
        FROM RPA3.[ReceivableNote] AS note
        {current_status_apply}
        WHERE note.[NoteType] = 1
          AND note.[NormalORGuarantee] = 1
          AND ISNULL(note.[Description], N'') NOT LIKE N'%ضمانت%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%تضمین%'
          AND ISNULL(note.[Description], N'') NOT LIKE N'%حسن انجام%'
        GROUP BY
            note.[State],
            current_status.[CurrentChequeState],
            current_status.[CurrentDocumentItemType],
            CASE WHEN current_status.[CurrentBankAccountRef] IS NULL THEN 0 ELSE 1 END,
            CASE
                WHEN ISNULL(current_status.[CurrentStatusDescription], N'') LIKE N'%مأمور وصول%'
                  OR ISNULL(current_status.[CurrentStatusDescription], N'') LIKE N'%مامور وصول%' THEN N'collector_text'
                WHEN ISNULL(current_status.[CurrentStatusDescription], N'') LIKE N'%واخواست%' THEN N'protested_text'
                WHEN ISNULL(current_status.[CurrentStatusDescription], N'') LIKE N'%مسترد%'
                  OR ISNULL(current_status.[CurrentStatusDescription], N'') LIKE N'%استرداد%' THEN N'returned_text'
                WHEN ISNULL(current_status.[CurrentStatusDescription], N'') LIKE N'%وصول%'
                  OR ISNULL(current_status.[CurrentStatusDescription], N'') LIKE N'%نقد%' THEN N'collected_text'
                WHEN ISNULL(current_status.[CurrentStatusDescription], N'') LIKE N'%واگذار%'
                  AND ISNULL(current_status.[CurrentStatusDescription], N'') LIKE N'%غیر%' THEN N'transferred_to_third_party_text'
                WHEN ISNULL(current_status.[CurrentStatusDescription], N'') LIKE N'%برگ دریافت چک%'
                  OR ISNULL(current_status.[CurrentStatusDescription], N'') LIKE N'%برگه دریافت چک%' THEN N'receipt_sheet_text'
                WHEN NULLIF(LTRIM(RTRIM(ISNULL(current_status.[CurrentStatusDescription], N''))), N'') IS NULL THEN N'empty_text'
                ELSE N'other_text'
            END
        ORDER BY [ChequeCount] DESC
        """
    )
    with _engine(engine).connect() as connection:
        rows = connection.execute(query).mappings().all()
    return {
        "status": "success",
        "source_system": "rahkaran_sql",
        "excel_used_as_data_source": False,
        "purpose": "calibrate_official_current_status",
        "signatures": [
            {
                "master_state": row["MasterState"],
                "transaction_state": row["CurrentChequeState"],
                "document_item_type": row["CurrentDocumentItemType"],
                "has_bank_account": bool(row["HasBankAccount"]),
                "description_class": row["DescriptionClass"],
                "count": int(row["ChequeCount"] or 0),
                "amount": _as_number(row["ChequeAmount"]),
                "min_due_date": _as_datetime(row["MinDueDate"]),
                "max_due_date": _as_datetime(row["MaxDueDate"]),
            }
            for row in rows
        ],
    }


# V122 diagnostic: Excel is validation only; dashboard remains SQL-only.
def compare_received_open_sql_with_rahkaran_excel(excel_bytes: bytes, engine: Engine | None = None) -> dict[str, Any]:
    from io import BytesIO
    from openpyxl import load_workbook

    allowed_statuses = {"نزد بانک", "نزد صندوق", "نزد مأمور وصول", "نزد مامور وصول"}
    def norm(v: Any) -> str:
        if v is None: return ""
        if isinstance(v, float) and v.is_integer(): v = int(v)
        x = str(v).strip().replace("ي", "ی").replace("ك", "ک").replace("\u200c", " ")
        return " ".join(x.split())

    wb = load_workbook(BytesIO(excel_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try: headers = [norm(x) for x in next(rows)]
    except StopIteration: raise ValueError("فایل Excel خالی است")
    aliases = {"serial":["شماره سریال","سریال","شماره چک"], "sayad":["شماره صیاد","صیاد","شناسه صیاد"], "status":["وضعیت فعلی"], "amount":["مبلغ","مبلغ به ارز عملیاتی"], "due":["تاریخ سررسید"]}
    def col(name):
        for a in aliases[name]:
            if a in headers: return headers.index(a)
        return None
    idx={k:col(k) for k in aliases}
    if idx["status"] is None or idx["amount"] is None: raise ValueError("ستون‌های وضعیت فعلی و مبلغ در Excel پیدا نشد")
    excel_open=[]
    excel_all=[]
    for r in rows:
        status=norm(r[idx["status"]]) if idx["status"] < len(r) else ""
        amount=r[idx["amount"]] if idx["amount"] < len(r) else 0
        try: amount_num=int(Decimal(str(amount or 0)))
        except Exception: amount_num=0
        item={"serial_number":norm(r[idx["serial"]]) if idx["serial"] is not None and idx["serial"] < len(r) else "", "sayad_number":norm(r[idx["sayad"]]) if idx["sayad"] is not None and idx["sayad"] < len(r) else "", "official_status":status.replace("نزد مامور وصول","نزد مأمور وصول"), "amount":amount_num, "due_date_excel":norm(r[idx["due"]]) if idx["due"] is not None and idx["due"] < len(r) else ""}
        excel_all.append(item)
        if status in allowed_statuses:
            excel_open.append(item)

    apply_sql=current_received_status_apply("note", "current_status")
    q=text(f"""
        SELECT note.[ReceivableNoteID] AS [ChequeID], note.[SerialNumber], note.[SayadNumber], note.[Amount], note.[DueDate], note.[State] AS [MasterState],
               current_status.[CurrentChequeState], current_status.[CurrentDocumentItemType], current_status.[CurrentBankAccountRef], current_status.[CurrentStatusDescription],
               current_status.[CurrentStatusDate], current_status.[CurrentStatusDocumentDate], current_status.[CurrentStatusTransactionID]
        FROM RPA3.[ReceivableNote] note
        {apply_sql}
        WHERE note.[NoteType]=1 AND note.[State] IN (1,2) AND note.[NormalORGuarantee]=1
          AND ISNULL(note.[Description],N'') NOT LIKE N'%ضمانت%' AND ISNULL(note.[Description],N'') NOT LIKE N'%تضمین%' AND ISNULL(note.[Description],N'') NOT LIKE N'%حسن انجام%'
        ORDER BY note.[DueDate], note.[ReceivableNoteID]
    """)
    with _engine(engine).connect() as c: sql_rows=c.execute(q).mappings().all()
    def sql_obj(r):
        return {"cheque_id":r["ChequeID"], "serial_number":norm(r["SerialNumber"]), "sayad_number":norm(r["SayadNumber"]), "amount":_as_number(r["Amount"]), "due_date":_as_datetime(r["DueDate"]), "master_state":r["MasterState"], "transaction_state":r["CurrentChequeState"], "document_item_type":r["CurrentDocumentItemType"], "bank_account_ref":r["CurrentBankAccountRef"], "current_description":r["CurrentStatusDescription"], "current_status_date":_as_datetime(r["CurrentStatusDate"]), "current_document_date":_as_datetime(r["CurrentStatusDocumentDate"]), "current_transaction_id":r["CurrentStatusTransactionID"]}
    sql_open=[sql_obj(r) for r in sql_rows]
    by_sayad={}; by_serial_amount={}
    for x in excel_open:
        if x["sayad_number"]: by_sayad.setdefault(x["sayad_number"],[]).append(x)
        if x["serial_number"]: by_serial_amount.setdefault((x["serial_number"],x["amount"]),[]).append(x)
    used=set(); matched=[]; sql_only=[]
    for sr in sql_open:
        cand=by_sayad.get(sr["sayad_number"],[]) if sr["sayad_number"] else []
        if not cand and sr["serial_number"]: cand=by_serial_amount.get((sr["serial_number"],int(sr["amount"] or 0)),[])
        chosen=next((x for x in cand if id(x) not in used),None)
        if chosen: used.add(id(chosen)); matched.append({**sr,"official_status":chosen["official_status"]})
        else: sql_only.append(sr)
    excel_only=[x for x in excel_open if id(x) not in used]

    # V124: annotate SQL-only rows with their status in the *full* official report,
    # even when that status is outside the three approved open holdings.
    all_by_sayad={}
    all_by_serial_amount={}
    for x in excel_all:
        if x["sayad_number"]: all_by_sayad.setdefault(x["sayad_number"], []).append(x)
        if x["serial_number"]: all_by_serial_amount.setdefault((x["serial_number"], x["amount"]), []).append(x)
    for sr in sql_only:
        cand=all_by_sayad.get(sr["sayad_number"], []) if sr["sayad_number"] else []
        if not cand and sr["serial_number"]:
            cand=all_by_serial_amount.get((sr["serial_number"], int(sr["amount"] or 0)), [])
        if cand:
            sr["official_status_any"] = cand[0].get("official_status")
            sr["official_due_date_excel"] = cand[0].get("due_date_excel")
        else:
            sr["official_status_any"] = None
            sr["official_due_date_excel"] = None

    def total(items): return int(sum(Decimal(str(x.get("amount") or 0)) for x in items))
    status_counts={}
    for x in sql_only:
        key=x.get("official_status_any") or "<not-found-in-official-report>"
        status_counts[key]=status_counts.get(key,0)+1
    return {"status":"success", "source":"sql_vs_official_excel_control", "dashboard_source_changed":False, "allowed_official_statuses":["نزد بانک","نزد صندوق","نزد مأمور وصول"], "sql_open":{"count":len(sql_open),"amount":total(sql_open)}, "excel_open":{"count":len(excel_open),"amount":total(excel_open)}, "matched":{"count":len(matched),"amount":total(matched)}, "sql_only":{"count":len(sql_only),"amount":total(sql_only),"official_status_counts":status_counts,"rows":sql_only}, "excel_only":{"count":len(excel_only),"amount":total(excel_only),"rows":excel_only}, "difference":{"count":len(sql_open)-len(excel_open),"amount":total(sql_open)-total(excel_open)}, "note":"V124: sql_only rows are annotated with official_status_any from the full Excel report. Excel remains validation-only and is not a Dashboard source."}

# V123 diagnostic: trace the complete SQL history of every cheque responsible for
# the SQL-vs-official-report discrepancy. Excel remains validation-only.
def analyze_received_cheque_official_discrepancies(excel_bytes: bytes, engine: Engine | None = None) -> dict[str, Any]:
    comparison = compare_received_open_sql_with_rahkaran_excel(excel_bytes, engine=engine)
    sql_only = comparison.get("sql_only", {}).get("rows", [])
    excel_only = comparison.get("excel_only", {}).get("rows", [])

    cheque_ids = {int(x["cheque_id"]) for x in sql_only if x.get("cheque_id") is not None}
    excel_only_resolution = []

    # Resolve the official-report-only rows back to ReceivableNote WITHOUT an
    # open-state filter. This is precisely what lets us see why the 9 collector
    # cheques were missed by State IN (1,2).
    if excel_only:
        clauses = []
        params: dict[str, Any] = {}
        for i, row in enumerate(excel_only):
            sayad = str(row.get("sayad_number") or "").strip()
            serial = str(row.get("serial_number") or "").strip()
            amount = int(row.get("amount") or 0)
            local = []
            if sayad:
                params[f"sayad_{i}"] = sayad
                local.append(f"LTRIM(RTRIM(CAST(note.[SayadNumber] AS nvarchar(100)))) = :sayad_{i}")
            if serial:
                params[f"serial_{i}"] = serial
                params[f"amount_{i}"] = amount
                local.append(f"(LTRIM(RTRIM(CAST(note.[SerialNumber] AS nvarchar(100)))) = :serial_{i} AND note.[Amount] = :amount_{i})")
            if local:
                clauses.append("(" + " OR ".join(local) + ")")
        if clauses:
            q = text(f"""
                SELECT note.[ReceivableNoteID] AS [ChequeID], note.[SerialNumber], note.[SayadNumber],
                       note.[Amount], note.[DueDate], note.[State] AS [MasterState],
                       note.[NormalORGuarantee], note.[Description]
                FROM RPA3.[ReceivableNote] note
                WHERE note.[NoteType] = 1 AND ({' OR '.join(clauses)})
            """)
            with _engine(engine).connect() as c:
                resolved = c.execute(q, params).mappings().all()
            for r in resolved:
                cid = int(r["ChequeID"])
                cheque_ids.add(cid)
                excel_only_resolution.append({
                    "cheque_id": cid,
                    "serial_number": str(r["SerialNumber"] or "").strip(),
                    "sayad_number": str(r["SayadNumber"] or "").strip(),
                    "amount": _as_number(r["Amount"]),
                    "due_date": _as_datetime(r["DueDate"]),
                    "master_state": r["MasterState"],
                    "normal_or_guarantee": r["NormalORGuarantee"],
                    "master_description": r["Description"],
                })

    histories: list[dict[str, Any]] = []
    if cheque_ids:
        ids = sorted(cheque_ids)
        bind = {f"id_{i}": cid for i, cid in enumerate(ids)}
        placeholders = ",".join(f":id_{i}" for i in range(len(ids)))
        q = text(f"""
            SELECT note.[ReceivableNoteID] AS [ChequeID], note.[SerialNumber], note.[SayadNumber],
                   note.[Amount], note.[DueDate], note.[State] AS [MasterState],
                   rnt.[ReceivableNoteTransactionID] AS [TransactionID],
                   rnt.[State] AS [TransactionState], rnt.[DocumentState],
                   rnt.[DocumentItemType], rnt.[BankAccountRef], rnt.[DocumentRef],
                   rnt.[DocumentNumber], rnt.[Date] AS [TransactionDate],
                   rnt.[DocumentDate], rnt.[Description] AS [TransactionDescription]
            FROM RPA3.[ReceivableNote] note
            LEFT JOIN RPA3.[ReceivableNoteTransaction] rnt
              ON rnt.[ReceivableNoteRef] = note.[ReceivableNoteID]
            WHERE note.[ReceivableNoteID] IN ({placeholders})
            ORDER BY note.[ReceivableNoteID], rnt.[Date], rnt.[DocumentDate], rnt.[ReceivableNoteTransactionID]
        """)
        with _engine(engine).connect() as c:
            rows = c.execute(q, bind).mappings().all()
        grouped: dict[int, dict[str, Any]] = {}
        sql_only_ids = {int(x["cheque_id"]) for x in sql_only if x.get("cheque_id") is not None}
        excel_only_ids = {int(x["cheque_id"]) for x in excel_only_resolution}
        for r in rows:
            cid = int(r["ChequeID"])
            item = grouped.setdefault(cid, {
                "cheque_id": cid,
                "side": "sql_only" if cid in sql_only_ids else "excel_only" if cid in excel_only_ids else "resolved",
                "serial_number": str(r["SerialNumber"] or "").strip(),
                "sayad_number": str(r["SayadNumber"] or "").strip(),
                "amount": _as_number(r["Amount"]),
                "due_date": _as_datetime(r["DueDate"]),
                "master_state": r["MasterState"],
                "transactions": [],
            })
            if r["TransactionID"] is not None:
                item["transactions"].append({
                    "transaction_id": r["TransactionID"],
                    "transaction_state": r["TransactionState"],
                    "document_state": r["DocumentState"],
                    "document_item_type": r["DocumentItemType"],
                    "bank_account_ref": r["BankAccountRef"],
                    "document_ref": r["DocumentRef"],
                    "document_number": r["DocumentNumber"],
                    "transaction_date": _as_datetime(r["TransactionDate"]),
                    "document_date": _as_datetime(r["DocumentDate"]),
                    "description": r["TransactionDescription"],
                })
        histories = list(grouped.values())

    # Compact signatures make the common business transition visible without
    # hiding the full history returned above.
    signatures: dict[str, dict[str, int]] = {"sql_only": {}, "excel_only": {}}
    for h in histories:
        side = h["side"]
        if side not in signatures:
            continue
        txs = h["transactions"]
        last = txs[-1] if txs else {}
        key = "master={}|tx={}|doc_item={}|bank={}|desc={}".format(
            h.get("master_state"), last.get("transaction_state"), last.get("document_item_type"),
            "yes" if last.get("bank_account_ref") is not None else "no",
            (str(last.get("description") or "").strip()[:80] or "<empty>")
        )
        signatures[side][key] = signatures[side].get(key, 0) + 1

    unresolved_excel_only = []
    resolved_keys = {(x["serial_number"], x["sayad_number"], int(x["amount"] or 0)) for x in excel_only_resolution}
    for x in excel_only:
        key = (str(x.get("serial_number") or "").strip(), str(x.get("sayad_number") or "").strip(), int(x.get("amount") or 0))
        if key not in resolved_keys:
            unresolved_excel_only.append(x)

    return {
        "status": "success",
        "purpose": "trace_exact_sql_vs_official_status_difference",
        "dashboard_source_changed": False,
        "control": {
            "sql_open": comparison.get("sql_open"),
            "excel_open": comparison.get("excel_open"),
            "sql_only": {"count": len(sql_only), "amount": comparison.get("sql_only", {}).get("amount")},
            "excel_only": {"count": len(excel_only), "amount": comparison.get("excel_only", {}).get("amount")},
            "net_difference": comparison.get("difference"),
        },
        "excel_only_sql_resolution": {
            "requested_count": len(excel_only),
            "resolved_count": len(excel_only_resolution),
            "unresolved_count": len(unresolved_excel_only),
            "resolved_rows": excel_only_resolution,
            "unresolved_rows": unresolved_excel_only,
        },
        "signatures": signatures,
        "history_count": len(histories),
        "histories": histories,
        "next_step": "Use the full histories/signatures to derive the SQL-only business-status rule; do not hard-code cheque IDs.",
    }

# V125 diagnostic: discover the SQL scope rule behind the 26 rows that are absent
# from Rahkaran's official ListData report. Excel is validation only.
def analyze_received_cheque_official_scope_rule(excel_bytes: bytes, engine: Engine | None = None) -> dict[str, Any]:
    comparison = compare_received_open_sql_with_rahkaran_excel(excel_bytes, engine=engine)
    sql_only = comparison.get("sql_only", {}).get("rows", [])
    ids = [int(x["cheque_id"]) for x in sql_only if x.get("cheque_id") is not None]
    if not ids:
        return {"status":"success", "version":"V125", "sql_only_count":0, "message":"هیچ رکورد SQL-only وجود ندارد."}

    bind_names = ",".join(f":id{i}" for i in range(len(ids)))
    params = {f"id{i}": v for i, v in enumerate(ids)}
    # Return every master field for the 26 rows. This avoids guessing which
    # Rahkaran master column controls the report scope.
    q_master = text(f"SELECT * FROM RPA3.[ReceivableNote] WHERE [ReceivableNoteID] IN ({bind_names}) ORDER BY [ReceivableNoteID]")
    # Also inspect the receipt/creation document and all transaction signatures.
    q_tx = text(f"""
        SELECT rnt.[ReceivableNoteRef] AS [ChequeID], rnt.*
        FROM RPA3.[ReceivableNoteTransaction] rnt
        WHERE rnt.[ReceivableNoteRef] IN ({bind_names})
        ORDER BY rnt.[ReceivableNoteRef], rnt.[Date], rnt.[DocumentDate], rnt.[ReceivableNoteTransactionID]
    """)
    with _engine(engine).connect() as c:
        masters = [dict(r) for r in c.execute(q_master, params).mappings().all()]
        txs = [dict(r) for r in c.execute(q_tx, params).mappings().all()]

    def safe(v: Any):
        if v is None or isinstance(v, (str, int, float, bool)): return v
        if isinstance(v, Decimal): return float(v)
        if hasattr(v, "isoformat"): return v.isoformat()
        return str(v)
    masters = [{k:safe(v) for k,v in r.items()} for r in masters]
    txs = [{k:safe(v) for k,v in r.items()} for r in txs]

    # Profile master columns across the 26 rows. Constant/low-cardinality fields
    # are the strongest candidates for the hidden official-report scope filter.
    profile=[]
    if masters:
        for col in masters[0].keys():
            vals=[r.get(col) for r in masters]
            counts={}
            for v in vals:
                key="<NULL>" if v is None else str(v)
                counts[key]=counts.get(key,0)+1
            if len(counts) <= 12:
                profile.append({"column":col,"distinct_count":len(counts),"values":dict(sorted(counts.items(), key=lambda kv:(-kv[1],kv[0]))[:12])})
        profile.sort(key=lambda x:(x["distinct_count"], x["column"]))

    tx_profile={}
    for r in txs:
        sig=f"state={r.get('State')}|doc_state={r.get('DocumentState')}|item={r.get('DocumentItemType')}|bank={'yes' if r.get('BankAccountRef') is not None else 'no'}"
        tx_profile[sig]=tx_profile.get(sig,0)+1

    return {
        "status":"success",
        "version":"V125",
        "purpose":"discover_rahkaran_official_report_scope_rule",
        "dashboard_source_changed":False,
        "excel_runtime_source":False,
        "control_target":{"count":comparison.get("excel_open",{}).get("count"),"amount":comparison.get("excel_open",{}).get("amount")},
        "sql_only":{"count":len(ids),"amount":comparison.get("sql_only",{}).get("amount"),"ids":ids},
        "master_low_cardinality_profile":profile,
        "transaction_signature_counts":dict(sorted(tx_profile.items(), key=lambda kv:-kv[1])),
        "master_rows":masters,
        "transaction_rows":txs,
        "next_step":"V126 will turn the discovered general SQL scope rule into the shared received-cheque predicate and validate 795 / 757785104980 without using Excel at runtime."
    }


def trace_received_cheque_mapping(
    cheque_id: int | None = None,
    serial_number: str | None = None,
    sayad_number: str | None = None,
    document_number: str | None = None,
    engine: Engine | None = None,
) -> dict[str, Any]:
    """Trace one received cheque end-to-end to detect UI/API/SQL mapping mismatches.

    Diagnostic only. It does not change dashboard business rules.
    At least one identifier must be supplied. document_number is matched against
    the approved Receipt linked to the cheque.
    """
    if not any([cheque_id, serial_number, sayad_number, document_number]):
        raise ValueError("حداقل یکی از cheque_id، serial_number، sayad_number یا document_number لازم است")

    where_parts = []
    params: dict[str, Any] = {}
    if cheque_id is not None:
        where_parts.append("note.[ReceivableNoteID] = :cheque_id")
        params["cheque_id"] = cheque_id
    if serial_number:
        where_parts.append("LTRIM(RTRIM(CONVERT(nvarchar(100), note.[SerialNumber]))) = :serial_number")
        params["serial_number"] = str(serial_number).strip()
    if sayad_number:
        where_parts.append("LTRIM(RTRIM(CONVERT(nvarchar(100), note.[SayadNumber]))) = :sayad_number")
        params["sayad_number"] = str(sayad_number).strip()
    if document_number:
        where_parts.append("EXISTS (SELECT 1 FROM RPA3.[ReceiptReceivableNote] rrn INNER JOIN RPA3.[Receipt] r ON r.[ReceiptID]=rrn.[ReceiptRef] WHERE rrn.[ReceivableNoteRef]=note.[ReceivableNoteID] AND CONVERT(nvarchar(100), r.[Number])=:document_number)")
        params["document_number"] = str(document_number).strip()

    query = text(f"""
        SELECT note.[ReceivableNoteID] AS [ChequeID], note.[SerialNumber], note.[Series],
               note.[SayadNumber], note.[Amount], note.[DueDate], note.[State] AS [MasterState],
               note.[NormalORGuarantee], note.[Description], note.[BankRef],
               bank.[Name] AS [BankName], note.[BankBranchName], note.[BankBranchCode],
               note.[CounterPartRef], counterpart.[Code] AS [CounterPartCode],
               counterpart.[Title] AS [CounterPartName]
        FROM RPA3.[ReceivableNote] note
        LEFT JOIN RPA3.[Bank] bank ON bank.[BankID]=note.[BankRef]
        LEFT JOIN FIN3.[DL] counterpart ON counterpart.[DLID]=note.[CounterPartRef]
        WHERE note.[NoteType]=1 AND ({' OR '.join(where_parts)})
        ORDER BY note.[ReceivableNoteID]
    """)

    with _engine(engine).connect() as connection:
        masters = connection.execute(query, params).mappings().all()
        ids = [int(r["ChequeID"]) for r in masters]
        if not ids:
            return {"status": "success", "diagnostic_only": True, "match_count": 0, "matches": [], "warning": "هیچ چکی با شناسه‌های ورودی پیدا نشد"}

        id_params = {f"id{i}": v for i, v in enumerate(ids)}
        placeholders = ",".join(f":id{i}" for i in range(len(ids)))
        tx_rows = connection.execute(text(f"""
            SELECT rnt.[ReceivableNoteRef] AS [ChequeID], rnt.[ReceivableNoteTransactionID] AS [TransactionID],
                   rnt.[State] AS [TransactionState], rnt.[DocumentState], rnt.[DocumentItemType],
                   rnt.[BankAccountRef], rnt.[DocumentRef], rnt.[DocumentNumber], rnt.[Date] AS [TransactionDate],
                   rnt.[DocumentDate], rnt.[Description]
            FROM RPA3.[ReceivableNoteTransaction] rnt
            WHERE rnt.[ReceivableNoteRef] IN ({placeholders})
            ORDER BY rnt.[ReceivableNoteRef], rnt.[Date], rnt.[DocumentDate], rnt.[ReceivableNoteTransactionID]
        """), id_params).mappings().all()
        receipt_rows = connection.execute(text(f"""
            SELECT rrn.[ReceivableNoteRef] AS [ChequeID], r.[ReceiptID], r.[Number] AS [ReceiptNumber],
                   r.[Date] AS [ReceiptDate], r.[ApproveState], r.[Description] AS [ReceiptDescription]
            FROM RPA3.[ReceiptReceivableNote] rrn
            INNER JOIN RPA3.[Receipt] r ON r.[ReceiptID]=rrn.[ReceiptRef]
            WHERE rrn.[ReceivableNoteRef] IN ({placeholders})
            ORDER BY rrn.[ReceivableNoteRef], r.[Date], r.[ReceiptID]
        """), id_params).mappings().all()

    tx_by_id: dict[int, list[dict[str, Any]]] = {i: [] for i in ids}
    for row in tx_rows:
        d = dict(row)
        for key in ("TransactionDate", "DocumentDate"):
            d[key] = _as_datetime(d.get(key))
        tx_by_id[int(row["ChequeID"])].append(d)
    receipts_by_id: dict[int, list[dict[str, Any]]] = {i: [] for i in ids}
    for row in receipt_rows:
        d = dict(row); d["ReceiptDate"] = _as_datetime(d.get("ReceiptDate"))
        receipts_by_id[int(row["ChequeID"])].append(d)

    matches = []
    for row in masters:
        cid = int(row["ChequeID"])
        master = dict(row)
        master["Amount"] = _as_number(master.get("Amount"))
        master["DueDate"] = _as_datetime(master.get("DueDate"))
        master["DueDateJalali"] = _as_jalali_date(row.get("DueDate"))
        txs = tx_by_id[cid]
        latest = txs[-1] if txs else None
        matches.append({
            "cheque_id": cid,
            "sql_master": master,
            "api_expected_core": {
                "cheque_id": cid,
                "serial_number": str(row.get("SerialNumber") or ""),
                "sayad_number": str(row.get("SayadNumber") or ""),
                "amount_rial": _as_number(row.get("Amount")),
                "amount_toman": _as_number(row.get("Amount")) / 10,
                "due_date": _as_datetime(row.get("DueDate")),
                "due_date_jalali": _as_jalali_date(row.get("DueDate")),
                "master_state": row.get("MasterState"),
                "counterparty": row.get("CounterPartName"),
                "bank": row.get("BankName"),
                "bank_branch": row.get("BankBranchName"),
            },
            "latest_transaction_any": latest,
            "transactions": txs,
            "linked_receipts": receipts_by_id[cid],
        })
    return {
        "status": "success", "diagnostic_only": True, "dashboard_source_changed": False,
        "input": {"cheque_id": cheque_id, "serial_number": serial_number, "sayad_number": sayad_number, "document_number": document_number},
        "match_count": len(matches), "multiple_match_warning": len(matches) > 1,
        "matches": matches,
        "check": "Compare sql_master.DueDate/DueDateJalali with api_expected_core and the frontend row. If they differ in UI only, the bug is frontend mapping/merge; if SQL master itself differs, inspect identifier collision or source selection."
    }


def _payment_order_category(description: str | None) -> str:
    text_value = (description or "").strip()
    rules = [
        ("تنخواه", ("تنخواه",)),
        ("حقوق و پرسنل", ("حقوق", "مساعده", "تسویه حساب", "تسويه حساب")),
        ("حمل و باربری", ("بارنامه", "حمل بار", "حمل و نقل", "راهداری")),
        ("واردات و بازرگانی", ("ترخیص", "ترخيص", "کوتاژ", "ثبت سفارش", "پرفرم", "گواهینامه بازرسی")),
        ("خرید و تأمین‌کننده", ("پ ف", "پیش فاکتور", "پيش فاکتور", "فاکتور", "بابت ظرف", "بابت قوطی", "بابت لیبل", "مواد اولیه")),
        ("خدمات و هزینه‌ها", ("تعمیر", "سرویس", "آنالیز", "آزمایشگاه", "اجاره", "ایونت", "مارکتینگ", "طراحی", "کارشناسی")),
    ]
    for label, tokens in rules:
        if any(token in text_value for token in tokens):
            return label
    return "سایر"


def get_company_payment_orders(limit: int = 1000, engine: Engine | None = None) -> dict[str, Any]:
    """حواله‌های پرداختی شرکت از PaymentOrder راهکاران، با تفکیک روش پرداخت و دسته مدیریتی."""
    safe_limit = max(1, min(int(limit), 5000))
    query = text(f"""
        SELECT TOP ({safe_limit})
            po.[PaymentOrderID], po.[Number] AS [PaymentOrderNumber],
            po.[Date] AS [OrderDate], po.[CreationDate], po.[ApproveDate],
            po.[CounterPartRef], cp.[Code] AS [CounterPartCode], cp.[Title] AS [CounterPartName],
            po.[State], po.[PaymentType], po.[SourceType], po.[SourceRef],
            po.[BranchRef], po.[FiscalYearRef], po.[Description], po.[CurrencyRef],
            po.[TotalOperationalCurrencyAmount],
            ISNULL(dep.[Amount], 0) AS [DepositAmount],
            ISNULL(cash.[Amount], 0) AS [CashAmount],
            ISNULL(note.[Amount], 0) AS [ChequeAmount]
        FROM RPA3.[PaymentOrder] po
        LEFT JOIN FIN3.[DL] cp ON cp.[DLID] = po.[CounterPartRef]
        OUTER APPLY (
            SELECT SUM(x.[Amount]) AS [Amount] FROM RPA3.[PaymentOrderDeposit] x
            WHERE x.[PaymentOrderRef] = po.[PaymentOrderID]
        ) dep
        OUTER APPLY (
            SELECT SUM(x.[Amount]) AS [Amount] FROM RPA3.[PaymentOrderCashMoney] x
            WHERE x.[PaymentOrderRef] = po.[PaymentOrderID]
        ) cash
        OUTER APPLY (
            SELECT SUM(x.[Amount]) AS [Amount] FROM RPA3.[PaymentOrderPayableNote] x
            WHERE x.[PaymentOrderRef] = po.[PaymentOrderID]
        ) note
        ORDER BY po.[Date] DESC, po.[PaymentOrderID] DESC
    """)
    with _engine(engine).connect() as connection:
        db_rows = connection.execute(query).mappings().all()

    rows: list[dict[str, Any]] = []
    for row in db_rows:
        deposit = _as_number(row["DepositAmount"])
        cash = _as_number(row["CashAmount"])
        cheque = _as_number(row["ChequeAmount"])
        calculated = deposit + cash + cheque
        if cheque and not (deposit or cash):
            method = "چکی"
        elif (deposit or cash) and not cheque:
            method = "نقد/بانکی"
        elif calculated:
            method = "ترکیبی"
        else:
            method = "نامشخص"
        rows.append({
            "payment_order_id": row["PaymentOrderID"],
            "payment_order_number": row["PaymentOrderNumber"],
            "order_date": _as_datetime(row["OrderDate"]),
            "order_date_jalali": _as_jalali_date(row["OrderDate"]),
            "creation_date": _as_datetime(row["CreationDate"]),
            "approve_date": _as_datetime(row["ApproveDate"]),
            "counterpart_ref": row["CounterPartRef"],
            "counterpart_code": row["CounterPartCode"],
            "counterpart_name": row["CounterPartName"],
            "state": row["State"],
            "state_label": f"وضعیت {row['State']}",
            "payment_type": row["PaymentType"],
            "source_type": row["SourceType"],
            "source_ref": row["SourceRef"],
            "branch_ref": row["BranchRef"],
            "fiscal_year_ref": row["FiscalYearRef"],
            "description": row["Description"],
            "currency_ref": row["CurrencyRef"],
            "operational_currency_amount": _as_number(row["TotalOperationalCurrencyAmount"]),
            "deposit_amount": deposit,
            "cash_amount": cash,
            "cheque_amount": cheque,
            "calculated_amount": calculated,
            "payment_method": method,
            "category": _payment_order_category(row["Description"]),
            "is_approved": row["ApproveDate"] is not None,
        })

    categories: dict[str, dict[str, Any]] = {}
    for item in rows:
        bucket = categories.setdefault(item["category"], {"count": 0, "amount": 0})
        bucket["count"] += 1
        bucket["amount"] += item["calculated_amount"]
    return {
        "status": "ok", "source": "rahkaran", "rows": rows,
        "summary": {
            "count": len(rows),
            "calculated_amount": sum(x["calculated_amount"] for x in rows),
            "approved_count": sum(1 for x in rows if x["is_approved"]),
            "waiting_count": sum(1 for x in rows if not x["is_approved"]),
            "bank_cash_count": sum(1 for x in rows if x["payment_method"] == "نقد/بانکی"),
            "cheque_count": sum(1 for x in rows if x["payment_method"] == "چکی"),
            "categories": categories,
        },
        "note": "دسته‌بندی مدیریتی بر پایه شرح حواله است؛ State خام راهکاران نیز برای کنترل نمایش داده می‌شود.",
    }
