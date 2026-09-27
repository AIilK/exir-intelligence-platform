from __future__ import annotations

import json
import math
import re
import shutil
from dataclasses import dataclass, asdict
from datetime import datetime, timezone
from pathlib import Path
from statistics import mean, median
from typing import Any

from openpyxl import load_workbook


PERSIAN_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789")
ARABIC_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")
MONTH_NAMES = {
    1: "فروردین", 2: "اردیبهشت", 3: "خرداد", 4: "تیر", 5: "مرداد", 6: "شهریور",
    7: "مهر", 8: "آبان", 9: "آذر", 10: "دی", 11: "بهمن", 12: "اسفند",
}


@dataclass
class Snapshot:
    jalali_date: str
    jalali_year: int
    jalali_month: int
    jalali_day: int
    filename: str
    liquidity_rial: float
    overdue_cheques_rial: float
    upcoming_cheques_rial: float
    loan_installments_rial: float
    obligations_rial: float
    remaining_cash_rial: float
    pressure_percent: float
    coverage_percent: float
    shortfall_rial: float
    source_sections: list[dict[str, Any]]


class MonthlyCashflowExcelService:
    """Stores daily Cash Flow workbooks and analyzes snapshots from the same Jalali month."""

    def __init__(self, root: Path | None = None):
        self.root = root or Path("data") / "cashflow_excel"
        self.root.mkdir(parents=True, exist_ok=True)

    @staticmethod
    def _number(value: Any) -> float:
        if value is None or value == "":
            return 0.0
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return float(value)
        text = str(value).translate(PERSIAN_DIGITS).translate(ARABIC_DIGITS)
        text = text.replace(",", "").replace("٬", "").strip()
        try:
            return float(text)
        except ValueError:
            return 0.0

    @staticmethod
    def _date_from_text(value: Any) -> tuple[int, int, int] | None:
        text = str(value or "").translate(PERSIAN_DIGITS).translate(ARABIC_DIGITS)
        match = re.search(r"(1[34]\d{2})\s*[/_.-]\s*(\d{1,2})\s*[/_.-]\s*(\d{1,2})", text)
        if not match:
            match = re.search(r"(1[34]\d{2})(\d{2})(\d{2})", text)
        if not match:
            return None
        year, month, day = map(int, match.groups())
        if 1 <= month <= 12 and 1 <= day <= 31:
            return year, month, day
        return None

    def _detect_date(self, workbook, filename: str) -> tuple[int, int, int]:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(min_row=1, max_row=min(8, sheet.max_row), values_only=True):
                for value in row:
                    found = self._date_from_text(value)
                    if found:
                        return found
        found = self._date_from_text(filename)
        if found:
            return found
        raise ValueError("تاریخ شمسی از داخل فایل یا نام فایل تشخیص داده نشد؛ نمونه نام صحیح: 14050603.xlsx")

    @staticmethod
    def _find_total_row(sheet) -> tuple[int, list[Any]]:
        candidates: list[tuple[int, list[Any]]] = []
        for row_no, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = list(row)
            if any(str(value or "").strip() == "جمع" for value in values):
                candidates.append((row_no, values))
        if not candidates:
            raise ValueError(f"ردیف جمع در شیت {sheet.title} پیدا نشد.")
        return candidates[-1]

    def _parse_section(self, sheet, position: int) -> dict[str, Any]:
        _, values = self._find_total_row(sheet)
        values += [None] * max(0, 8 - len(values))
        liquidity = self._number(values[3])
        overdue = self._number(values[4])
        upcoming = self._number(values[5])
        loan = self._number(values[6]) if position == 0 else 0.0
        remaining_index = 7 if position == 0 else 6
        remaining = self._number(values[remaining_index])
        return {
            "sheet": sheet.title,
            "section": ("حساب‌های رسمی", "حساب‌های غیررسمی", "زرین کالای کادوس")[position],
            "liquidity_rial": liquidity,
            "overdue_cheques_rial": overdue,
            "upcoming_cheques_rial": upcoming,
            "loan_installments_rial": loan,
            "remaining_cash_rial": remaining,
        }

    def parse(self, path: Path, original_filename: str | None = None) -> Snapshot:
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            year, month, day = self._detect_date(workbook, original_filename or path.name)
            if len(workbook.worksheets) < 3:
                raise ValueError("فایل Cash Flow باید حداقل سه شیت حساب‌های رسمی، غیررسمی و زرین داشته باشد.")
            sections = [self._parse_section(workbook.worksheets[i], i) for i in range(3)]
        finally:
            workbook.close()

        liquidity = sum(x["liquidity_rial"] for x in sections)
        overdue = sum(x["overdue_cheques_rial"] for x in sections)
        upcoming = sum(x["upcoming_cheques_rial"] for x in sections)
        loan = sum(x["loan_installments_rial"] for x in sections)
        obligations = overdue + upcoming + loan
        remaining = liquidity - obligations
        pressure = obligations / liquidity * 100 if liquidity > 0 else (100.0 if obligations else 0.0)
        coverage = min(100.0, liquidity / obligations * 100) if obligations > 0 else 100.0
        return Snapshot(
            jalali_date=f"{year:04d}/{month:02d}/{day:02d}", jalali_year=year, jalali_month=month, jalali_day=day,
            filename=original_filename or path.name, liquidity_rial=liquidity,
            overdue_cheques_rial=overdue, upcoming_cheques_rial=upcoming, loan_installments_rial=loan,
            obligations_rial=obligations, remaining_cash_rial=remaining, pressure_percent=pressure,
            coverage_percent=coverage, shortfall_rial=max(0.0, -remaining), source_sections=sections,
        )

    def save_upload(self, temporary_path: Path, original_filename: str) -> dict[str, Any]:
        snapshot = self.parse(temporary_path, original_filename)
        month_dir = self.root / str(snapshot.jalali_year) / f"{snapshot.jalali_month:02d}"
        month_dir.mkdir(parents=True, exist_ok=True)
        # One authoritative snapshot per Jalali date. Re-uploading that date replaces it.
        for old in month_dir.glob(f"{snapshot.jalali_year:04d}{snapshot.jalali_month:02d}{snapshot.jalali_day:02d}__*"):
            old.unlink()
        safe_name = re.sub(r"[^A-Za-z0-9_.-]+", "_", Path(original_filename).name)
        target = month_dir / f"{snapshot.jalali_year:04d}{snapshot.jalali_month:02d}{snapshot.jalali_day:02d}__{safe_name}"
        shutil.copy2(temporary_path, target)
        metadata = asdict(snapshot)
        metadata["stored_file"] = target.name
        metadata["uploaded_at"] = datetime.now(timezone.utc).isoformat()
        (month_dir / f"{target.stem}.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
        return {"status": "success", "snapshot": metadata, "month_analysis": self.analyze(snapshot.jalali_year, snapshot.jalali_month)}

    def _month_snapshots(self, year: int, month: int) -> list[dict[str, Any]]:
        month_dir = self.root / str(year) / f"{month:02d}"
        snapshots: dict[str, dict[str, Any]] = {}
        if not month_dir.exists():
            return []
        for path in month_dir.glob("*.json"):
            try:
                item = json.loads(path.read_text(encoding="utf-8"))
                snapshots[item["jalali_date"]] = item
            except (OSError, ValueError, KeyError, json.JSONDecodeError):
                continue
        return sorted(snapshots.values(), key=lambda x: (x["jalali_year"], x["jalali_month"], x["jalali_day"]))

    def stored_workbook(self, year: int, month: int, stored_file: str) -> Path:
        """Resolve only an authoritative workbook stored inside the requested month."""
        if Path(stored_file).name != stored_file or not stored_file.lower().endswith(".xlsx"):
            raise ValueError("نام فایل Excel نامعتبر است.")
        month_dir = (self.root / str(year) / f"{month:02d}").resolve()
        candidate = (month_dir / stored_file).resolve()
        if candidate.parent != month_dir or not candidate.is_file():
            raise FileNotFoundError("فایل Excel پیدا نشد.")
        if not any(row.get("stored_file") == stored_file for row in self._month_snapshots(year, month)):
            raise FileNotFoundError("این فایل Snapshot معتبر ماه انتخاب‌شده نیست.")
        return candidate

    def workbook_preview(self, year: int, month: int, stored_file: str) -> dict[str, Any]:
        path = self.stored_workbook(year, month, stored_file)
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            sheets = []
            for sheet in workbook.worksheets:
                rows = []
                for values in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 0, 250), max_col=min(sheet.max_column or 0, 40), values_only=True):
                    row = []
                    for value in values:
                        if isinstance(value, datetime):
                            value = value.isoformat()
                        elif value is not None and not isinstance(value, (str, int, float, bool)):
                            value = str(value)
                        row.append(value)
                    while row and row[-1] is None:
                        row.pop()
                    rows.append(row)
                while rows and not rows[-1]:
                    rows.pop()
                sheets.append({"name": sheet.title, "rows": rows, "shown_row_count": len(rows), "total_row_count": sheet.max_row or 0, "shown_column_count": max((len(row) for row in rows), default=0), "truncated": (sheet.max_row or 0) > 250 or (sheet.max_column or 0) > 40})
        finally:
            workbook.close()
        return {"status": "success", "stored_file": stored_file, "sheets": sheets}

    @staticmethod
    def _slope(values: list[float]) -> float:
        n = len(values)
        if n < 2:
            return 0.0
        xs = list(range(n)); xbar = mean(xs); ybar = mean(values)
        denominator = sum((x - xbar) ** 2 for x in xs)
        return sum((x - xbar) * (y - ybar) for x, y in zip(xs, values)) / denominator if denominator else 0.0

    def analyze(self, year: int, month: int) -> dict[str, Any]:
        rows = self._month_snapshots(year, month)
        if not rows:
            return {"status": "empty", "jalali_year": year, "jalali_month": month, "month_name": MONTH_NAMES.get(month), "message": "برای این ماه فایلی ثبت نشده است.", "snapshots": []}
        obligations = [float(x["obligations_rial"]) for x in rows]
        liquidity = [float(x["liquidity_rial"]) for x in rows]
        pressure = [float(x["pressure_percent"]) for x in rows]
        latest = rows[-1]
        total_liquidity = sum(liquidity)
        weighted_pressure = sum(obligations) / total_liquidity * 100 if total_liquidity else 0.0
        slope = self._slope(pressure)
        remaining_days = max(0, 31 - int(latest["jalali_day"]))
        projected_end_pressure = max(0.0, min(200.0, pressure[-1] + slope * remaining_days))
        estimated_coverage = 100.0 if projected_end_pressure <= 100 else max(0.0, 10000 / projected_end_pressure)
        peak = max(rows, key=lambda x: float(x["pressure_percent"]))
        shortfall_days = sum(1 for x in rows if float(x["shortfall_rial"]) > 0)
        status = "critical" if shortfall_days or projected_end_pressure >= 70 else "attention" if projected_end_pressure >= 30 else "healthy"
        trend = "افزایشی" if slope > 0.25 else "کاهشی" if slope < -0.25 else "باثبات"
        good = [f"در {len(rows) - shortfall_days} گزارش از {len(rows)} گزارش کسری نقدینگی ثبت نشده است."]
        risks = []
        actions = []
        if slope > 0.25:
            risks.append("فشار تعهدات چکی در طول ماه روند افزایشی دارد.")
            actions.append({"action": "موجودی حساب‌های پرداخت را پیش از سررسیدهای پرتراکم تقویت کنید.", "why": "شیب فشار چک‌ها نسبت به گزارش‌های ابتدای ماه افزایشی است."})
        if float(peak["pressure_percent"]) >= 40:
            risks.append(f"در {peak['jalali_date']} فشار پرداخت به {peak['pressure_percent']:.1f}٪ موجودی رسیده است.")
            actions.append({"action": "سررسیدهای روز اوج را با ورودی‌های قطعی تطبیق دهید.", "why": "تمرکز پرداخت در یک روز، ریسک عملیاتی ایجاد می‌کند حتی اگر مانده کل مثبت باشد."})
        if shortfall_days:
            risks.append(f"در {shortfall_days} گزارش کسری نقدینگی مشاهده شده است.")
        if not actions:
            actions.append({"action": "پوشش نقدی فعلی حفظ و فایل روزانه بدون وقفه بارگذاری شود.", "why": "در داده‌های ماه تا امروز کسری مشاهده نشده و فشار پایان ماه کنترل‌شده برآورد شده است."})
        return {
            "status": "success", "jalali_year": year, "jalali_month": month, "month_name": MONTH_NAMES.get(month),
            "period_label": f"{MONTH_NAMES.get(month, month)} {year}", "snapshot_count": len(rows), "latest_snapshot": latest,
            "summary": {
                "average_obligations_rial": mean(obligations), "median_obligations_rial": median(obligations),
                "latest_obligations_rial": obligations[-1], "latest_liquidity_rial": liquidity[-1],
                "weighted_average_pressure_percent": weighted_pressure, "average_coverage_percent": mean(float(x["coverage_percent"]) for x in rows),
                "peak_pressure_percent": float(peak["pressure_percent"]), "peak_pressure_date": peak["jalali_date"],
                "shortfall_days": shortfall_days,
            },
            "forecast": {
                "method": "روند خطی فشار تعهدات روی Snapshotهای همان ماه؛ سناریوی توضیح‌پذیر و نه نتیجه قطعی",
                "trend": trend, "pressure_slope_per_snapshot": slope, "projected_end_month_pressure_percent": projected_end_pressure,
                "estimated_payment_capacity_percent": min(100.0, estimated_coverage),
                "confidence": "medium" if len(rows) >= 5 else "low", "status": status,
            },
            "agent_analysis": {
                "headline": f"وضعیت پوشش چک‌های {MONTH_NAMES.get(month)}: {'مناسب' if status == 'healthy' else 'نیازمند توجه' if status == 'attention' else 'بحرانی'}",
                "human_summary": f"تا روز {latest['jalali_day']} ماه، {len(rows)} گزارش بررسی شد. متوسط فشار وزنی چک‌ها {weighted_pressure:.1f}٪ و ظرفیت پرداخت پیش‌بینی‌شده تا پایان ماه {min(100.0, estimated_coverage):.1f}٪ است.",
                "good_signals": good, "risks": risks, "recommended_actions": actions,
            },
            "snapshots": rows,
        }

    def months(self) -> list[dict[str, Any]]:
        result = []
        for year_dir in self.root.iterdir():
            if not year_dir.is_dir() or not year_dir.name.isdigit():
                continue
            for month_dir in year_dir.iterdir():
                if month_dir.is_dir() and month_dir.name.isdigit():
                    rows = self._month_snapshots(int(year_dir.name), int(month_dir.name))
                    if rows:
                        result.append({"jalali_year": int(year_dir.name), "jalali_month": int(month_dir.name), "month_name": MONTH_NAMES.get(int(month_dir.name)), "snapshot_count": len(rows), "latest_date": rows[-1]["jalali_date"]})
        return sorted(result, key=lambda x: (x["jalali_year"], x["jalali_month"]), reverse=True)

    def latest_snapshot(self) -> dict[str, Any] | None:
        available = self.months()
        if not available:
            return None
        latest_month = available[0]
        rows = self._month_snapshots(
            latest_month["jalali_year"],
            latest_month["jalali_month"],
        )
        return rows[-1] if rows else None

    def historical_comparison(self, months_limit: int = 6) -> dict[str, Any]:
        """Compare the actual daily Excel snapshots of the most recent months.

        This is deliberately based on the uploaded source workbooks rather than
        synthetic backfilling.  A month with no uploaded Cash Flow is omitted
        instead of being represented as zero.
        """
        safe_limit = max(2, min(int(months_limit), 24))
        available = self.months()[:safe_limit]
        periods: list[dict[str, Any]] = []
        for month in reversed(available):
            analysis = self.analyze(month["jalali_year"], month["jalali_month"])
            if analysis.get("status") != "success":
                continue
            summary = analysis.get("summary") or {}
            latest = analysis.get("latest_snapshot") or {}
            periods.append({
                "jalali_year": month["jalali_year"],
                "jalali_month": month["jalali_month"],
                "period_label": analysis.get("period_label"),
                "snapshot_count": analysis.get("snapshot_count", 0),
                "latest_snapshot_date": latest.get("jalali_date"),
                "latest_liquidity_rial": summary.get("latest_liquidity_rial", 0),
                "average_obligations_rial": summary.get("average_obligations_rial", 0),
                "weighted_average_pressure_percent": summary.get("weighted_average_pressure_percent", 0),
                "average_coverage_percent": summary.get("average_coverage_percent", 0),
                "peak_pressure_percent": summary.get("peak_pressure_percent", 0),
                "peak_pressure_date": summary.get("peak_pressure_date"),
                "shortfall_days": summary.get("shortfall_days", 0),
                "trend": (analysis.get("forecast") or {}).get("trend"),
            })

        current = periods[-1] if periods else None
        previous = periods[-2] if len(periods) >= 2 else None
        previous_pressure = float(previous.get("weighted_average_pressure_percent") or 0) if previous else None
        current_pressure = float(current.get("weighted_average_pressure_percent") or 0) if current else None
        return {
            "status": "success" if periods else "empty",
            "source": "daily_cashflow_excel_snapshots",
            "period_count": len(periods),
            "periods": periods,
            "current_vs_previous": None if not current or not previous else {
                "current_period": current["period_label"],
                "previous_period": previous["period_label"],
                "pressure_change_percent_point": round(current_pressure - previous_pressure, 2),
                "liquidity_change_rial": round(
                    float(current.get("latest_liquidity_rial") or 0)
                    - float(previous.get("latest_liquidity_rial") or 0),
                    2,
                ),
                "shortfall_day_change": int(current.get("shortfall_days") or 0)
                - int(previous.get("shortfall_days") or 0),
            },
            "forecast_use": {
                "used_as": "historical risk context",
                "rule": "روند ماه‌های گذشته برای هشدار و تفسیر پیش‌بینی استفاده می‌شود؛ مبلغ قطعی Cash Flow فقط از موجودی Excel و اسناد SQL خواندنی محاسبه می‌شود.",
                "confidence": "medium" if len(periods) >= 3 else "low",
            },
        }
