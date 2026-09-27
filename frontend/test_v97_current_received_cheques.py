from collections import Counter
from copy import deepcopy
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
import pytest

from app.services.karamad_manual_import_service import KaramadManualImportService, _text

SAMPLE = Path(__file__).resolve().parent / 'test_fixtures/v97_current_received.xlsx'

@pytest.fixture
def service(tmp_path):
    return KaramadManualImportService(tmp_path/'in', tmp_path/'archive', tmp_path/'state.json')


def test_real_snapshot_has_exact_statuses_dates_and_no_footer_record(service):
    result = service.import_workbook(SAMPLE, 'tmpE8A6.tmp.xlsx')
    rows = service.cheque_rows('received_cheques', include_current_portfolio=True)
    assert len(rows) == len({r['cheque_id'] for r in rows}) == 831
    assert sum(r['amount'] for r in rows) == 467665127966
    assert Counter(r['state_label'] for r in rows) == {
        'واگذار شده': 698, 'برگشتی نزد مشتری': 94, 'نزد صندوق': 23, 'وضعیت نامشخص': 16,
    }
    assert all(r['collection_date'] is None for r in rows)
    first = rows[0]
    assert first['due_date_jalali'] == '1404/06/05'
    assert first['registration_date_jalali'] == '1403/09/04'
    assert first['assignment_date_jalali'] == '1403/09/21'
    assert first['serial_number'] == '488986'
    assert first['sayad_number'] == '2128030156242386'
    assert first['state'] == 'assigned_for_collection'
    assert 'transfer_number' not in first
    assert result['import']['usage_scope'] == 'cheques_only'
    assert service.import_workbook(SAMPLE, SAMPLE.name)['duplicate_file']


def test_snapshot_does_not_change_existing_history_actuals_or_forecast_inputs(service):
    state = service._empty_state()
    legacy = {'direction':'inflow', 'source_kind':'received_cheques', 'source_kinds':['received_cheques'],
              'transfer_id':'71050', 'transfer_number':'L-1', 'classification':'operational',
              'amount_rial':1000, 'registration_date_jalali':'1405/06/10', 'registration_month':'1405/06',
              'level4_name':'Existing customer', 'branch':'Existing branch', 'detected_schema':'executed_movement'}
    state['records']['inflow:71050'] = legacy
    service._save(state)
    records_before = deepcopy(service._state()['records'])
    customers_before = service.customer_summary()
    detail_before = service.customer_detail('Existing customer')
    movements_before = service.actual_movements(date(2026,1,1), date(2027,1,1))
    forecast_before = service.cheque_rows('received_cheques')
    service.import_workbook(SAMPLE, SAMPLE.name, 'received_cheques_current')
    assert service._state()['records'] == records_before
    assert service.customer_summary() == customers_before
    assert service.customer_detail('Existing customer') == detail_before
    assert service.actual_movements(date(2026,1,1), date(2027,1,1)) == movements_before
    assert service.cheque_rows('received_cheques') == forecast_before
    assert service.report()['pagination']['total_count'] == 1
    assert service.summary()['operational_inflow_rial'] == 1000
    current = service.cheque_rows('received_cheques', include_current_portfolio=True)
    assert len(current) == 831
    assert not any(r.get('serial_number') == 'L-1' for r in current)


def test_new_snapshot_replaces_and_known_old_file_cannot_roll_it_back(service, tmp_path, monkeypatch):
    service.import_workbook(SAMPLE, SAMPLE.name)
    raw = deepcopy(service._state()['current_received_portfolio']['records'][:2])
    raw[0]['state_label'] = 'نزد صندوق'
    updated = tmp_path/'updated.xlsx'
    updated.write_bytes(b'fixture-v2')
    monkeypatch.setattr(service, 'parse', lambda path, kind: (deepcopy(raw), []))
    service.import_workbook(updated, 'چک دریافتی جاری.xlsx')
    assert len(service.cheque_rows('received_cheques', include_current_portfolio=True)) == 2
    assert service.cheque_rows('received_cheques', include_current_portfolio=True)[0]['state_label'] == 'نزد صندوق'
    assert service.import_workbook(SAMPLE, SAMPLE.name)['duplicate_file']
    assert len(service.cheque_rows('received_cheques', include_current_portfolio=True)) == 2


class SheetRows:
    def __init__(self, rows): self.rows = rows
    def iter_rows(self, min_row, values_only): return iter(self.rows)


def test_duplicate_id_invalid_total_and_invalid_date_rejected(service):
    book = load_workbook(SAMPLE, read_only=True, data_only=True)
    try:
        data = list(book.active.values)
    finally:
        book.close()
    headers = [_text(x) for x in data[0]]
    with pytest.raises(ValueError, match="تکراری"):
        service._parse_current_cheques(SheetRows([data[1], data[1]]), 1, headers)
    with pytest.raises(ValueError, match="تطبیق ندارد"):
        service._parse_current_cheques(SheetRows([data[1], data[-1]]), 1, headers)
    invalid = list(data[1]); invalid[9] = "1405/13/01"
    with pytest.raises(ValueError):
        service._parse_current_cheques(SheetRows([invalid]), 1, headers)


def test_wrong_source_rejected_and_previous_snapshot_survives(service, tmp_path):
    service.import_workbook(SAMPLE, SAMPLE.name)
    before = service.state_path.read_bytes()
    with pytest.raises(ValueError, match="فهرست جاری"):
        service.import_workbook(SAMPLE, SAMPLE.name, 'paid_transfers')
    invalid = tmp_path/'broken.xlsx'; invalid.write_bytes(b'not-a-workbook')
    with pytest.raises(Exception):
        service.import_workbook(invalid, invalid.name)
    assert service.state_path.read_bytes() == before


def test_treasury_display_explicitly_opts_into_current_snapshot(service, monkeypatch):
    from app.services import karamad_manual_import_service, treasury_service
    service.import_workbook(SAMPLE, SAMPLE.name)
    monkeypatch.setattr(karamad_manual_import_service, 'KaramadManualImportService', lambda: service)
    assert len(treasury_service._karamad_cheques('received_cheques')) == 831
    assert service.cheque_rows('received_cheques') == []
