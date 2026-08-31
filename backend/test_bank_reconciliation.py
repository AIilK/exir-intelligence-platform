from collections import Counter
from datetime import date
from decimal import Decimal
from io import BytesIO
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from app.services.bank_reconciliation_service import (
    _detect_statement_bank_name,
    _derive_statement_account_core,
    _is_fee_text,
    _match_bank_rows,
    parse_bank_statement,
)
from app.services.reconciliation_store import (
    get_reconciliation_rows,
    get_reconciliation_summary,
    save_reconciliation_report,
)
from app.services.reconciliation_export_service import (
    build_reconciliation_workbook,
)


def bank_row(
    row_number: int,
    amount: str,
    *,
    description: str = "واریز مشتری آبرام",
    reference: str = "ABC123",
):
    return {
        "row_number": row_number,
        "date": date(2026, 8, 17),
        "amount": Decimal(amount),
        "direction": "deposit",
        "description": description,
        "reference": reference,
        "balance": None,
    }


def transaction(
    transaction_id: int,
    amount: str,
    *,
    description: str = "واریز مشتری آبرام پیگیری ABC123",
):
    return {
        "transaction_id": transaction_id,
        "booking_date": date(2026, 8, 17),
        "settlement_date": None,
        "description": description,
        "amount": Decimal(amount),
        "effect": 1,
        "reference_component": None,
        "reference_entity": None,
        "reference_ref": None,
    }


class BankReconciliationMatchingTests(unittest.TestCase):
    def test_bank_shahr_filename_wins_over_counterparty_bank_in_rows(self):
        self.assertEqual(
            _detect_statement_bank_name(
                "بانکشهر 6316 اکسیر.xls",
                "صورتحساب",
                [
                    ["نام سپرده", "اکسیر کادوس"],
                    ["تعیین وضعیت چک عهده بانک تجارت"],
                ],
            ),
            "شهر",
        )

    def test_tejarat_statement_account_maps_to_rahkaran_number(self):
        self.assertEqual(
            _derive_statement_account_core("0103560567794", "تجارت"),
            "356056779",
        )
        self.assertEqual(
            _derive_statement_account_core("0103560567877", "تجارت"),
            "356056787",
        )
        self.assertEqual(
            _derive_statement_account_core("0103560567877", "پاسارگاد"),
            "",
        )

    def test_unique_exact_match_is_posted(self):
        rows, used = _match_bank_rows(
            [bank_row(2, "100000")],
            [transaction(10, "100000")],
            tolerance_days=2,
        )
        self.assertEqual(rows[0]["business_status"], "posted")
        self.assertEqual(used, {10})

    def test_internal_transfer_deposit_has_separate_business_status(self):
        row = bank_row(
            2,
            "1800000000",
            description="انتقال از حساب 6316 به تجارت 787 اکسیر",
        )
        transfer = transaction(
            10,
            "1800000000",
            description="انتقال از 6316 به 787 اکسیر",
        )
        transfer["transaction_type"] = "bank_transfer_in"

        rows, used = _match_bank_rows([row], [transfer], tolerance_days=2)

        self.assertEqual(rows[0]["business_status"], "internal_transfer")
        self.assertEqual(rows[0]["business_status_fa"], "انتقال داخلی شرکت")
        self.assertEqual(rows[0]["match_method"], "internal_transfer_policy")
        self.assertIn(
            "نه سندخورده و نه بدون سند",
            " | ".join(rows[0]["reasons"]),
        )
        self.assertEqual(used, {10})

    def test_internal_transfer_withdrawal_has_separate_business_status(self):
        row = bank_row(
            2,
            "1650000000",
            description="انتقال به حساب دیگر شرکت",
        )
        row["direction"] = "withdrawal"
        transfer = transaction(
            10,
            "1650000000",
            description="انتقال بین حساب‌ها",
        )
        transfer.update(
            {
                "transaction_type": "bank_transfer_out",
                "effect": 2,
            }
        )

        rows, used = _match_bank_rows([row], [transfer], tolerance_days=2)

        self.assertEqual(rows[0]["business_status"], "internal_transfer")
        self.assertEqual(rows[0]["match_method"], "internal_transfer_policy")
        self.assertEqual(used, {10})

    def test_issued_cheque_clearance_matches_bank_withdrawal(self):
        row = bank_row(
            2,
            "560000000",
            description="برداشت وجه چک به شماره سریال 4371597994",
            reference="1132918004472",
        )
        row["direction"] = "withdrawal"
        issued = transaction(
            10,
            "560000000",
            description="پاس شدن چک صادرشده",
        )
        issued.update(
            {
                "transaction_type": "issued_cheque_status",
                "effect": 2,
                "item_number": "597994",
            }
        )

        rows, used = _match_bank_rows([row], [issued], tolerance_days=2)

        self.assertEqual(rows[0]["business_status"], "posted")
        self.assertEqual(used, {10})

    def test_issued_cheque_exact_serial_allows_delayed_rahkaran_status(self):
        row = bank_row(
            2,
            "5000000000",
            description="پرداخت انتقالی چک شماره",
            reference="0000908534168",
        )
        row.update(
            {
                "direction": "withdrawal",
                "bank_document_number": "402316",
                "date": date(2026, 7, 15),
            }
        )
        issued = transaction(
            10,
            "5000000000",
            description="چک صادرشده",
        )
        issued.update(
            {
                "transaction_type": "issued_cheque_status",
                "effect": 2,
                "item_number": "402316",
                "booking_date": date(2026, 8, 21),
            }
        )

        rows, used = _match_bank_rows([row], [issued], tolerance_days=2)

        self.assertEqual(rows[0]["business_status"], "posted")
        self.assertEqual(rows[0]["date_difference_days"], 37)
        self.assertEqual(used, {10})

    def test_tejarat_corrective_entry_reverses_original_bank_row(self):
        withdrawal = bank_row(
            2,
            "12000",
            description="شماره سند 350881 بابت ثبت چک",
            reference="1127118705781",
        )
        withdrawal["direction"] = "withdrawal"
        reversal = bank_row(
            3,
            "12000",
            description="برداشت از حساب - سند اصلاحی",
            reference="1127118705781",
        )
        reversal["direction"] = "deposit"

        rows, used = _match_bank_rows(
            [withdrawal, reversal],
            [],
            tolerance_days=2,
        )

        self.assertTrue(
            all(row["business_status"] == "reversed" for row in rows)
        )
        self.assertEqual(used, set())

    def test_tejarat_cheque_registration_cost_is_a_fee(self):
        self.assertTrue(_is_fee_text("شماره سند بابت ثبت چک ش ص"))
        self.assertTrue(_is_fee_text("درخواست ابطال برگ چک"))

    def test_amount_difference_requires_review(self):
        rows, used = _match_bank_rows(
            [bank_row(2, "100000")],
            [transaction(10, "200000")],
            tolerance_days=2,
        )
        self.assertEqual(rows[0]["business_status"], "needs_review")
        self.assertEqual(rows[0]["amount_difference"], 100000.0)
        self.assertEqual(used, set())

    def test_sepah_xls_header_and_footer_are_handled(self):
        rows = [
            ["صورت‌حساب سپرده", "310123456789"],
            ["IR930123456789012345678901"],
            [],
            [],
            [],
            [],
            [],
            [
                "ردیف",
                "تاریخ",
                "زمان",
                "عملیات",
                "شرح",
                "واریز (ریال)",
                "برداشت (ریال)",
                "مانده (ریال) ",
                "شعبه",
                "شماره سند",
                "شماره برگه/شماره چک",
                "شناسه واریز",
                "یادداشت",
                "شماره پیگیری",
            ],
            [
                1,
                "1405/05/25",
                "13:47:06",
                "طرف بدهکار",
                "کارمزد پایا",
                0,
                31728,
                1000000,
                "شعبه",
                260001,
                "-",
                "",
                "",
                505001,
            ],
            ["", "", "", "", "مجموع", 100000, 31728],
        ]
        with patch(
            "app.services.bank_reconciliation_service._read_xls",
            return_value=(rows, "صورت‌حساب"),
        ):
            parsed, info = parse_bank_statement(
                content=b"xls",
                filename="بانک سپه فراز.xls",
            )

        self.assertEqual(info["header_row"], 8)
        self.assertEqual(info["sheet_name"], "صورت‌حساب")
        self.assertEqual(len(parsed), 1)
        self.assertEqual(parsed[0]["direction"], "withdrawal")
        self.assertEqual(parsed[0]["bank_document_number"], "260001")
        self.assertEqual(parsed[0]["reference"], "505001")

    def test_bank_shahr_deposit_number_is_detected(self):
        rows = [
            [
                "نام سپرده",
                "قرض الحسنه جاری حقوقی اکسیر کادوس اشتهارد",
                "",
                "شماره سپرده:",
                "1001000836316",
            ],
            ["نوع ارز: ریال ایران", "", "", "شماره شبا:", "IR470610000001001000836316"],
            ["گزارش صورتحساب از تاریخ 1405/01/01 تا تاریخ 1405/05/25"],
            ["ردیف", "شرح", "تاریخ", "زمان", "واریز", "برداشت", "موجودی"],
            [1, "واریز چک", "1405/05/25", "05:09:19", 267000000, 0, 22227813059],
        ]
        with patch(
            "app.services.bank_reconciliation_service._read_xls",
            return_value=(rows, "صورتحساب"),
        ):
            parsed, info = parse_bank_statement(
                content=b"xls",
                filename="بانکشهر 6316 اکسیر.xls",
            )

        self.assertEqual(info["header_row"], 4)
        self.assertEqual(info["statement_account_number"], "1001000836316")
        self.assertEqual(
            info["statement_iban"],
            "IR470610000001001000836316",
        )
        self.assertEqual(len(parsed), 1)

    def test_tejarat_vertical_account_number_is_detected(self):
        rows = [
            [],
            [None, None, None, None, "نشانی بانک تجارت"],
            [],
            [],
            [],
            [],
            [],
            [],
            [],
            [None, None, None, None, None, "از تاریخ 1405/01/01 تا تاریخ 1405/05/27"],
            [],
            [],
            [None, None, None, None, None, None, "خلاصه وضعیت دوره"],
            [
                None,
                "مانده انتهایی دوره",
                "مانده ابتدایی دوره",
                "بستانکار",
                "بدهکار",
                "تعداد تراکنش",
                "شماره حساب",
            ],
            [None, "0", "0", "0", "0", "1", "0103560567794"],
            [],
            [],
            [
                None,
                "RRN",
                "کد رهگیری",
                "شناسه پرداخت دوم",
                "شناسه پرداخت اول",
                "شماره سند",
                "شرح سند",
                "شرح تراکنش",
                "شرح عملیات",
                "کد شعبه",
                "واریز",
                "برداشت",
                "موجودی حساب",
                "زمان",
                "تاریخ",
                "شماره حساب",
            ],
            [
                None,
                "1148818607528",
                "",
                "",
                "",
                "804557",
                "منظومه",
                "برداشت وجه چک به حساب IR130180000000103560567794",
                "",
                "1465",
                "0",
                "1,671,600,000",
                "14,562,375",
                "09:19:24",
                "1405/05/24",
                "0103560567794",
            ],
        ]
        with patch(
            "app.services.bank_reconciliation_service._read_xlsx",
            return_value=(rows, "tejaratbank-account"),
        ):
            parsed, info = parse_bank_statement(
                content=b"xlsx",
                filename="تجارت فراز779.xlsx",
            )

        self.assertEqual(info["header_row"], 18)
        self.assertEqual(info["statement_account_number"], "0103560567794")
        self.assertEqual(
            info["statement_iban"],
            "IR130180000000103560567794",
        )
        self.assertEqual(info["statement_bank_name"], "تجارت")
        self.assertEqual(info["detected_columns"]["reference"], "RRN")
        self.assertEqual(len(parsed), 1)
        self.assertEqual(parsed[0]["reference"], "1148818607528")
        self.assertIn("برداشت وجه چک", parsed[0]["description"])

    def test_no_candidate_is_unposted(self):
        rows, used = _match_bank_rows(
            [bank_row(2, "100000", description="ناشناس", reference="")],
            [transaction(10, "200000", description="فروشنده دیگر")],
            tolerance_days=2,
        )
        self.assertEqual(rows[0]["business_status"], "unposted")
        self.assertEqual(used, set())

    def test_duplicate_bank_rows_are_not_auto_posted(self):
        rows, used = _match_bank_rows(
            [bank_row(2, "100000"), bank_row(3, "100000")],
            [transaction(10, "100000")],
            tolerance_days=2,
        )
        self.assertTrue(
            all(row["business_status"] == "needs_review" for row in rows)
        )
        self.assertEqual(used, set())

    def test_same_day_bank_fees_match_one_aggregate_document(self):
        fee_rows = [
            bank_row(2, "100000", description="برداشت کارمزد پایا", reference=""),
            bank_row(3, "200000", description="برداشت کارمزد ساتنا", reference=""),
        ]
        for row in fee_rows:
            row["direction"] = "withdrawal"
        fee_transaction = transaction(
            10,
            "300000",
            description="ثبت تجمیعی کارمزد بانکی",
        )
        fee_transaction.update(
            {
                "transaction_type": "payment",
                "document_id": 500,
                "document_number": "900",
                "item_id": 10,
                "item_number": "1",
                "effect": 2,
            }
        )

        rows, used = _match_bank_rows(
            fee_rows,
            [fee_transaction],
            tolerance_days=2,
        )

        self.assertTrue(
            all(row["business_status"] == "posted" for row in rows)
        )
        self.assertTrue(
            all(row["match_method"] == "daily_fee_group" for row in rows)
        )
        self.assertTrue(
            all(row["match_group_bank_total"] == 300000.0 for row in rows)
        )
        self.assertEqual(used, {10})

    def test_independent_fee_is_split_before_matching_daily_remainder(self):
        amounts = ("26602", "28741", "16450", "3300", "40000", "18000")
        fee_rows = []
        for offset, amount in enumerate(amounts, start=2):
            description = (
                "برداشت کارمزد دریافت کارمزد واگذاری چک عادی"
                if amount == "18000"
                else "برداشت کارمزد پایا"
            )
            row = bank_row(offset, amount, description=description, reference="")
            row["direction"] = "withdrawal"
            fee_rows.append(row)

        aggregate = transaction(10, "115093", description="بابت کارمزد")
        aggregate.update(
            {
                "transaction_type": "payment",
                "document_id": 500,
                "document_number": "32222",
                "item_id": 10,
                "item_number": "1",
                "effect": 2,
            }
        )
        independent = transaction(
            11,
            "18000",
            description="کارمزد واگذاری چک عادی",
        )
        independent.update(
            {
                "transaction_type": "payment",
                "document_id": 501,
                "document_number": "32188",
                "item_id": 11,
                "item_number": "1",
                "effect": 2,
            }
        )

        rows, used = _match_bank_rows(
            fee_rows,
            [aggregate, independent],
            tolerance_days=2,
        )

        self.assertTrue(all(row["business_status"] == "posted" for row in rows))
        self.assertEqual(used, {10, 11})
        aggregate_rows = [
            row
            for row in rows
            if row["matched_transaction"]["document_number"] == "32222"
        ]
        self.assertEqual(len(aggregate_rows), 5)
        self.assertTrue(
            all(row["match_group_bank_total"] == 115093.0 for row in aggregate_rows)
        )

    def test_consumed_candidate_does_not_remain_possible_match(self):
        first = bank_row(2, "100000", description="واریز اول", reference="")
        second = bank_row(3, "100000", description="واریز دوم متفاوت", reference="")
        only_transaction = transaction(10, "100000")

        rows, used = _match_bank_rows(
            [first, second],
            [only_transaction],
            tolerance_days=2,
        )

        self.assertEqual(
            Counter(row["business_status"] for row in rows),
            Counter({"posted": 1, "unposted": 1}),
        )
        self.assertEqual(used, {10})

    def test_bank_shahr_fee_descriptions_without_fee_word_are_detected(self):
        self.assertTrue(
            _is_fee_text(
                "نگهداري اماني چک هاي با سررسيد بيش از 30 روز"
            )
        )
        self.assertTrue(
            _is_fee_text("انتقال وجه بين بانکي ساتنا(حضوري )")
        )
        self.assertTrue(_is_fee_text("برداشت هزینه حسابرسی"))
        self.assertTrue(_is_fee_text("برداشت دستور پرداخت سمت متعهد"))
        self.assertFalse(_is_fee_text("ثبت دستور پرداخت ساتنا"))

    def test_cheque_serial_in_bank_description_resolves_same_amount_ambiguity(self):
        row = bank_row(
            8,
            "500000000",
            description="تعیین وضعیت چک /9917/9463881405/05/07",
            reference="",
        )
        selected = transaction(10, "500000000", description="وصولی")
        selected["transaction_type"] = "received_cheque_status"
        selected["item_number"] = "946388 9917 1234567890123456"
        other = transaction(11, "500000000", description="وصولی")
        other["transaction_type"] = "received_cheque_status"
        other["item_number"] = "539434 30007 9999999999999999"

        rows, used = _match_bank_rows(
            [row],
            [selected, other],
            tolerance_days=2,
        )

        self.assertEqual(rows[0]["business_status"], "posted")
        self.assertEqual(rows[0]["matched_transaction"]["transaction_id"], 10)
        self.assertEqual(used, {10})

    def test_exact_cheque_serial_allows_extended_posting_date(self):
        row = bank_row(
            8,
            "600000000",
            description="تعیین وضعیت چک /1936/37864021",
            reference="",
        )
        selected = transaction(10, "600000000", description="وصولی")
        selected.update(
            {
                "transaction_type": "received_cheque_status",
                "item_number": "378640/21 1936 5790040149812494",
                "booking_date": date(2026, 8, 13),
            }
        )
        other = transaction(11, "600000000", description="وصولی")
        other.update(
            {
                "transaction_type": "received_cheque_status",
                "item_number": "999999 1936 1111111111111111",
                "booking_date": date(2026, 8, 13),
            }
        )

        rows, used = _match_bank_rows(
            [row],
            [selected, other],
            tolerance_days=2,
        )

        self.assertEqual(rows[0]["business_status"], "posted")
        self.assertEqual(rows[0]["matched_transaction"]["transaction_id"], 10)
        self.assertEqual(rows[0]["date_difference_days"], 4)
        self.assertEqual(used, {10})

    def test_regular_transfer_to_received_cheque_requires_review(self):
        row = bank_row(
            8,
            "200000000",
            description="انتقال وجه از سپرده مشتری به حساب شرکت",
            reference="",
        )
        cheque = transaction(10, "200000000", description="وصولی")
        cheque["transaction_type"] = "received_cheque_status"
        cheque["item_number"] = "200000 1234567890123456"

        rows, used = _match_bank_rows([row], [cheque], tolerance_days=2)

        self.assertEqual(rows[0]["business_status"], "needs_review")
        self.assertIn("تأیید خزانه", " ".join(rows[0]["reasons"]))
        self.assertEqual(used, set())

    def test_review_candidate_is_kept_when_exact_candidate_was_consumed(self):
        exact_row = bank_row(
            8,
            "200000000",
            description="واریز مشتری",
            reference="",
        )
        probable_row = bank_row(
            9,
            "200000000",
            description="انتقال وجه از سپرده مشتری",
            reference="",
        )
        receipt = transaction(10, "200000000", description="واریز مشتری")
        cheque = transaction(11, "200000000", description="وصولی")
        cheque.update(
            {
                "transaction_type": "received_cheque_status",
                "item_number": "123456 1234 1111111111111111",
            }
        )

        rows, used = _match_bank_rows(
            [exact_row, probable_row],
            [receipt, cheque],
            tolerance_days=2,
        )

        self.assertEqual(
            Counter(row["business_status"] for row in rows),
            Counter({"posted": 1, "needs_review": 1}),
        )
        self.assertEqual(used, {10})

    def test_same_day_equal_amount_batch_is_posted_with_unique_consumption(self):
        first_row = bank_row(20, "200000", description="پرداخت اول", reference="")
        second_row = bank_row(21, "200000", description="پرداخت دوم", reference="")
        first_row["direction"] = "withdrawal"
        second_row["direction"] = "withdrawal"
        first_transaction = transaction(30, "200000", description="سند اول")
        second_transaction = transaction(31, "200000", description="سند دوم")
        for item in (first_transaction, second_transaction):
            item["effect"] = 2

        rows, used = _match_bank_rows(
            [first_row, second_row],
            [first_transaction, second_transaction],
            tolerance_days=2,
        )

        self.assertTrue(all(row["business_status"] == "posted" for row in rows))
        self.assertTrue(
            all(row["match_method"] == "same_day_amount_batch" for row in rows)
        )
        self.assertEqual(used, {30, 31})
        self.assertEqual(
            len({row["matched_transaction"]["transaction_id"] for row in rows}),
            2,
        )

    def test_tolerance_batch_is_posted_only_when_component_is_balanced(self):
        first_row = bank_row(20, "200000", description="پرداخت اول", reference="")
        second_row = bank_row(21, "200000", description="پرداخت دوم", reference="")
        first_row["direction"] = "withdrawal"
        second_row["direction"] = "withdrawal"
        first_transaction = transaction(30, "200000", description="سند اول")
        second_transaction = transaction(31, "200000", description="سند دوم")
        for item in (first_transaction, second_transaction):
            item["effect"] = 2
            item["booking_date"] = date(2026, 8, 16)

        rows, used = _match_bank_rows(
            [first_row, second_row],
            [first_transaction, second_transaction],
            tolerance_days=2,
        )

        self.assertTrue(all(row["business_status"] == "posted" for row in rows))
        self.assertTrue(
            all(row["match_method"] == "tolerance_amount_batch" for row in rows)
        )
        self.assertEqual(used, {30, 31})

    def test_ambiguous_reversal_batch_is_reconciled(self):
        withdrawals = []
        reversals = []
        for row_number in (30, 31):
            row = bank_row(
                row_number,
                "350000",
                description="برداشت کارمزد کارمزد ساتنا",
                reference="",
            )
            row["direction"] = "withdrawal"
            withdrawals.append(row)
        for row_number in (32, 33):
            reversals.append(
                bank_row(
                    row_number,
                    "350000",
                    description="برگشت عملیات برداشت کارمزد کارمزد ساتنا",
                    reference="",
                )
            )

        rows, used = _match_bank_rows(
            [*withdrawals, *reversals],
            [],
            tolerance_days=2,
        )

        self.assertTrue(all(row["business_status"] == "reversed" for row in rows))
        self.assertTrue(
            all(row["match_method"] == "bank_reversal_batch" for row in rows)
        )
        self.assertEqual(used, set())

    def test_sepah_reversal_suffix_is_removed_before_batch_matching(self):
        withdrawals = []
        reversals = []
        for row_number in (172, 174):
            reversals.append(
                bank_row(
                    row_number,
                    "350000",
                    description=(
                        "برگشت عمليات برداشت کارمزد کارمزد ساتنا "
                        "| بازگشت عملیات"
                    ),
                    reference="",
                )
            )
        for row_number in (178, 180):
            row = bank_row(
                row_number,
                "350000",
                description=(
                    "برداشت کارمزد کارمزد ساتنا | طرف بدهکار"
                ),
                reference="",
            )
            row["direction"] = "withdrawal"
            withdrawals.append(row)

        rows, used = _match_bank_rows(
            [*reversals, *withdrawals],
            [],
            tolerance_days=2,
        )

        self.assertTrue(all(row["business_status"] == "reversed" for row in rows))
        self.assertTrue(
            all(row["match_method"] == "bank_reversal_batch" for row in rows)
        )
        self.assertEqual(used, set())

    def test_same_day_withdrawal_and_reversal_are_reconciled(self):
        withdrawal = bank_row(
            10,
            "3000000000",
            description="ثبت دستور پرداخت ساتنا 0503250611535758",
            reference="",
        )
        withdrawal["direction"] = "withdrawal"
        reversal = bank_row(
            11,
            "3000000000",
            description="برگشت عملیات ثبت دستور پرداخت ساتنا 0503250611535758",
            reference="",
        )

        rows, used = _match_bank_rows(
            [withdrawal, reversal],
            [],
            tolerance_days=2,
        )

        self.assertTrue(all(row["business_status"] == "reversed" for row in rows))
        self.assertTrue(
            all(row["match_method"] == "bank_reversal_pair" for row in rows)
        )
        self.assertEqual(used, set())

class ReconciliationStoreTests(unittest.TestCase):
    def test_report_is_saved_and_filtered(self):
        report = {
            "status": "success",
            "file": {"filename": "bank.csv"},
            "account": {"account_id": 1, "name": "بانک"},
            "matching_rules": {},
            "summary": {"business_counts": {"unposted": 1}},
            "rows": [
                {"row_number": 2, "business_status": "unposted"},
                {"row_number": 3, "business_status": "posted"},
            ],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            database_path = f"{temp_dir}/reconciliation.db"
            with patch(
                "app.services.reconciliation_store.settings.treasury_reconciliation_db",
                database_path,
            ):
                stored = save_reconciliation_report(report)
                summary = get_reconciliation_summary(
                    stored["reconciliation_id"]
                )
                rows = get_reconciliation_rows(
                    stored["reconciliation_id"],
                    "unposted",
                )

        self.assertEqual(summary["summary"]["business_counts"]["unposted"], 1)
        self.assertEqual(rows["returned_count"], 1)
        self.assertEqual(rows["rows"][0]["row_number"], 2)


class ReconciliationExportTests(unittest.TestCase):
    def test_excel_contains_summary_and_review_sheets(self):
        report = {
            "status": "success",
            "reconciliation_id": "a" * 32,
            "created_at": "2026-08-17T10:00:00+00:00",
            "file": {
                "filename": "bank.xls",
                "statement_iban": "IR123",
                "statement_period": "1405/01/01 تا 1405/05/26",
            },
            "account": {"account_id": 82, "number": "3101"},
            "bank_account": {
                "bank_account_id": 82,
                "number": "3101",
                "international_number": "IR123",
                "bank_name": "بانک سپه",
                "bank_branch_name": "شعبه مرکزی",
                "bank_branch_code": "100",
            },
            "summary": {
                "bank_row_count": 2,
                "bank_total_amount": 300000,
                "business_counts": {
                    "posted": 1,
                    "reversed": 0,
                    "internal_transfer": 0,
                    "needs_review": 0,
                    "unposted": 1,
                },
                "business_amounts": {
                    "posted": 100000,
                    "reversed": 0,
                    "internal_transfer": 0,
                    "needs_review": 0,
                    "unposted": 200000,
                },
                "direction_summary": {
                    "deposit": {"count": 1, "amount": 100000},
                    "withdrawal": {"count": 1, "amount": 200000},
                },
                "business_direction_summary": {
                    "posted": {
                        "deposit": {"count": 1, "amount": 100000},
                        "withdrawal": {"count": 0, "amount": 0},
                    },
                    "reversed": {
                        "deposit": {"count": 0, "amount": 0},
                        "withdrawal": {"count": 0, "amount": 0},
                    },
                    "internal_transfer": {
                        "deposit": {"count": 0, "amount": 0},
                        "withdrawal": {"count": 0, "amount": 0},
                    },
                    "needs_review": {
                        "deposit": {"count": 0, "amount": 0},
                        "withdrawal": {"count": 0, "amount": 0},
                    },
                    "unposted": {
                        "deposit": {"count": 0, "amount": 0},
                        "withdrawal": {"count": 1, "amount": 200000},
                    },
                },
            },
            "rows": [
                {
                    "row_number": 9,
                    "bank_date": "2026-08-16",
                    "direction": "deposit",
                    "amount": 100000,
                    "description": "واریز",
                    "business_status": "posted",
                    "business_status_fa": "سندخورده",
                    "confidence": 100,
                    "reasons": ["مبلغ و تاریخ برابر است"],
                    "matched_transaction": {
                        "document_number": "123",
                        "booking_date": "2026-08-16",
                        "amount": 100000,
                        "transaction_type": "receipt",
                        "item_id": 10,
                        "description": "سند دریافت",
                    },
                },
                {
                    "row_number": 10,
                    "bank_date": "2026-08-16",
                    "direction": "withdrawal",
                    "amount": 200000,
                    "description": "برداشت",
                    "business_status": "unposted",
                    "business_status_fa": "بدون سند",
                    "confidence": 0,
                    "reasons": ["سند متناظر پیدا نشد"],
                    "matched_transaction": None,
                },
            ],
            "unmatched_erp_transactions": [
                {
                    "transaction_id": "payment:20",
                    "transaction_type": "payment",
                    "document_number": "456",
                    "booking_date": "2026-08-16",
                    "amount": 300000,
                    "item_id": 20,
                    "item_number": "1",
                    "description": "سند پرداخت",
                    "counterpart_account_name": "تأمین‌کننده",
                }
            ],
        }

        output = build_reconciliation_workbook(report)
        workbook = load_workbook(BytesIO(output.getvalue()), data_only=False)

        self.assertEqual(
            workbook.sheetnames,
            [
                "خلاصه",
                "سندخورده",
                "برگشت یا خنثی‌شده",
                "انتقال داخلی شرکت",
                "نیازمند بررسی",
                "بدون سند احتمالی",
                "اسناد تطبیق‌قطعی‌نشده",
            ],
        )
        self.assertEqual(
            workbook["خلاصه"]["A1"].value,
            "گزارش مغایرت‌گیری حساب 3101 بانک سپه",
        )
        self.assertEqual(
            workbook["خلاصه"]["B5"].value,
            "حساب 3101 بانک سپه",
        )
        self.assertEqual(workbook["خلاصه"]["B6"].value, "بانک سپه")
        self.assertEqual(workbook["خلاصه"]["B7"].value, "3101")
        self.assertEqual(workbook["خلاصه"]["B13"].value, 1)
        self.assertEqual(workbook["خلاصه"]["C25"].value, 0)
        self.assertEqual(workbook["خلاصه"]["E25"].value, 200000)
        self.assertEqual(workbook["خلاصه"]["A28"].value, "نرخ تطبیق قطعی")
        self.assertEqual(
            workbook["خلاصه"]["A30"].value,
            "پوشش با تلورانس و موارد احتمالی",
        )
        self.assertEqual(workbook["خلاصه"]["B30"].value, 0.5)
        self.assertEqual(
            workbook["خلاصه"]["A31"].value,
            "نرخ تعیین تکلیف قطعی",
        )
        self.assertEqual(workbook["خلاصه"]["B31"].value, 0.5)
        self.assertEqual(workbook["سندخورده"]["Q2"].value, "123")
        self.assertEqual(workbook["بدون سند احتمالی"]["A2"].value, 10)
        self.assertEqual(
            workbook["اسناد تطبیق‌قطعی‌نشده"]["B2"].value,
            "456",
        )


if __name__ == "__main__":
    unittest.main()
