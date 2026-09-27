from __future__ import annotations

from datetime import datetime, timezone
from hashlib import sha256
import json
from pathlib import Path
import re
import shutil
from threading import RLock
from typing import Any

from openpyxl import load_workbook


PERSIAN_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789")
ARABIC_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")

HEADER_MAP = {
    "ردیف": "row_number",
    "شناسه حواله": "transfer_id",
    "شماره حواله": "transfer_number",
    "شعبه": "branch",
    "شماره سند": "document_number",
    "تاریخ ثبت": "registration_date_jalali",
    "مبلغ": "amount_rial",
    "کارمزد": "fee_rial",
    "بانک": "bank",
    "کد معین": "subsidiary_code",
    "کد سطح4": "level4_code",
    "کد سطح5": "level5_code",
    "کد سطح6": "level6_code",
    "نام معین": "subsidiary_name",
    "نام سطح4": "level4_name",
    "نام سطح5": "level5_name",
    "نام سطح6": "level6_name",
    "بابت": "purpose",
    "توضیحات": "description",
    "شماره فاکتور": "invoice_number",
    "تاریخ سند": "document_date_jalali",
    "شماره درخواست پرداخت": "payment_request_number",
}
REQUIRED_HEADERS = {
    "شناسه حواله",
    "شماره حواله",
    "تاریخ ثبت",
    "مبلغ",
    "بانک",
    "بابت",
}
_LOCK = RLock()


def _header(value: Any) -> str:
    return " ".join(str(value or "").replace("\u200c", " ").split()).strip()


def _identifier(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _number(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    normalized = (
        str(value)
        .translate(PERSIAN_DIGITS)
        .translate(ARABIC_DIGITS)
        .replace(",", "")
        .replace("٬", "")
        .strip()
    )
    try:
        return float(normalized)
    except ValueError as exc:
        raise ValueError(f"مبلغ نامعتبر است: {value}") from exc


def _jalali_date(value: Any, field_name: str, row_number: int) -> str | None:
    text = str(value or "").translate(PERSIAN_DIGITS).translate(ARABIC_DIGITS).strip()
    if not text:
        return None
    match = re.search(r"(1[34]\d{2})\s*[/_.-]\s*(\d{1,2})\s*[/_.-]\s*(\d{1,2})", text)
    if not match:
        raise ValueError(f"{field_name} در ردیف {row_number} شمسی معتبر نیست: {value}")
    year, month, day = map(int, match.groups())
    if not (1 <= month <= 12 and 1 <= day <= 31):
        raise ValueError(f"{field_name} در ردیف {row_number} معتبر نیست: {value}")
    return f"{year:04d}/{month:02d}/{day:02d}"


class KaramadPaymentTransferService:
    """Manual Excel source for historical paid transfers exported from KarAmand."""

    def __init__(self, root: Path | None = None):
        self.root = root or Path("data") / "karamad_payment_transfers"
        self.root.mkdir(parents=True, exist_ok=True)
        self.archive = self.root / "imports"
        self.archive.mkdir(parents=True, exist_ok=True)
        self.state_file = self.root / "state.json"

    def parse(self, path: Path) -> list[dict[str, Any]]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            header_row = None
            headers: list[str] = []
            for row_number, values in enumerate(
                sheet.iter_rows(min_row=1, max_row=min(10, sheet.max_row or 10), values_only=True),
                start=1,
            ):
                normalized = [_header(value) for value in values]
                if REQUIRED_HEADERS.issubset(set(normalized)):
                    header_row, headers = row_number, normalized
                    break
            if header_row is None:
                missing = "، ".join(sorted(REQUIRED_HEADERS))
                raise ValueError(f"سربرگ فایل کارآمد پیدا نشد. ستون‌های ضروری: {missing}")

            unknown = [name for name in headers if name and name not in HEADER_MAP]
            if unknown:
                raise ValueError(f"ستون ناشناخته در فایل کارآمد: {'، '.join(unknown)}")

            column_keys = [HEADER_MAP.get(name) for name in headers]
            records: list[dict[str, Any]] = []
            seen_ids: set[str] = set()
            for excel_row, values in enumerate(
                sheet.iter_rows(min_row=header_row + 1, values_only=True),
                start=header_row + 1,
            ):
                if not any(value not in (None, "") for value in values):
                    continue
                raw = {
                    key: values[index] if index < len(values) else None
                    for index, key in enumerate(column_keys)
                    if key
                }
                transfer_id = _identifier(raw.get("transfer_id"))
                if not transfer_id:
                    raise ValueError(f"شناسه حواله در ردیف {excel_row} خالی است.")
                if transfer_id in seen_ids:
                    raise ValueError(f"شناسه حواله {transfer_id} داخل همین فایل تکراری است.")
                seen_ids.add(transfer_id)
                registration_date = _jalali_date(
                    raw.get("registration_date_jalali"), "تاریخ ثبت", excel_row
                )
                amount = _number(raw.get("amount_rial"))
                fee = _number(raw.get("fee_rial"))
                if amount < 0 or fee < 0:
                    raise ValueError(f"مبلغ یا کارمزد منفی در ردیف {excel_row} مجاز نیست.")
                record = {
                    key: _identifier(value) if key not in {"amount_rial", "fee_rial"} else value
                    for key, value in raw.items()
                }
                record.update(
                    {
                        "transfer_id": transfer_id,
                        "registration_date_jalali": registration_date,
                        "registration_month": registration_date[:7] if registration_date else None,
                        "document_date_jalali": _jalali_date(
                            raw.get("document_date_jalali"), "تاریخ سند", excel_row
                        ),
                        "amount_rial": amount,
                        "fee_rial": fee,
                        "source_row": excel_row,
                    }
                )
                records.append(record)
            if not records:
                raise ValueError("فایل کارآمد هیچ حواله‌ای ندارد.")
            return records
        finally:
            workbook.close()

    def _state(self) -> dict[str, Any]:
        if not self.state_file.exists():
            return {"records": {}, "imports": []}
        try:
            state = json.loads(self.state_file.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise RuntimeError("مخزن حواله‌های کارآمد قابل خواندن نیست.") from exc
        state.setdefault("records", {})
        state.setdefault("imports", [])
        return state

    def _write_state(self, state: dict[str, Any]) -> None:
        temporary = self.state_file.with_suffix(".tmp")
        temporary.write_text(
            json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        temporary.replace(self.state_file)

    @staticmethod
    def _business_record(record: dict[str, Any]) -> dict[str, Any]:
        return {
            key: value
            for key, value in record.items()
            if key not in {"source_filename", "imported_at", "source_row"}
        }

    def import_workbook(self, path: Path, original_filename: str) -> dict[str, Any]:
        records = self.parse(path)
        content_hash = sha256(path.read_bytes()).hexdigest()
        imported_at = datetime.now(timezone.utc).isoformat()
        safe_filename = re.sub(r"[^A-Za-z0-9_.-]+", "_", Path(original_filename).name)
        with _LOCK:
            state = self._state()
            existing = state["records"]
            inserted = updated = unchanged = 0
            for record in records:
                transfer_id = record["transfer_id"]
                previous = existing.get(transfer_id)
                enriched = {
                    **record,
                    "source_filename": original_filename,
                    "imported_at": imported_at,
                }
                if previous is None:
                    inserted += 1
                elif self._business_record(previous) == self._business_record(enriched):
                    unchanged += 1
                else:
                    updated += 1
                existing[transfer_id] = enriched

            archive_name = f"{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}__{safe_filename}"
            shutil.copy2(path, self.archive / archive_name)
            import_log = {
                "filename": original_filename,
                "stored_file": archive_name,
                "content_sha256": content_hash,
                "imported_at": imported_at,
                "rows_in_file": len(records),
                "inserted_count": inserted,
                "updated_count": updated,
                "unchanged_count": unchanged,
            }
            state["imports"].append(import_log)
            state["imports"] = state["imports"][-100:]
            self._write_state(state)
        return {
            "status": "success",
            "import": import_log,
            "summary": self.summary(),
            "message": (
                f"{len(records)} ردیف خوانده شد؛ {inserted} جدید، "
                f"{updated} به‌روزشده و {unchanged} بدون تغییر."
            ),
        }

    def _filtered_records(
        self, month: str | None = None, search: str | None = None
    ) -> list[dict[str, Any]]:
        records = list(self._state()["records"].values())
        normalized_month = (month or "").strip()
        if normalized_month and not re.fullmatch(r"1[34]\d{2}/(?:0[1-9]|1[0-2])", normalized_month):
            raise ValueError("ماه باید به‌شکل 1405/06 باشد.")
        if normalized_month:
            records = [row for row in records if row.get("registration_month") == normalized_month]
        normalized_search = " ".join((search or "").casefold().split())
        if normalized_search:
            searchable = (
                "transfer_id", "transfer_number", "branch", "document_number",
                "bank", "purpose", "description", "invoice_number",
                "payment_request_number", "level4_name", "level5_name", "level6_name",
            )
            records = [
                row
                for row in records
                if normalized_search
                in " ".join(str(row.get(key) or "") for key in searchable).casefold()
            ]
        return sorted(
            records,
            key=lambda row: (
                row.get("registration_date_jalali") or "",
                int(row["transfer_id"]) if str(row["transfer_id"]).isdigit() else 0,
            ),
            reverse=True,
        )

    def summary(self, month: str | None = None) -> dict[str, Any]:
        records = self._filtered_records(month=month)
        months: dict[str, dict[str, Any]] = {}
        for row in self._state()["records"].values():
            key = row.get("registration_month") or "نامشخص"
            item = months.setdefault(key, {"month": key, "count": 0, "amount_rial": 0.0})
            item["count"] += 1
            item["amount_rial"] += float(row.get("amount_rial") or 0)
        dates = [row.get("registration_date_jalali") for row in records if row.get("registration_date_jalali")]
        state = self._state()
        return {
            "record_count": len(records),
            "total_amount_rial": sum(float(row.get("amount_rial") or 0) for row in records),
            "total_fee_rial": sum(float(row.get("fee_rial") or 0) for row in records),
            "min_registration_date_jalali": min(dates) if dates else None,
            "max_registration_date_jalali": max(dates) if dates else None,
            "selected_month": month,
            "available_months": sorted(months.values(), key=lambda item: item["month"], reverse=True),
            "last_import": state["imports"][-1] if state["imports"] else None,
            "cashflow_included": False,
            "cashflow_rule": "حواله پرداخت‌شده تاریخی است و دوباره از Cash Flow آینده کم نمی‌شود.",
        }

    def report(
        self,
        month: str | None = None,
        search: str | None = None,
        limit: int = 200,
        offset: int = 0,
    ) -> dict[str, Any]:
        limit, offset = max(1, min(int(limit), 2000)), max(0, int(offset))
        records = self._filtered_records(month=month, search=search)
        page = records[offset : offset + limit]
        return {
            "status": "success",
            "report_type": "karamad_paid_payment_transfers",
            "source": "KarAmand Excel",
            "filters": {"month": month, "search": search},
            "summary": self.summary(month=month),
            "pagination": {
                "limit": limit,
                "offset": offset,
                "returned_count": len(page),
                "total_count": len(records),
                "has_more": offset + len(page) < len(records),
            },
            "transfers": page,
        }
