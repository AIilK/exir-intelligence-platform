"""پرداخت‌های ثبت‌نشده/نامشخص مشتریان کارآمد (V172) — بدون صورت‌حساب بانک.

مشتری می‌گوید «پرداخت کرده‌ام» ولی بدهی‌اش در کارآمد مانده است. بدون صورت‌حساب بانکِ هیبرید،
سه جای داده‌ای که چنین پولی ممکن است در آن گم شده باشد بررسی می‌شود:

1. واریز بی‌صاحب کارآمد: دریافت نقدی/حواله (tblCashD با PriceN≠0، tblDraftD) که به هیچ مشتری‌ای
   ثبت نشده — بدون تفصیلی، روی حساب «هیبرید من‌ها» (کلاس ۲۸)، «اشخاص»، «سایر» یا حساب‌های
   «کنترل وصول/علی‌الحساب/نامشخص». حساب بانک/صندوق (کلاس ۱ و ۲) انتقال داخلی است و حذف می‌شود.
2. ثبت به نام مشتری دیگر: شرح دریافت شماره فاکتورهایی را می‌آورد («فاکتور 1145 و 1146») که در همان
   شعبه مال مشتری دیگری است.
3. واریز مستقیم به حساب‌های شرکت در راهکاران (دسته «واریز مستقیم مشتری») که در کارآمد معادل هم‌مبلغ
   (±۱۰ روز) ندارد و نام واریزکننده با یک بدهکار کارآمد یکی است.

برای هر مورد، مشتری‌های بدهکارِ همان شعبه با این شواهد رتبه‌بندی می‌شوند: شماره فاکتور در شرح، برابری
مبلغ با مانده یک فاکتور باز یا کل بدهی، و شباهت نام. خروجی فقط پیشنهاد بررسی است، نه ثبت خودکار.
"""

from __future__ import annotations

import re
from collections import defaultdict
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Any

from sqlalchemy import text

from app.utils.jalali import format_jalali_date

TRANSFER_CLASSES = {1, 2}  # حساب‌های بانکی، صندوق‌ها
UNASSIGNED_CLASSES = {3, 16, 28}  # اشخاص، سایر، هیبرید من‌ها
SUSPENSE_WORDS = ("کنترل وصول", "علی الحساب", "علی‌الحساب", "نامشخص", "متفرقه", "واریزی نامشخص")
GENERIC_NAME_WORDS = {
    "آرایشی", "ارایشی", "بهداشتی", "گالری", "فروشگاه", "داروخانه", "دکتر", "عطر", "عمده", "فروشی", "شعبه",
    "هیبرید", "انبار", "پخش", "شرکت", "زنجیره", "زنجیره‌ای", "سوپر", "مارکت", "لوازم", "آقای", "خانم", "محمد",
    "علی", "رضا", "حسین", "مهدی", "احمد", "سید", "بابت", "فاکتور", "واریز", "تهران",
}
AMOUNT_TOLERANCE_RIAL = 10_000
# برابری مبلغ فقط برای مبالغ معنادار و وقتی مبلغ در شعبه یکتا باشد شاهد حساب می‌شود.
MIN_AMOUNT_EVIDENCE_RIAL = 5_000_000
MAX_SAME_AMOUNT_CUSTOMERS = 2
INVOICE_LOOKBACK_DAYS = 365
# پول ورودی از خود شرکت/شرکا یا انتقال بین حساب‌ها؛ پرداخت مشتری نیست.
INTERNAL_FUNDING_WORDS = ("جاری", "کادوس", "شرکا", "انتقال", "تنخواه", "کارمزد", "سود سپرده", "وام", "تسهیلات")


def is_internal_funding(*texts: Any) -> bool:
    from app.services.karamad_live_cash_draft_service import _classify

    joined = " ".join(_clean(t) for t in texts)
    return any(w in joined for w in INTERNAL_FUNDING_WORDS) or _classify(joined, "")[0] != "operational"


INVOICE_PATTERN = re.compile(
    r"(?:فاکتور|فاكتور|فاک|ف\.)\s*(?:های|ها|شماره|ش\.?)?\s*[:：\-]?\s*((?:\d{1,6}\s*(?:[-،,/و]|\sو\s|\s)?\s*)+)")


def _clean(value: Any) -> str:
    value = str(value or "").replace("ي", "ی").replace("ك", "ک").replace("‌", " ")
    return re.sub(r"\s+", " ", value).strip()


def _tokens(value: Any) -> set[str]:
    return {t for t in re.split(r"[\s\-\(\)/.،,_:]+", _clean(value))
            if len(t) >= 3 and not t.isdigit() and t not in GENERIC_NAME_WORDS}


def invoice_codes(*texts: Any) -> set[str]:
    """Invoice numbers explicitly written after «فاکتور»; dates/years and long IDs are ignored."""

    codes: set[str] = set()
    for value in texts:
        cleaned = _clean(value)
        cleaned = re.sub(r"\d{4}/\d{1,2}/\d{1,2}", " ", cleaned)  # dates like 1405/04/17
        for match in INVOICE_PATTERN.finditer(cleaned):
            for number in re.findall(r"\d+", match.group(1)):
                # شماره فاکتور ۲ تا ۶ رقمی؛ شناسه‌های بلند و سال‌ها شماره فاکتور نیستند.
                if not 2 <= len(number) <= 6 or re.fullmatch(r"1[34]\d\d", number):
                    continue
                codes.add(str(int(number)))
    return codes


def _name_score(name: str, *texts: Any) -> float:
    target = _tokens(name)
    if not target:
        return 0.0
    found = set().union(*(_tokens(t) for t in texts))
    return len(target & found) / len(target)


class KaramadUnrecordedPaymentService:
    def __init__(self, engine=None):
        self._engine = engine

    @property
    def engine(self):
        if self._engine is None:
            from app.database.karamad_sqlserver import get_karamad_sqlserver_engine
            self._engine = get_karamad_sqlserver_engine()
        return self._engine

    # ------------------------------------------------------------------ data
    def _load(self, since: date) -> dict[str, Any]:
        from app.services.karamad_sales_network_service import ledger_positions

        params = {"since": since}
        with self.engine.connect() as c:
            branches = {int(r[0]): r[1] for r in c.execute(text("SELECT ID, Name FROM dbo.[tblBranch]"))}
            customers = {int(r["DLRef"]): dict(r) for r in c.execute(text(
                "SELECT ID, DLRef, Name, Code, BranchRef, Mobile FROM dbo.[tblCustomer] WHERE DLRef IS NOT NULL"
            )).mappings()}
            invoices = [dict(r) for r in c.execute(text("""
                SELECT f.ID, f.Code, f.DateE, f.BranchRef, cu.DLRef, f.FactorPriceP, v.Name AS VisitorName,
                       f.FactorPriceP - (ISNULL(p.CashN,0) + ISNULL(p.CashE,0) + ISNULL(p.CashT,0) + ISNULL(p.Draft,0)
                         + ISNULL(p.Cheque,0) + ISNULL(p.FactorB,0) + ISNULL(p.SharedFactor,0) + ISNULL(p.SharedCustomerDebt,0)
                         + ISNULL(p.SharedReturnCheque,0) + ISNULL(p.SharedReturn,0) + ISNULL(p.Barter,0)) AS Remaining
                FROM dbo.[tblFactorF] f
                JOIN dbo.[tblCustomer] cu ON cu.ID = f.CustomerRef
                LEFT JOIN dbo.[vwFactorFPayoff] p ON p.ID = f.ID
                LEFT JOIN dbo.[tblVisitor] v ON v.ID = f.VisitorRef
                WHERE f.DateE >= DATEADD(day, -400, :since)
            """), params).mappings()]
            inflows = [dict(r) for r in c.execute(text("""
                SELECT k.kind, k.inx, k.Code, k.BookDate, k.Amount, k.Behalf, k.Description, k.BranchRef, k.DLRef,
                       dl.Name AS DLName, dl.ClassRef, bk.Name AS BankName, u.Name AS UserName
                FROM (
                    SELECT N'حواله بانکی' AS kind, inx, CAST(Code AS nvarchar(50)) AS Code, BookDate, Price AS Amount,
                           Behalf, Description, BranchRef, DLRef, BankRef, UserRef FROM dbo.[tblDraftD]
                    UNION ALL
                    SELECT N'دریافت نقدی', inx, CAST(Code AS nvarchar(50)), BookDate, PriceN, Behalf, Description,
                           BranchRef, DLRef, NULL, UserRef FROM dbo.[tblCashD] WHERE ISNULL(PriceN, 0) <> 0
                ) k
                LEFT JOIN dbo.[tblDL] dl ON dl.ID = k.DLRef
                LEFT JOIN dbo.[tblUser] u ON u.ID = k.UserRef
                OUTER APPLY (SELECT TOP (1) a.Name FROM dbo.[tblAccBank] a WHERE a.Code = k.BankRef) bk
                WHERE k.BookDate >= :since
            """), params).mappings()]
        positions = ledger_positions()
        return {"branches": branches, "customers": customers, "invoices": invoices,
                "inflows": inflows, "positions": positions}

    # -------------------------------------------------------------- matching
    def build(self, since: date | None = None) -> dict[str, Any]:
        today = date.today()
        since = since or (today - timedelta(days=365))
        data = self._load(since)
        branches, customers, positions = data["branches"], data["customers"], data["positions"]

        debt = {dl: float(p.get("balance_rial") or 0) for dl, p in positions.items()}
        debtors = {dl: c for dl, c in customers.items() if debt.get(dl, 0) > AMOUNT_TOLERANCE_RIAL}
        # شماره فاکتور هر سال مالی از نو شروع می‌شود؛ پس برای هر (شعبه، شماره) همه فاکتورها نگه داشته
        # می‌شوند و در زمان تطبیق فقط فاکتور قبل از تاریخ واریز (حداکثر یک سال) معتبر است.
        invoices_by_code: dict[tuple[int, str], list[dict[str, Any]]] = defaultdict(list)
        open_by_customer: dict[int, list[dict[str, Any]]] = defaultdict(list)
        last_visitor: dict[int, str] = {}
        for inv in sorted(data["invoices"], key=lambda r: (r["DateE"] or date.min)):
            branch, dl = int(inv["BranchRef"] or 0), int(inv["DLRef"])
            invoices_by_code[(branch, str(inv["Code"]))].append(inv)
            if inv.get("VisitorName"):
                last_visitor[dl] = inv["VisitorName"]
            if dl in debtors and float(inv["Remaining"] or 0) > AMOUNT_TOLERANCE_RIAL:
                open_by_customer[dl].append(inv)
        debtors_by_branch: dict[int, list[int]] = defaultdict(list)
        for dl, cust in debtors.items():
            debtors_by_branch[int(cust["BranchRef"] or 0)].append(dl)

        def relevant_invoices(branch: int, code: str, when: date | None) -> list[dict[str, Any]]:
            """Invoices with this number issued in the year before the payment, newest first."""
            rows = []
            for inv in invoices_by_code.get((branch, code), []):
                issued = inv["DateE"]
                if when and issued and not (when - timedelta(days=INVOICE_LOOKBACK_DAYS) <= issued <= when):
                    continue
                rows.append(inv)
            return sorted(rows, key=lambda i: i["DateE"] or date.min, reverse=True)

        amount_matches: dict[tuple[int, int], set[int]] = defaultdict(set)
        for dl, cust in debtors.items():
            branch = int(cust["BranchRef"] or 0)
            amount_matches[(branch, round(debt.get(dl, 0) / AMOUNT_TOLERANCE_RIAL))].add(dl)
            for inv in open_by_customer.get(dl, []):
                amount_matches[(branch, round(float(inv["Remaining"]) / AMOUNT_TOLERANCE_RIAL))].add(dl)

        def distinctive_amount(branch: int, amount: float) -> bool:
            if amount < MIN_AMOUNT_EVIDENCE_RIAL:
                return False
            key = round(amount / AMOUNT_TOLERANCE_RIAL)
            same = set().union(*(amount_matches.get((branch, key + d), set()) for d in (-1, 0, 1)))
            return 0 < len(same) <= MAX_SAME_AMOUNT_CUSTOMERS

        def candidates(branch: int, amount: float, when: date | None, *texts: Any,
                       exclude: int | None = None) -> list[dict[str, Any]]:
            codes = invoice_codes(*texts)
            found: dict[int, dict[str, Any]] = {}

            def add(dl: int, score: int, reason: str) -> None:
                if dl == exclude or dl not in debtors:
                    return
                item = found.setdefault(dl, {"dl_ref": dl, "score": 0, "reasons": []})
                item["score"] = min(100, item["score"] + score)
                item["reasons"].append(reason)

            for code in codes:
                # فقط فاکتوری که هنوز تسویه نشده شاهد پرداخت ثبت‌نشده است.
                for inv in relevant_invoices(branch, code, when)[:1]:
                    if float(inv["Remaining"] or 0) > AMOUNT_TOLERANCE_RIAL:
                        add(int(inv["DLRef"]), 60, f"فاکتور {code} ({format_jalali_date(inv['DateE'])}) در شرح آمده و هنوز تسویه نشده است")
            use_amount = distinctive_amount(branch, amount)
            for dl in debtors_by_branch.get(branch, []):
                if use_amount and abs(debt.get(dl, 0) - amount) <= AMOUNT_TOLERANCE_RIAL:
                    add(dl, 30, "مبلغ با کل بدهی فعلی مشتری برابر است")
                for inv in open_by_customer.get(dl, []) if use_amount else []:
                    if abs(float(inv["Remaining"]) - amount) <= AMOUNT_TOLERANCE_RIAL and (when is None or not inv["DateE"] or inv["DateE"] <= when):
                        add(dl, 35, f"مبلغ با مانده فاکتور {inv['Code']} ({format_jalali_date(inv['DateE'])}) برابر است")
                        break
                score = _name_score(debtors[dl]["Name"], *texts)
                if score >= 0.6:
                    add(dl, 30, "نام مشتری در شرح واریز آمده است")
            # واریزکننده ممکن است مشتری هیبرید دیگری باشد (مثلاً واریز به حساب ملت تهران)؛
            # در شعبه‌های دیگر فقط تطابق نام قوی (حداقل دو کلمه متمایز، ≥۷۵٪) پذیرفته می‌شود.
            text_tokens = set().union(*(_tokens(t) for t in texts))
            if len(text_tokens) >= 2:
                for dl, cust in debtors.items():
                    if int(cust["BranchRef"] or 0) == branch:
                        continue
                    target = _tokens(cust["Name"])
                    if len(target) >= 2 and len(target & text_tokens) / len(target) >= 0.75:
                        add(dl, 30, f"نام مشتری (هیبرید {branches.get(int(cust['BranchRef'] or 0)) or '—'}) در شرح واریز آمده است")
            ranked = sorted(found.values(), key=lambda x: -x["score"])
            return [x for x in ranked if x["score"] >= 30][:3]

        def customer_view(dl: int) -> dict[str, Any]:
            cust = customers.get(dl) or {}
            opens = open_by_customer.get(dl, [])
            oldest = min((i["DateE"] for i in opens if i["DateE"]), default=None)
            return {"dl_ref": dl, "name": cust.get("Name"), "code": cust.get("Code"),
                    "branch": branches.get(int(cust.get("BranchRef") or 0)), "mobile": cust.get("Mobile"),
                    "visitor": last_visitor.get(dl), "debt_rial": debt.get(dl, 0.0),
                    "open_invoice_count": len(opens), "oldest_open_invoice_jalali": format_jalali_date(oldest)}

        def inflow_view(r: dict[str, Any]) -> dict[str, Any]:
            return {"kind": r["kind"], "number": r["Code"], "date_jalali": format_jalali_date(r["BookDate"]),
                    "amount_rial": float(r["Amount"] or 0), "branch": branches.get(int(r["BranchRef"] or 0)),
                    "account": r["DLName"] or "بدون حساب طرف", "bank": r.get("BankName"),
                    "behalf": _clean(r["Behalf"]), "description": _clean(r["Description"]), "registered_by": r["UserName"]}

        # 1) unassigned money
        unassigned = []
        internal_funding = {"count": 0, "amount_rial": 0.0}
        for r in data["inflows"]:
            dl = r["DLRef"]
            if dl is not None and int(dl) in customers:
                continue
            klass = r["ClassRef"]
            suspense = any(w in _clean(r["DLName"]) for w in SUSPENSE_WORDS)
            if dl is not None and klass in TRANSFER_CLASSES and not suspense:
                continue
            if dl is not None and klass not in UNASSIGNED_CLASSES and not suspense:
                continue
            if is_internal_funding(r["Behalf"], r["Description"]):
                internal_funding["count"] += 1
                internal_funding["amount_rial"] += float(r["Amount"] or 0)
                continue
            branch = int(r["BranchRef"] or 0)
            unassigned.append({**inflow_view(r), "candidates": [
                {**customer_view(c["dl_ref"]), "score": c["score"], "reasons": c["reasons"]}
                for c in candidates(branch, float(r["Amount"] or 0), r["BookDate"], r["Behalf"], r["Description"],
                                    r["DLName"] if klass == 3 else "")]})

        # 2) credited to another customer
        misallocated = []
        for r in data["inflows"]:
            dl = r["DLRef"]
            if dl is None or int(dl) not in customers:
                continue
            branch = int(r["BranchRef"] or 0)
            codes = invoice_codes(r["Behalf"], r["Description"])
            if not codes:
                continue
            named = [inv for code in codes for inv in relevant_invoices(branch, code, r["BookDate"])]
            # اگر هر کدام از فاکتورهای نام‌برده مال همین مشتری باشد، ثبت درست است.
            if not named or any(int(inv["DLRef"]) == int(dl) for inv in named):
                continue
            others = sorted({int(inv["DLRef"]) for inv in named
                             if int(inv["DLRef"]) in debtors and float(inv["Remaining"] or 0) > AMOUNT_TOLERANCE_RIAL})
            if not others:
                continue
            misallocated.append({**inflow_view(r), "credited_to": customer_view(int(dl)), "invoice_codes": sorted(codes),
                                 "candidates": [{**customer_view(o), "score": 80,
                                                 "reasons": ["فاکتور نام‌برده در شرح مال این مشتری است و هنوز تسویه نشده"]} for o in others]})

        # 3) direct deposits to the company's Rahkaran accounts with no Karamad receipt
        direct = self._rahkaran_direct(data, debtors, debt, open_by_customer, customer_view, since)

        found_by_customer: dict[int, list[dict[str, Any]]] = defaultdict(list)
        for source, rows in (("واریز بی‌صاحب کارآمد", unassigned), ("ثبت به نام مشتری دیگر", misallocated),
                             ("واریز مستقیم به حساب شرکت (راهکاران)", direct)):
            for row in rows:
                for cand in row["candidates"]:
                    found_by_customer[cand["dl_ref"]].append({"source": source, "score": cand["score"],
                                                              "reasons": cand["reasons"], "payment": row})
        customers_with_evidence = []
        for dl, items in found_by_customer.items():
            items.sort(key=lambda x: -x["score"])
            customers_with_evidence.append({**customer_view(dl), "evidence_count": len(items),
                                            "evidence_amount_rial": sum(x["payment"]["amount_rial"] for x in items),
                                            "best_score": items[0]["score"], "evidence": items})
        customers_with_evidence.sort(key=lambda x: (-x["best_score"], -x["debt_rial"]))

        by_branch: dict[str, dict[str, Any]] = defaultdict(lambda: {"debtors": 0, "debt_rial": 0.0, "unassigned": 0,
                                                                    "unassigned_rial": 0.0, "customers_with_evidence": 0})
        for dl in debtors:
            b = by_branch[branches.get(int(debtors[dl]["BranchRef"] or 0)) or "—"]
            b["debtors"] += 1
            b["debt_rial"] += debt[dl]
        for row in unassigned:
            b = by_branch[row["branch"] or "—"]
            b["unassigned"] += 1
            b["unassigned_rial"] += row["amount_rial"]
        for row in customers_with_evidence:
            by_branch[row["branch"] or "—"]["customers_with_evidence"] += 1

        return {
            "status": "success",
            "generated_at": datetime.now().isoformat(timespec="minutes"),
            "since_jalali": format_jalali_date(since),
            "summary": {
                "debtor_count": len(debtors), "debt_rial": sum(debt[d] for d in debtors),
                "unassigned_count": len(unassigned), "unassigned_rial": sum(r["amount_rial"] for r in unassigned),
                "unassigned_with_candidate": sum(1 for r in unassigned if r["candidates"]),
                "misallocated_count": len(misallocated), "misallocated_rial": sum(r["amount_rial"] for r in misallocated),
                "direct_count": len(direct), "direct_rial": sum(r["amount_rial"] for r in direct),
                "customers_with_evidence": len(customers_with_evidence),
                "internal_funding_excluded_count": internal_funding["count"],
                "internal_funding_excluded_rial": internal_funding["amount_rial"],
            },
            "by_branch": [{"branch": k, **v} for k, v in sorted(by_branch.items(), key=lambda kv: -kv[1]["debt_rial"])],
            "unassigned": sorted(unassigned, key=lambda r: (-(r["candidates"][0]["score"] if r["candidates"] else 0), -r["amount_rial"])),
            "misallocated": misallocated,
            "direct": direct,
            "customers": customers_with_evidence,
        }

    def _rahkaran_direct(self, data, debtors, debt, open_by_customer, customer_view, since) -> list[dict[str, Any]]:
        try:
            from app.services.treasury_service import get_customer_b2b_remittances
            rows = get_customer_b2b_remittances(limit=20000).get("rows") or []
        except Exception:  # noqa: BLE001 — Rahkaran unavailable: Karamad findings still stand
            return []
        karamad_by_amount: dict[int, list[date]] = defaultdict(list)
        for r in data["inflows"]:
            karamad_by_amount[round(float(r["Amount"] or 0))].append(r["BookDate"])
        out = []
        for r in rows:
            day = str(r.get("deposit_date") or "")[:10]
            if r.get("category") != "واریز مستقیم مشتری" or not day or day < since.isoformat():
                continue
            when = date.fromisoformat(day)
            amount = float(r.get("amount_rial") or 0)
            if any(abs((d - when).days) <= 10 for d in karamad_by_amount.get(round(amount), []) if d):
                continue
            payer = r.get("customer_name") or ""
            cands = []
            for dl, cust in debtors.items():
                score = _name_score(cust["Name"], payer, r.get("description"))
                if score < 0.75:
                    continue
                points, reasons = 40, ["نام واریزکننده در راهکاران با نام مشتری کارآمد یکی است"]
                if abs(debt.get(dl, 0) - amount) <= AMOUNT_TOLERANCE_RIAL:
                    points += 30
                    reasons.append("مبلغ با کل بدهی فعلی مشتری برابر است")
                elif any(abs(float(i["Remaining"]) - amount) <= AMOUNT_TOLERANCE_RIAL for i in open_by_customer.get(dl, [])):
                    points += 30
                    reasons.append("مبلغ با مانده یک فاکتور باز مشتری برابر است")
                else:
                    continue  # تشابه نام به‌تنهایی (مثلاً یک نام خانوادگی مشترک) کافی نیست
                cands.append({**customer_view(dl), "score": points, "reasons": reasons})
            if cands:
                cands.sort(key=lambda x: -x["score"])
                out.append({"kind": "واریز راهکاران", "number": r.get("receipt_number"),
                            "date_jalali": r.get("deposit_date_jalali"), "amount_rial": amount,
                            "branch": None, "account": payer, "bank": r.get("bank_branch_name"),
                            "behalf": "", "description": _clean(r.get("description")), "registered_by": None,
                            "candidates": cands[:3]})
        return out

    # ----------------------------------------------------------------- excel
    def workbook(self, since: date | None = None) -> BytesIO:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        report = self.build(since)
        wb = Workbook()
        header_fill, header_font = PatternFill("solid", fgColor="1F4E78"), Font(color="FFFFFF", bold=True)

        def sheet(title: str, headers: list[str], rows: list[list[Any]], widths: list[int], money_cols: tuple[int, ...] = ()):
            ws = wb.create_sheet(title)
            ws.sheet_view.rightToLeft = True
            ws.append(headers)
            for cell in ws[1]:
                cell.fill, cell.font = header_fill, header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for row in rows:
                ws.append(row)
            for col in money_cols:
                for cell in ws[get_column_letter(col)][1:]:
                    cell.number_format = "#,##0"
            for i, w in enumerate(widths, start=1):
                ws.column_dimensions[get_column_letter(i)].width = w
            for r in ws.iter_rows(min_row=2):
                for cell in r:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            ws.freeze_panes = "A2"
            if rows:
                ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
            return ws

        toman = lambda rial: round(float(rial or 0) / 10)  # noqa: E731
        s = report["summary"]
        summary = wb.active
        summary.title = "خلاصه و راهنما"
        summary.sheet_view.rightToLeft = True
        lines = [
            ["پرداخت‌های ثبت‌نشده احتمالی مشتریان کارآمد", ""],
            [f"تاریخ گزارش: {format_jalali_date(date.today())} — بازه بررسی: از {report['since_jalali']}", ""],
            ["", ""],
            ["مشتری بدهکار", s["debtor_count"]], ["جمع بدهی (تومان)", toman(s["debt_rial"])],
            ["واریز بی‌صاحب در کارآمد", s["unassigned_count"]], ["مبلغ واریز بی‌صاحب (تومان)", toman(s["unassigned_rial"])],
            ["واریز بی‌صاحب با مشتری پیشنهادی", s["unassigned_with_candidate"]],
            ["واریز ثبت‌شده به نام مشتری دیگر", s["misallocated_count"]],
            ["مبلغ ثبت به نام مشتری دیگر (تومان)", toman(s["misallocated_rial"])],
            ["واریز مستقیم به حساب شرکت بدون معادل در کارآمد", s["direct_count"]],
            ["مشتری دارای شواهد پرداخت", s["customers_with_evidence"]],
            ["دریافت بی‌حساب داخلی (جاری/کادوس/شرکا/انتقال) — کنار گذاشته شد", s["internal_funding_excluded_count"]],
            ["", ""],
            ["راهنما", ""],
            ["واریز بی‌صاحب", "پولی که در کارآمد دریافت شده ولی به هیچ مشتری ثبت نشده (بدون حساب طرف، روی حساب هیبرید، اشخاص، سایر یا کنترل وصول)."],
            ["ثبت به نام مشتری دیگر", "شرح دریافت شماره فاکتورهای مشتری دیگری در همان شعبه را آورده است."],
            ["واریز مستقیم به شرکت", "واریز مشتری به حساب‌های شرکت در راهکاران که در کارآمد معادل هم‌مبلغ (±۱۰ روز) ندارد و نام واریزکننده با بدهکار کارآمد یکی است."],
            ["امتیاز", "شماره فاکتور در شرح ۶۰، برابری مبلغ با مانده فاکتور ۳۵، برابری با کل بدهی ۳۰، نام مشتری در شرح ۳۰ (حداکثر ۱۰۰). فقط پیشنهاد بررسی است، نه ثبت خودکار."],
            ["محدودیت", "پول نقد یا چکی که به ویزیتور/موزع داده شده و به بانک نرسیده، و واریزهایی که در هیچ سیستمی ثبت نشده‌اند، بدون صورت‌حساب بانک قابل ردیابی نیستند."],
        ]
        for line in lines:
            summary.append(line)
        summary["A1"].font = Font(bold=True, size=14)
        for cell in summary["B"][3:12]:
            cell.number_format = "#,##0"
        summary.column_dimensions["A"].width = 42
        summary.column_dimensions["B"].width = 110
        for r in summary.iter_rows(min_row=15):
            for cell in r:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        sheet("به تفکیک هیبرید", ["هیبرید / شعبه", "مشتری بدهکار", "جمع بدهی (تومان)", "واریز بی‌صاحب", "مبلغ بی‌صاحب (تومان)", "مشتری دارای شواهد"],
              [[b["branch"], b["debtors"], toman(b["debt_rial"]), b["unassigned"], toman(b["unassigned_rial"]), b["customers_with_evidence"]]
               for b in report["by_branch"]], [30, 14, 20, 14, 20, 18], money_cols=(3, 5))

        def cand_text(cands: list[dict[str, Any]], i: int, key: str) -> Any:
            if len(cands) <= i:
                return ""
            c = cands[i]
            return {"name": f"{c['name']} (کد {c['code']})", "score": c["score"], "why": "؛ ".join(c["reasons"]),
                    "debt": toman(c["debt_rial"])}[key]

        payment_headers = ["نوع", "شماره سند", "تاریخ", "مبلغ (تومان)", "هیبرید / شعبه", "حساب ثبت‌شده", "بانک", "بابت", "شرح", "ثبت‌کننده"]

        def payment_cols(r: dict[str, Any]) -> list[Any]:
            return [r["kind"], r["number"], r["date_jalali"], toman(r["amount_rial"]), r["branch"], r["account"],
                    r["bank"], r["behalf"], r["description"], r["registered_by"]]

        sheet("واریز بی‌صاحب", payment_headers + ["مشتری پیشنهادی ۱", "امتیاز", "دلیل", "بدهی مشتری (تومان)", "مشتری پیشنهادی ۲", "امتیاز ۲"],
              [payment_cols(r) + [cand_text(r["candidates"], 0, "name"), cand_text(r["candidates"], 0, "score"),
                                  cand_text(r["candidates"], 0, "why"), cand_text(r["candidates"], 0, "debt"),
                                  cand_text(r["candidates"], 1, "name"), cand_text(r["candidates"], 1, "score")]
               for r in report["unassigned"]],
              [12, 12, 12, 16, 22, 26, 26, 30, 30, 18, 32, 9, 46, 16, 30, 9], money_cols=(4, 14))
        sheet("ثبت به نام مشتری دیگر", payment_headers + ["ثبت‌شده به نام", "شماره فاکتورهای شرح", "مشتری صاحب فاکتور", "بدهی صاحب فاکتور (تومان)"],
              [payment_cols(r) + [f"{r['credited_to']['name']} (کد {r['credited_to']['code']})", "، ".join(r["invoice_codes"]),
                                  cand_text(r["candidates"], 0, "name"), cand_text(r["candidates"], 0, "debt")]
               for r in report["misallocated"]],
              [12, 12, 12, 16, 22, 26, 26, 30, 30, 18, 32, 20, 32, 18], money_cols=(4, 14))
        sheet("واریز مستقیم به شرکت", payment_headers + ["مشتری پیشنهادی", "امتیاز", "دلیل", "بدهی مشتری (تومان)"],
              [payment_cols(r) + [cand_text(r["candidates"], 0, "name"), cand_text(r["candidates"], 0, "score"),
                                  cand_text(r["candidates"], 0, "why"), cand_text(r["candidates"], 0, "debt")]
               for r in report["direct"]],
              [12, 12, 12, 16, 22, 26, 26, 30, 30, 18, 32, 9, 46, 16], money_cols=(4, 14))
        sheet("بدهکاران دارای شواهد", ["مشتری", "کد", "هیبرید / شعبه", "ویزیتور آخر", "موبایل", "بدهی فعلی (تومان)", "فاکتور باز",
                                      "قدیمی‌ترین فاکتور باز", "تعداد شواهد", "جمع مبلغ شواهد (تومان)", "بهترین امتیاز", "بهترین شاهد"],
              [[c["name"], c["code"], c["branch"], c["visitor"], c["mobile"], toman(c["debt_rial"]), c["open_invoice_count"],
                c["oldest_open_invoice_jalali"], c["evidence_count"], toman(c["evidence_amount_rial"]), c["best_score"],
                f"{c['evidence'][0]['source']}: {c['evidence'][0]['payment']['kind']} {c['evidence'][0]['payment']['date_jalali']} "
                f"به مبلغ {toman(c['evidence'][0]['payment']['amount_rial']):,} تومان — " + "؛ ".join(c["evidence"][0]["reasons"])]
               for c in report["customers"]],
              [32, 12, 22, 26, 14, 18, 10, 14, 10, 20, 10, 80], money_cols=(6, 10))

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output
