"""مشتریان کارآمد با پرداخت سندنخورده (V174).

مشتری پرداخت کرده و دریافت نقد/حواله‌اش در کارآمد ثبت شده (tblCashD / tblDraftD)، ولی سند
حسابداری‌اش هنوز صادر نشده (VoucherRef خالی). کارآمد برای چنین سندی سطر موقت در
tblVoucherLinesTemp دارد، پس مانده داشبورد (قطعی + موقت) پرداخت را کم کرده، ولی مانده حسابداری
(فقط اسناد قطعی) هنوز مشتری را بدهکار نشان می‌دهد.

برای هر مشتری: بدهی حسابداری (قطعی)، جمع پرداخت‌های سندنخورده و بدهی پس از ثبت سند آن‌ها.
فقط گزارش است؛ سندی ثبت نمی‌شود.
"""

from __future__ import annotations

import re
from collections import defaultdict
from datetime import date, datetime
from io import BytesIO
from typing import Any

from sqlalchemy import text

from app.services.karamad_customer_account_service import (
    LEDGER_READ_ISOLATION,
    _LEDGER_LINES,
    _RECEIVABLE_SL_CODES,
)
from app.utils.jalali import format_jalali_date

TOLERANCE_RIAL = 10_000


def _clean(value: Any) -> str:
    value = str(value or "").replace("ي", "ی").replace("ك", "ک").replace("‌", " ")
    return re.sub(r"\s+", " ", value).strip()


def _day(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    return value if isinstance(value, date) else None


class KaramadUnpostedPaymentService:
    def __init__(self, engine=None):
        self._engine = engine

    @property
    def engine(self):
        if self._engine is None:
            from app.database.karamad_sqlserver import get_karamad_sqlserver_engine
            self._engine = get_karamad_sqlserver_engine()
        return self._engine

    def build(self) -> dict[str, Any]:
        codes = ", ".join(f"N'{c}'" for c in _RECEIVABLE_SL_CODES)
        with self.engine.connect().execution_options(isolation_level=LEDGER_READ_ISOLATION) as c:
            fy = c.execute(text(
                "SELECT TOP (1) [ID], [DateStart] FROM dbo.[tblFiscalYear] "
                "WHERE CAST(GETDATE() AS date) BETWEEN [DateStart] AND [DateEnd] ORDER BY [DateStart] DESC"
            )).mappings().first()
            balances = c.execute(text(f"""
                WITH active_fy AS (SELECT {int(fy['ID'])} AS FiscalYearRef),
                receivable_sl AS (SELECT [ID] FROM dbo.[tblSL] WHERE [Code] IN ({codes})),
                ledger AS ({_LEDGER_LINES}),
                lines AS (
                    SELECT l.[DLRef] AS DL, l.[Debtor], l.[Creditor], l.[IsTemp] FROM ledger l JOIN receivable_sl s ON s.[ID] = l.[SLRef] WHERE l.[DLRef] IS NOT NULL
                    UNION ALL SELECT l.[DL2Ref], l.[Debtor], l.[Creditor], l.[IsTemp] FROM ledger l JOIN receivable_sl s ON s.[ID] = l.[SLRef] WHERE l.[DL2Ref] IS NOT NULL
                    UNION ALL SELECT l.[DL3Ref], l.[Debtor], l.[Creditor], l.[IsTemp] FROM ledger l JOIN receivable_sl s ON s.[ID] = l.[SLRef] WHERE l.[DL3Ref] IS NOT NULL
                )
                SELECT DL, SUM(CASE WHEN IsTemp = 0 THEN Debtor - Creditor ELSE 0 END) AS Posted,
                           SUM(CASE WHEN IsTemp = 1 THEN Debtor - Creditor ELSE 0 END) AS Temp
                FROM lines GROUP BY DL
            """)).mappings().all()
            receipts = c.execute(text("""
                SELECT k.kind, k.inx, k.Code, k.BookDate, k.Amount, k.Behalf, k.Description, k.DLRef,
                       k.Confirmed, b.Name AS Branch, bk.Name AS BankName, u.Name AS UserName,
                       cu.Name AS Customer, cu.Code AS CustomerCode, cu.Mobile
                FROM (
                    SELECT N'حواله بانکی' AS kind, inx, CAST(Code AS nvarchar(50)) AS Code, BookDate, Price AS Amount,
                           Behalf, Description, DLRef, BranchRef, BankRef, UserRef, VoucherRef, Confirmed FROM dbo.[tblDraftD]
                    UNION ALL
                    SELECT N'دریافت نقدی', inx, CAST(Code AS nvarchar(50)), BookDate, PriceN, Behalf, Description,
                           DLRef, BranchRef, NULL, UserRef, VoucherRef, Confirmed FROM dbo.[tblCashD] WHERE ISNULL(PriceN, 0) <> 0
                ) k
                JOIN dbo.[tblCustomer] cu ON cu.[DLRef] = k.DLRef
                LEFT JOIN dbo.[tblUser] u ON u.[ID] = k.UserRef
                OUTER APPLY (SELECT TOP (1) x.[Name] FROM dbo.[tblBranch] x WHERE x.[Code] = k.BranchRef) b
                OUTER APPLY (SELECT TOP (1) a.[Name] FROM dbo.[tblAccBank] a WHERE a.[Code] = k.BankRef) bk
                WHERE k.VoucherRef IS NULL AND k.BookDate >= :since
            """), {"since": fy["DateStart"]}).mappings().all()
            visitors = {int(r[0]): r[1] for r in c.execute(text("""
                SELECT cu.DLRef, v.Name FROM dbo.[tblCustomer] cu
                CROSS APPLY (SELECT TOP (1) f.VisitorRef FROM dbo.[tblFactorF] f WHERE f.CustomerRef = cu.ID ORDER BY f.DateE DESC) last
                JOIN dbo.[tblVisitor] v ON v.ID = last.VisitorRef WHERE cu.DLRef IS NOT NULL
            """))}

        posted = {int(r["DL"]): float(r["Posted"] or 0) for r in balances}
        temp = {int(r["DL"]): float(r["Temp"] or 0) for r in balances}
        today = date.today()
        payments = []
        for r in receipts:
            day = _day(r["BookDate"])
            payments.append({
                "dl_ref": int(r["DLRef"]), "customer": _clean(r["Customer"]), "customer_code": r["CustomerCode"],
                "mobile": r["Mobile"], "branch": _clean(r["Branch"]), "kind": r["kind"], "number": r["Code"],
                "date_jalali": format_jalali_date(day), "days_waiting": (today - day).days if day else None,
                "amount_rial": float(r["Amount"] or 0), "bank": _clean(r["BankName"]), "behalf": _clean(r["Behalf"]),
                "description": _clean(r["Description"]), "registered_by": r["UserName"],
                "confirmed": bool(r["Confirmed"]),
            })

        by_customer: dict[int, list[dict[str, Any]]] = defaultdict(list)
        for p in payments:
            by_customer[p["dl_ref"]].append(p)
        customers = []
        for dl, items in by_customer.items():
            unposted = sum(p["amount_rial"] for p in items)
            posted_debt = posted.get(dl, 0.0)
            customers.append({
                "dl_ref": dl, "customer": items[0]["customer"], "customer_code": items[0]["customer_code"],
                "mobile": items[0]["mobile"], "branch": items[0]["branch"], "visitor": visitors.get(dl),
                "posted_debt_rial": posted_debt, "dashboard_balance_rial": posted_debt + temp.get(dl, 0.0),
                "unposted_count": len(items), "unposted_rial": unposted,
                "debt_after_posting_rial": posted_debt - unposted,
                "oldest_unposted_jalali": min((p["date_jalali"] for p in items if p["date_jalali"]), default=None),
                "max_days_waiting": max((p["days_waiting"] or 0) for p in items),
                "debtor_in_accounting": posted_debt > TOLERANCE_RIAL,
            })
        customers.sort(key=lambda x: (-x["unposted_rial"]))
        debtors = [c for c in customers if c["debtor_in_accounting"]]

        branches: dict[str, dict[str, Any]] = defaultdict(lambda: {"payments": 0, "rial": 0.0, "customers": set(),
                                                                   "debtors": set(), "oldest_days": 0})
        for p in payments:
            b = branches[p["branch"] or "—"]
            b["payments"] += 1
            b["rial"] += p["amount_rial"]
            b["customers"].add(p["dl_ref"])
            if posted.get(p["dl_ref"], 0) > TOLERANCE_RIAL:
                b["debtors"].add(p["dl_ref"])
            b["oldest_days"] = max(b["oldest_days"], p["days_waiting"] or 0)
        by_branch = sorted(({"branch": k, "payments": v["payments"], "rial": v["rial"], "customers": len(v["customers"]),
                             "debtors": len(v["debtors"]), "oldest_days": v["oldest_days"]} for k, v in branches.items()),
                           key=lambda x: -x["rial"])
        return {
            "status": "success",
            "fiscal_year_start_jalali": format_jalali_date(_day(fy["DateStart"])),
            "summary": {
                "payments": len(payments), "rial": sum(p["amount_rial"] for p in payments),
                "customers": len(customers), "debtors_in_accounting": len(debtors),
                "debtors_unposted_rial": sum(c["unposted_rial"] for c in debtors),
                "debtors_posted_debt_rial": sum(c["posted_debt_rial"] for c in debtors),
                "debtors_cleared_by_posting": sum(1 for c in debtors if c["debt_after_posting_rial"] <= TOLERANCE_RIAL),
            },
            "by_branch": by_branch,
            "customers": customers,
            "payments": sorted(payments, key=lambda p: -(p["days_waiting"] or 0)),
        }

    def workbook(self) -> BytesIO:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        report = self.build()
        toman = lambda rial: round(float(rial or 0) / 10)  # noqa: E731
        wb = Workbook()
        head_fill, head_font = PatternFill("solid", fgColor="1F4E78"), Font(color="FFFFFF", bold=True)

        def sheet(title, headers, rows, widths, money=()):
            ws = wb.create_sheet(title)
            ws.sheet_view.rightToLeft = True
            ws.append(headers)
            for cell in ws[1]:
                cell.fill, cell.font = head_fill, head_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for row in rows:
                ws.append(row)
            for col in money:
                for cell in ws[get_column_letter(col)][1:]:
                    cell.number_format = "#,##0"
            for i, w in enumerate(widths, start=1):
                ws.column_dimensions[get_column_letter(i)].width = w
            for r in ws.iter_rows(min_row=2):
                for cell in r:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            ws.freeze_panes = "A2"
            if rows:
                ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

        s = report["summary"]
        ws = wb.active
        ws.title = "خلاصه و راهنما"
        ws.sheet_view.rightToLeft = True
        for line in [
            ["مشتریان کارآمد با پرداخت سندنخورده", ""],
            [f"تاریخ گزارش {format_jalali_date(date.today())} — از ابتدای سال مالی {report['fiscal_year_start_jalali']}", ""],
            ["", ""],
            ["دریافت سندنخورده از مشتریان (تعداد)", s["payments"]],
            ["دریافت سندنخورده از مشتریان (تومان)", toman(s["rial"])],
            ["مشتری دارای دریافت سندنخورده", s["customers"]],
            ["…که در حسابداری (اسناد قطعی) هنوز بدهکارند", s["debtors_in_accounting"]],
            ["جمع بدهی حسابداری همین بدهکاران (تومان)", toman(s["debtors_posted_debt_rial"])],
            ["جمع پرداخت سندنخورده همین بدهکاران (تومان)", toman(s["debtors_unposted_rial"])],
            ["بدهکارانی که با صدور سند، بدهی‌شان صفر یا بستانکار می‌شود", s["debtors_cleared_by_posting"]],
            ["", ""], ["راهنما", ""],
            ["دریافت سندنخورده", "دریافت نقد یا حواله مشتری که در کارآمد ثبت شده ولی سند حسابداری‌اش هنوز صادر نشده است."],
            ["بدهی حسابداری", "مانده حساب مشتری فقط از اسناد قطعی (بدون سطرهای موقت)؛ همان چیزی که گزارش حسابداری نشان می‌دهد."],
            ["بدهی پس از ثبت سند", "بدهی حسابداری منهای دریافت‌های سندنخورده؛ اگر سند این دریافت‌ها صادر شود، بدهی به این عدد می‌رسد."],
            ["مانده داشبورد", "مانده با سطرهای موقت (اسناد صادرنشده)؛ معمولاً پرداخت‌های سندنخورده را از قبل کم کرده است."],
        ]:
            ws.append(line)
        ws["A1"].font = Font(bold=True, size=14)
        for cell in ws["B"][3:10]:
            cell.number_format = "#,##0"
        ws.column_dimensions["A"].width = 55
        ws.column_dimensions["B"].width = 100
        for r in ws.iter_rows(min_row=12):
            for cell in r:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        sheet("به تفکیک هیبرید", ["هیبرید / شعبه", "دریافت سندنخورده", "مبلغ (تومان)", "مشتری", "بدهکار در حسابداری",
                                  "قدیمی‌ترین (روز)"],
              [[b["branch"], b["payments"], toman(b["rial"]), b["customers"], b["debtors"], b["oldest_days"]]
               for b in report["by_branch"]], [28, 16, 20, 10, 18, 14], money=(3,))
        sheet("بدهکاران با پرداخت سندنخورده",
              ["مشتری", "کد", "هیبرید / شعبه", "ویزیتور آخر", "موبایل", "بدهی حسابداری (تومان)", "پرداخت سندنخورده (تومان)",
               "تعداد", "بدهی پس از ثبت سند (تومان)", "مانده داشبورد (تومان)", "قدیمی‌ترین پرداخت سندنخورده", "روز انتظار"],
              [[c["customer"], c["customer_code"], c["branch"], c["visitor"], c["mobile"], toman(c["posted_debt_rial"]),
                toman(c["unposted_rial"]), c["unposted_count"], toman(c["debt_after_posting_rial"]),
                toman(c["dashboard_balance_rial"]), c["oldest_unposted_jalali"], c["max_days_waiting"]]
               for c in report["customers"] if c["debtor_in_accounting"]],
              [34, 12, 22, 26, 14, 20, 20, 8, 22, 20, 16, 10], money=(6, 7, 9, 10))
        sheet("ریز دریافت‌های سندنخورده",
              ["مشتری", "کد", "هیبرید / شعبه", "نوع", "شماره", "تاریخ", "روز انتظار", "مبلغ (تومان)", "بانک", "بابت", "شرح",
               "ثبت‌کننده", "تأیید شده"],
              [[p["customer"], p["customer_code"], p["branch"], p["kind"], p["number"], p["date_jalali"], p["days_waiting"],
                toman(p["amount_rial"]), p["bank"], p["behalf"], p["description"], p["registered_by"],
                "بله" if p["confirmed"] else "خیر"] for p in report["payments"]],
              [32, 12, 22, 14, 10, 12, 10, 18, 26, 36, 30, 22, 10], money=(8,))
        out = BytesIO()
        wb.save(out)
        out.seek(0)
        return out
