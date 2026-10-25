"""هشدار اسناد انبار قیمت‌گذاری‌نشده — راهکاران (فقط خواندنی).

مبنا همان ستون «وضعیت قیمت گذاری» فهرست اسناد انبار راهکاران است:
LGS3.vwInventoryVoucherWithAllStates.PricingState (1 = قیمت گذاری نشده، 2 = قیمت گذاری شده).
اسناد باطل‌شده (State = 3) حساب نمی‌شوند. جهت سند از نوع سند (Direction: 1 رسید/ورودی، 2 حواله/خروجی).
انبارهای کارخانه = سایت‌های کادوس، اکسیر کادوس و فراز بهداشت (Plant 1، 2، 3).
"""

from __future__ import annotations

import threading
import time
from datetime import date, datetime
from io import BytesIO
from typing import Any

import jdatetime
from sqlalchemy import text

from app.utils.jalali import format_jalali_date

FACTORY_PLANTS = (1, 2, 3)
UNPRICED, PRICED, VOIDED = 1, 2, 3
CACHE_SECONDS = 300
JALALI_MONTHS = ("فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور", "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند")
VOUCHER_STATES = {1: "ثبت شده", 2: "تایید شده", 3: "باطل شده"}

# انبارهای راکد و ضایعات در این هشدار حساب نمی‌شوند (درخواست واحد مالی).
EXCLUDED_STORES_SQL = "s.Name NOT LIKE N'%راکد%' AND s.Name NOT LIKE N'%ضایعات%'"

_cache: dict[Any, tuple[float, Any]] = {}
_lock = threading.Lock()

# Fixed internal SQL; only :from_date / :store_id / :direction are parameters.
_BASE = """
    FROM LGS3.vwInventoryVoucherWithAllStates v
    JOIN LGS3.InventoryVoucherSpecification sp ON sp.InventoryVoucherSpecificationID = v.InventoryVoucherSpecificationRef
    JOIN LGS3.Store s ON s.StoreID = v.StoreRef
    LEFT JOIN LGS3.Plant pl ON pl.PlantID = s.PlantRef
    WHERE v.Date >= :from_date AND v.State <> 3
      AND s.Name NOT LIKE N'%راکد%' AND s.Name NOT LIKE N'%ضایعات%'
"""


def year_start() -> date:
    """اول سال شمسی جاری."""
    return jdatetime.date(jdatetime.date.today().year, 1, 1).togregorian()


def _as_date(value: Any) -> date | None:
    if value is None:
        return None
    return value.date() if isinstance(value, datetime) else value


class InventoryPricingService:
    def __init__(self, engine=None):
        self._engine = engine

    @property
    def engine(self):
        if self._engine is None:
            from app.database.sqlserver import get_sqlserver_engine
            self._engine = get_sqlserver_engine()
        return self._engine

    def _cached(self, key, build):
        with _lock:
            hit = _cache.get(key)
            if hit and time.time() - hit[0] < CACHE_SECONDS:
                return hit[1]
        value = build()
        with _lock:
            _cache[key] = (time.time(), value)
        return value

    # ------------------------------------------------------------------ summary
    def summary(self, from_date: date | None = None) -> dict[str, Any]:
        start = from_date or year_start()
        return self._cached(("summary", start), lambda: self._summary(start))

    def _summary(self, start: date) -> dict[str, Any]:
        today = date.today()
        with self.engine.connect() as c:
            stores = c.execute(text(f"""
                SELECT s.StoreID, s.Code, s.Name, s.PlantRef, pl.Name AS PlantName,
                       COUNT(*) AS docs,
                       SUM(CASE WHEN v.PricingState = {UNPRICED} THEN 1 ELSE 0 END) AS unpriced,
                       SUM(CASE WHEN v.PricingState = {UNPRICED} AND sp.Direction = 1 THEN 1 ELSE 0 END) AS unpriced_in,
                       SUM(CASE WHEN v.PricingState = {UNPRICED} AND sp.Direction = 2 THEN 1 ELSE 0 END) AS unpriced_out,
                       MIN(CASE WHEN v.PricingState = {UNPRICED} THEN v.Date END) AS oldest
                {_BASE}
                GROUP BY s.StoreID, s.Code, s.Name, s.PlantRef, pl.Name
            """), {"from_date": start}).mappings().all()
            monthly_rows = c.execute(text(f"""
                SELECT v.Date, v.PricingState, COUNT(*) AS n {_BASE} GROUP BY v.Date, v.PricingState
            """), {"from_date": start}).mappings().all()
            types = c.execute(text(f"""
                SELECT sp.Name, sp.Direction, COUNT(*) AS unpriced {_BASE} AND v.PricingState = {UNPRICED}
                GROUP BY sp.Name, sp.Direction ORDER BY unpriced DESC
            """), {"from_date": start}).mappings().all()

        store_list = []
        for r in stores:
            oldest = _as_date(r["oldest"])
            store_list.append({
                "store_id": int(r["StoreID"]),
                "store_code": str(r["Code"] or "").strip(),
                "store_name": str(r["Name"] or "").strip(),
                "plant_name": str(r["PlantName"] or "").strip(),
                "is_factory": r["PlantRef"] in FACTORY_PLANTS,
                "document_count": int(r["docs"] or 0),
                "unpriced_count": int(r["unpriced"] or 0),
                "unpriced_receipt_count": int(r["unpriced_in"] or 0),
                "unpriced_issue_count": int(r["unpriced_out"] or 0),
                "priced_count": int(r["docs"] or 0) - int(r["unpriced"] or 0),
                "oldest_unpriced_date_jalali": format_jalali_date(oldest) or "",
                "oldest_unpriced_days": (today - oldest).days if oldest else None,
            })
        store_list.sort(key=lambda x: (not x["is_factory"], x["plant_name"], -x["unpriced_count"]))

        months: dict[tuple[int, int], dict[str, int]] = {}
        for r in monthly_rows:
            j = jdatetime.date.fromgregorian(date=_as_date(r["Date"]))
            m = months.setdefault((j.year, j.month), {"documents": 0, "unpriced": 0})
            m["documents"] += int(r["n"] or 0)
            if r["PricingState"] == UNPRICED:
                m["unpriced"] += int(r["n"] or 0)
        monthly = [{
            "year": y, "month": mo, "label": f"{JALALI_MONTHS[mo - 1]} {y}",
            "document_count": v["documents"], "unpriced_count": v["unpriced"],
            "priced_percent": round(100 * (v["documents"] - v["unpriced"]) / v["documents"], 1) if v["documents"] else 0.0,
        } for (y, mo), v in sorted(months.items())]

        def total(rows):
            return {
                "document_count": sum(x["document_count"] for x in rows),
                "unpriced_count": sum(x["unpriced_count"] for x in rows),
                "unpriced_receipt_count": sum(x["unpriced_receipt_count"] for x in rows),
                "unpriced_issue_count": sum(x["unpriced_issue_count"] for x in rows),
            }

        return {
            "status": "success",
            "from_date_jalali": format_jalali_date(start),
            "generated_at_jalali": format_jalali_date(today),
            "total": total(store_list),
            "factory_total": total([x for x in store_list if x["is_factory"]]),
            "stores": store_list,
            "monthly": monthly,
            "types": [{"title": str(t["Name"] or "").strip(), "direction": "receipt" if t["Direction"] == 1 else "issue",
                       "unpriced_count": int(t["unpriced"] or 0)} for t in types],
            "note": "مبنا ستون «وضعیت قیمت گذاری» فهرست اسناد انبار راهکاران است؛ اسناد باطل‌شده و انبارهای راکد و ضایعات حساب نمی‌شوند.",
        }

    # ---------------------------------------------------------------- documents
    def documents(self, from_date: date | None = None, store_id: int | None = None,
                  direction: str | None = None, factory_only: bool = False,
                  pricing: str = "unpriced") -> list[dict[str, Any]]:
        start = from_date or year_start()
        key = ("docs", start, store_id, direction, factory_only, pricing)
        return self._cached(key, lambda: self._documents(start, store_id, direction, factory_only, pricing))

    def _documents(self, start: date, store_id: int | None, direction: str | None, factory_only: bool,
                   pricing: str = "unpriced") -> list[dict[str, Any]]:
        where, params = "", {"from_date": start}
        if pricing in ("unpriced", "priced"):
            where += f" AND v.PricingState = {UNPRICED if pricing == 'unpriced' else PRICED}"
        if store_id is not None:
            where += " AND v.StoreRef = :store_id"
            params["store_id"] = int(store_id)
        if direction in ("receipt", "issue"):
            where += " AND sp.Direction = :direction"
            params["direction"] = 1 if direction == "receipt" else 2
        if factory_only:
            where += f" AND s.PlantRef IN ({','.join(str(p) for p in FACTORY_PLANTS)})"
        today = date.today()
        with self.engine.connect() as c:
            rows = c.execute(text(f"""
                SELECT v.InventoryVoucherID, v.Number, v.Date, v.State, v.PricingState, sp.Name AS Title, sp.Direction,
                       s.Name AS StoreName, s.Code AS StoreCode, pl.Name AS PlantName,
                       cs.Name AS CounterpartStore, v.CounterpartEntityText, v.CounterpartDLCode, v.Description,
                       u.Name AS CreatorName,
                       (SELECT COUNT(*) FROM LGS3.InventoryVoucherItem i WHERE i.InventoryVoucherRef = v.InventoryVoucherID) AS ItemCount
                FROM LGS3.vwInventoryVoucherWithAllStates v
                JOIN LGS3.InventoryVoucherSpecification sp ON sp.InventoryVoucherSpecificationID = v.InventoryVoucherSpecificationRef
                JOIN LGS3.Store s ON s.StoreID = v.StoreRef
                LEFT JOIN LGS3.Plant pl ON pl.PlantID = s.PlantRef
                LEFT JOIN LGS3.Store cs ON cs.StoreID = v.CounterpartStoreRef
                LEFT JOIN SYS3.[User] u ON u.UserID = v.Creator
                WHERE v.Date >= :from_date AND v.State <> {VOIDED} AND {EXCLUDED_STORES_SQL} {where}
                ORDER BY v.Date, v.Number
            """), params).mappings().all()
        result = []
        for r in rows:
            d = _as_date(r["Date"])
            result.append({
                "voucher_id": int(r["InventoryVoucherID"]),
                "number": str(r["Number"] or ""),
                "title": str(r["Title"] or "").strip(),
                "direction": "receipt" if r["Direction"] == 1 else "issue",
                "store_name": str(r["StoreName"] or "").strip(),
                "store_code": str(r["StoreCode"] or "").strip(),
                "plant_name": str(r["PlantName"] or "").strip(),
                "date_jalali": format_jalali_date(d) or "",
                "days_waiting": (today - d).days if d else None,
                "counterpart": str(r["CounterpartEntityText"] or "").strip(),
                "counterpart_code": str(r["CounterpartDLCode"] or "").strip(),
                "counterpart_store": str(r["CounterpartStore"] or "").strip(),
                "state": VOUCHER_STATES.get(r["State"], str(r["State"])),
                "priced": r["PricingState"] == PRICED,
                "description": str(r["Description"] or "").strip(),
                "creator": str(r["CreatorName"] or "").strip(),
                "item_count": int(r["ItemCount"] or 0),
            })
        return result

    # ------------------------------------------------------------------ excel
    def workbook(self, from_date: date | None = None) -> BytesIO:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        summary = self.summary(from_date)
        docs = self.documents(from_date)
        head = PatternFill("solid", fgColor="1F4E78")
        wb = Workbook()

        def sheet(ws, headers, rows, widths):
            ws.sheet_view.rightToLeft = True
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = head
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for row in rows:
                ws.append(row)
            for i, w in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(i)].width = w
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

        ws = wb.active
        ws.title = "خلاصه انبارها"
        sheet(ws, ["کارخانه / سایت", "کد انبار", "انبار", "کل اسناد", "قیمت گذاری نشده", "رسید (ورودی)", "حواله (خروجی)",
                   "قیمت گذاری شده", "قدیمی‌ترین سند بی‌قیمت", "روز انتظار"],
              [[x["plant_name"], x["store_code"], x["store_name"], x["document_count"], x["unpriced_count"],
                x["unpriced_receipt_count"], x["unpriced_issue_count"], x["priced_count"],
                x["oldest_unpriced_date_jalali"], x["oldest_unpriced_days"]] for x in summary["stores"]],
              [18, 10, 34, 11, 15, 13, 13, 14, 18, 11])
        sheet(wb.create_sheet("ماه به ماه"), ["ماه", "کل اسناد", "قیمت گذاری نشده", "درصد قیمت‌گذاری‌شده"],
              [[m["label"], m["document_count"], m["unpriced_count"], m["priced_percent"]] for m in summary["monthly"]],
              [18, 12, 16, 18])
        sheet(wb.create_sheet("ریز اسناد قیمت گذاری نشده"),
              ["شماره سند", "عنوان سند", "رسید / حواله", "کارخانه / سایت", "انبار", "تاریخ سند", "روز انتظار", "طرف مقابل",
               "کد تفصیلی طرف مقابل", "انبار مقابل", "وضعیت", "تعداد ردیف", "صادر کننده", "توضیحات"],
              [[d["number"], d["title"], "رسید" if d["direction"] == "receipt" else "حواله", d["plant_name"], d["store_name"],
                d["date_jalali"], d["days_waiting"], d["counterpart"], d["counterpart_code"], d["counterpart_store"],
                d["state"], d["item_count"], d["creator"], d["description"]] for d in docs],
              [11, 26, 10, 16, 30, 12, 10, 30, 14, 26, 11, 10, 22, 50])
        info = wb.create_sheet("توضیحات")
        info.sheet_view.rightToLeft = True
        for line in [f"اسناد انبار قیمت‌گذاری‌نشده راهکاران از {summary['from_date_jalali']} تا {summary['generated_at_jalali']}",
                     summary["note"],
                     "انبارهای کارخانه: سایت‌های کادوس، اکسیر کادوس و فراز بهداشت."]:
            info.append([line])
        info.column_dimensions["A"].width = 110
        out = BytesIO()
        wb.save(out)
        out.seek(0)
        return out
