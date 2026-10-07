from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.utils.jalali import format_jalali_date, format_jalali_datetime


_STATUS_LABELS = {
    "posted": "سندخورده",
    "reversed": "برگشت یا خنثی‌شده",
    "internal_transfer": "انتقال داخلی شرکت",
    "needs_review": "نیازمند بررسی",
    "unposted": "بدون سند احتمالی",
}

_STATUS_COLORS = {
    "posted": "D9EAD3",
    "reversed": "D9EAF7",
    "internal_transfer": "DDEBF7",
    "needs_review": "FFF2CC",
    "unposted": "F4CCCC",
}

_BANK_ROW_HEADERS = [
    "ردیف فایل بانک",
    "تاریخ شمسی بانک",
    "جهت",
    "مبلغ بانک (ریال)",
    "شرح بانک",
    "شماره پیگیری",
    "شماره سند بانک",
    "شناسه واریز",
    "وضعیت",
    "روش تطبیق",
    "تعداد ردیف گروه",
    "جمع گروه بانک (ریال)",
    "درجه اطمینان",
    "دلیل نتیجه",
    "اختلاف مبلغ",
    "اختلاف تاریخ (روز)",
    "شماره سند راهکاران",
    "تاریخ شمسی سند راهکاران",
    "مبلغ سند راهکاران",
    "نوع سند راهکاران",
    "شناسه ردیف راهکاران",
    "شرح سند راهکاران",
    "نوع کارمزد",
    "ردیف گردش اصلی کارمزد",
    "کارمزد این گردش (ریال)",
    "مانده بانک بعد از ردیف",
    "شماره سند حسابداری",
    "مانده دفتر راهکاران تا این سند",
    "مبنای مانده دفتر",
    "مانده پایان روز بانک",
    "مانده پایان روز دفتر",
    "اختلاف مانده روز",
]

_AMOUNT_COLUMNS = ("D", "L", "O", "S", "Y", "Z", "AB", "AD", "AE", "AF")

_ERP_ONLY_HEADERS = [
    "نوع سند",
    "شماره سند راهکاران",
    "تاریخ شمسی سند",
    "مبلغ (ریال)",
    "شناسه ردیف",
    "شماره ردیف سند",
    "شرح سند",
    "حساب مقابل",
    "شناسه تراکنش",
]

_THIN_GRAY = Side(style="thin", color="D9E1F2")
_TABLE_BORDER = Border(
    left=_THIN_GRAY,
    right=_THIN_GRAY,
    top=_THIN_GRAY,
    bottom=_THIN_GRAY,
)
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FILL = PatternFill("solid", fgColor="D9EAF7")


def _set_sheet_defaults(sheet) -> None:
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "A2"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.outlinePr.summaryBelow = False


def _style_header(sheet, row: int, column_count: int) -> None:
    for cell in sheet[row][:column_count]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = _TABLE_BORDER
    sheet.row_dimensions[row].height = 32


def _apply_table_style(sheet, start_row: int, end_row: int, end_column: int) -> None:
    for row in sheet.iter_rows(
        min_row=start_row,
        max_row=end_row,
        min_col=1,
        max_col=end_column,
    ):
        for cell in row:
            cell.border = _TABLE_BORDER
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )


def _set_widths(sheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _direction_label(direction: str | None) -> str:
    return {
        "deposit": "واریز",
        "withdrawal": "برداشت",
    }.get(direction or "", direction or "")


def _transaction_type_label(transaction_type: str | None) -> str:
    return {
        "receipt": "دریافت",
        "payment": "پرداخت",
    }.get(transaction_type or "", transaction_type or "")


def _matched_transaction(row: dict[str, Any]) -> dict[str, Any]:
    transaction = row.get("matched_transaction")
    return transaction if isinstance(transaction, dict) else {}


def _bank_row_values(row: dict[str, Any]) -> list[Any]:
    transaction = _matched_transaction(row)
    reasons = row.get("reasons") or []
    return [
        row.get("row_number"),
        row.get("bank_date_jalali") or format_jalali_date(row.get("bank_date")),
        _direction_label(row.get("direction")),
        row.get("amount"),
        row.get("description", ""),
        row.get("reference", ""),
        row.get("bank_document_number", ""),
        row.get("deposit_id", ""),
        row.get("business_status_fa")
        or _STATUS_LABELS.get(row.get("business_status"), ""),
        {
            "daily_fee_group": "تجمیع کارمزد روزانه",
            "one_to_one": "تطبیق یک‌به‌یک",
            "bank_reversal_pair": "جفت برداشت و برگشت",
            "bank_reversal_batch": "گروه برداشت و برگشت",
            "same_day_amount_batch": "تطبیق گروهی همان روز",
            "tolerance_amount_batch": "تطبیق گروهی بازه تاریخ",
            "internal_transfer_policy": "انتقال بین حساب‌های شرکت",
        }.get(row.get("match_method"), ""),
        row.get("match_group_row_count"),
        row.get("match_group_bank_total"),
        row.get("confidence"),
        " | ".join(str(reason) for reason in reasons),
        row.get("amount_difference"),
        row.get("date_difference_days"),
        transaction.get("document_number", ""),
        transaction.get("booking_date_jalali")
        or format_jalali_date(transaction.get("booking_date"))
        or "",
        transaction.get("amount"),
        _transaction_type_label(transaction.get("transaction_type")),
        transaction.get("item_id", ""),
        transaction.get("description", ""),
        row.get("fee_kind") or "",
        row.get("fee_parent_row_number"),
        row.get("fee_amount") or None,
        row.get("balance"),
        row.get("book_voucher_number") or "",
        row.get("book_balance_after_document"),
        row.get("book_balance_basis") or "",
        row.get("bank_day_end_balance"),
        row.get("book_day_end_balance"),
        row.get("day_balance_difference"),
    ]


def _add_bank_status_sheet(
    workbook: Workbook,
    report: dict[str, Any],
    status: str,
) -> None:
    title = _STATUS_LABELS[status]
    sheet = workbook.create_sheet(title)
    _set_sheet_defaults(sheet)
    sheet.sheet_properties.tabColor = _STATUS_COLORS[status]
    sheet.append(_BANK_ROW_HEADERS)
    _style_header(sheet, 1, len(_BANK_ROW_HEADERS))

    rows = [
        row
        for row in report.get("rows", [])
        if row.get("business_status") == status
    ]
    for row in rows:
        sheet.append(_bank_row_values(row))

    if rows:
        _apply_table_style(sheet, 2, len(rows) + 1, len(_BANK_ROW_HEADERS))
        last_column = get_column_letter(len(_BANK_ROW_HEADERS))
        sheet.auto_filter.ref = f"A1:{last_column}{len(rows) + 1}"
        for column in _AMOUNT_COLUMNS:
            for cell in sheet[column][1:]:
                cell.number_format = "#,##0"
        for column in ("K", "M"):
            for cell in sheet[column][1:]:
                cell.number_format = "0"

    _set_widths(
        sheet,
        [10, 14, 11, 18, 46, 20, 22, 18, 18, 22, 15, 21, 13, 42, 17, 17, 20, 18, 19, 17, 18, 48,
         22, 14, 16, 20, 14, 22, 18, 20, 20, 18],
    )


def _add_erp_only_sheet(workbook: Workbook, report: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("اسناد تطبیق‌قطعی‌نشده")
    _set_sheet_defaults(sheet)
    sheet.sheet_properties.tabColor = "B4C6E7"
    sheet.append(_ERP_ONLY_HEADERS)
    _style_header(sheet, 1, len(_ERP_ONLY_HEADERS))

    transactions = report.get("unmatched_erp_transactions") or []
    for transaction in transactions:
        sheet.append(
            [
                _transaction_type_label(transaction.get("transaction_type")),
                transaction.get("document_number", ""),
                transaction.get("booking_date_jalali")
                or format_jalali_date(transaction.get("booking_date"))
                or "",
                transaction.get("amount"),
                transaction.get("item_id", ""),
                transaction.get("item_number", ""),
                transaction.get("description", ""),
                transaction.get("counterpart_account_name", ""),
                transaction.get("transaction_id", ""),
            ]
        )

    if transactions:
        _apply_table_style(sheet, 2, len(transactions) + 1, len(_ERP_ONLY_HEADERS))
        sheet.auto_filter.ref = f"A1:I{len(transactions) + 1}"
        for cell in sheet["D"][1:]:
            cell.number_format = "#,##0"
    _set_widths(sheet, [13, 20, 18, 20, 18, 18, 52, 32, 24])


def _add_daily_balance_sheet(workbook: Workbook, report: dict[str, Any]) -> None:
    book_balance = report.get("book_balance") or {}
    sheet = workbook.create_sheet("مانده روزانه")
    _set_sheet_defaults(sheet)
    sheet.sheet_properties.tabColor = "C6E0B4"
    if not book_balance.get("available"):
        sheet.append([book_balance.get("message") or "مانده دفتر راهکاران در دسترس نیست."])
        _set_widths(sheet, [80])
        return

    headers = [
        "تاریخ شمسی",
        "مانده پایان روز بانک",
        "مانده پایان روز دفتر راهکاران",
        "اختلاف (بانک − دفتر)",
        "وضعیت",
    ]
    sheet.append(headers)
    _style_header(sheet, 1, len(headers))
    daily = book_balance.get("daily") or []
    for item in daily:
        sheet.append(
            [
                item.get("date_jalali"),
                item.get("bank_balance"),
                item.get("book_balance"),
                item.get("difference"),
                "برابر" if item.get("matched") else "مغایر",
            ]
        )
    if daily:
        _apply_table_style(sheet, 2, len(daily) + 1, len(headers))
        for column in ("B", "C", "D"):
            for cell in sheet[column][1:]:
                cell.number_format = "#,##0"
        for row_index, item in enumerate(daily, start=2):
            sheet.cell(row_index, 5).fill = PatternFill(
                "solid",
                fgColor="D9EAD3" if item.get("matched") else "F4CCCC",
            )
    note_row = len(daily) + 3
    sheet.cell(
        note_row,
        1,
        f"تفصیلی دفتر: {book_balance.get('dl_code', '')} — {book_balance.get('dl_title', '')}",
    )
    if book_balance.get("first_unmatched_date_jalali"):
        sheet.cell(
            note_row + 1,
            1,
            "مانده بانک و دفتر تا "
            f"{book_balance.get('last_matched_date_jalali') or '-'} برابر است؛ "
            f"اولین روز مغایر: {book_balance['first_unmatched_date_jalali']}",
        )
    _set_widths(sheet, [16, 24, 28, 22, 12])


def _add_fee_sheet(workbook: Workbook, report: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("کارمزدها")
    _set_sheet_defaults(sheet)
    sheet.sheet_properties.tabColor = "F8CBAD"
    headers = [
        "ردیف فایل بانک",
        "تاریخ شمسی",
        "نوع کارمزد",
        "مبلغ کارمزد (ریال)",
        "ردیف گردش اصلی",
        "مبلغ گردش اصلی",
        "درصد کارمزد",
        "شرح گردش اصلی",
        "وضعیت",
        "سند راهکاران",
        "توضیح",
    ]
    sheet.append(headers)
    _style_header(sheet, 1, len(headers))
    fees = [row for row in report.get("rows", []) if row.get("is_fee")]
    for row in fees:
        transaction = _matched_transaction(row)
        sheet.append(
            [
                row.get("row_number"),
                row.get("bank_date_jalali") or format_jalali_date(row.get("bank_date")),
                row.get("fee_kind"),
                row.get("amount"),
                row.get("fee_parent_row_number"),
                row.get("fee_parent_amount"),
                row.get("fee_rate_percent"),
                row.get("fee_parent_description", ""),
                row.get("business_status_fa")
                or _STATUS_LABELS.get(row.get("business_status"), ""),
                transaction.get("document_number", ""),
                row.get("fee_note", ""),
            ]
        )
    if fees:
        _apply_table_style(sheet, 2, len(fees) + 1, len(headers))
        sheet.auto_filter.ref = f"A1:K{len(fees) + 1}"
        for column in ("D", "F"):
            for cell in sheet[column][1:]:
                cell.number_format = "#,##0"
        for cell in sheet["G"][1:]:
            cell.number_format = "0.0000"

    summary = (report.get("summary") or {}).get("fee_summary") or {}
    start_row = len(fees) + 3
    sheet.cell(start_row, 1, "جمع به تفکیک نوع کارمزد")
    sheet.cell(start_row, 1).font = Font(bold=True)
    for offset, item in enumerate(summary.get("by_kind") or [], start=1):
        sheet.cell(start_row + offset, 1, item.get("kind"))
        sheet.cell(start_row + offset, 2, item.get("count"))
        amount_cell = sheet.cell(start_row + offset, 3, item.get("amount"))
        amount_cell.number_format = "#,##0"
    _set_widths(sheet, [22, 14, 26, 18, 14, 20, 12, 48, 16, 20, 44])


def _add_summary_sheet(workbook: Workbook, report: dict[str, Any]) -> None:
    sheet = workbook.active
    sheet.title = "خلاصه"
    _set_sheet_defaults(sheet)
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.tabColor = "5B9BD5"
    file_info = report.get("file", {})
    bank_account = report.get("bank_account") or report.get("account", {})
    account_number = str(
        file_info.get("statement_account_number")
        or bank_account.get("number")
        or ""
    )
    account_digits = "".join(character for character in account_number if character.isdigit())
    account_last_four = str(
        bank_account.get("account_last_four")
        or (account_digits[-4:] if account_digits else "")
    )
    bank_name = str(bank_account.get("bank_name") or "")
    account_display_label = str(
        bank_account.get("display_label")
        or " ".join(
            value
            for value in (
                f"حساب {account_last_four}" if account_last_four else "حساب بانکی",
                bank_name,
            )
            if value
        )
    )
    sheet.merge_cells("A1:E1")
    sheet["A1"] = f"گزارش مغایرت‌گیری {account_display_label}"
    sheet["A1"].font = Font(bold=True, size=16, color="1F1F1F")
    sheet["A1"].fill = _TITLE_FILL
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 32

    metadata = [
        ("شناسه گزارش", report.get("reconciliation_id", "")),
        ("نام فایل", file_info.get("filename", "")),
        ("عنوان حساب", account_display_label),
        ("نام بانک", bank_account.get("bank_name", "")),
        ("شماره حساب کامل", account_number),
        (
            "شماره شبا",
            bank_account.get("international_number")
            or file_info.get("statement_iban", ""),
        ),
        (
            "شعبه",
            " - ".join(
                value
                for value in (
                    str(bank_account.get("bank_branch_name") or ""),
                    str(bank_account.get("bank_branch_code") or ""),
                )
                if value
            ),
        ),
        (
            "دوره صورتحساب شمسی",
            file_info.get("statement_period_jalali")
            or file_info.get("statement_period", ""),
        ),
        (
            "تاریخ ایجاد گزارش",
            format_jalali_datetime(report.get("created_at"))
            or report.get("created_at", ""),
        ),
    ]
    for row_index, (label, value) in enumerate(metadata, start=3):
        sheet.cell(row_index, 1, label).font = Font(bold=True)
        sheet.cell(row_index, 1).fill = PatternFill("solid", fgColor="EAF2F8")
        value_cell = sheet.cell(row_index, 2, value)
        sheet.merge_cells(
            start_row=row_index,
            start_column=2,
            end_row=row_index,
            end_column=5,
        )
        is_ltr_value = label in {
            "شناسه گزارش",
            "شماره حساب کامل",
            "شماره شبا",
            "تاریخ ایجاد گزارش",
        }
        value_cell.alignment = Alignment(
            horizontal="right",
            vertical="center",
            wrap_text=True,
            readingOrder=1 if is_ltr_value else 2,
        )
        sheet.row_dimensions[row_index].height = (
            34 if label == "دوره صورتحساب شمسی" else 24
        )

    sheet.append([])
    table_row = 12
    headers = ["وضعیت", "تعداد", "مبلغ (ریال)", "سهم از تعداد"]
    for column, header in enumerate(headers, start=1):
        sheet.cell(table_row, column, header)
    _style_header(sheet, table_row, len(headers))

    summary = report.get("summary", {})
    counts = summary.get("business_counts", {})
    amounts = summary.get("business_amounts", {})
    total_count = int(summary.get("bank_row_count") or 0)
    statuses = (
        "posted",
        "reversed",
        "internal_transfer",
        "needs_review",
        "unposted",
    )
    for offset, status in enumerate(statuses, start=1):
        row_index = table_row + offset
        count = int(counts.get(status) or 0)
        sheet.cell(row_index, 1, _STATUS_LABELS[status])
        sheet.cell(row_index, 2, count)
        sheet.cell(row_index, 3, amounts.get(status) or 0)
        sheet.cell(row_index, 4, count / total_count if total_count else 0)
        sheet.cell(row_index, 1).fill = PatternFill(
            "solid",
            fgColor=_STATUS_COLORS[status],
        )

    total_row = table_row + len(statuses) + 1
    sheet.cell(total_row, 1, "جمع")
    sheet.cell(total_row, 2, total_count)
    sheet.cell(total_row, 3, summary.get("bank_total_amount") or 0)
    sheet.cell(total_row, 4, 1 if total_count else 0)
    for cell in sheet[total_row][:4]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")

    _apply_table_style(sheet, table_row + 1, total_row, 4)
    for cell in sheet["B"][table_row:total_row]:
        cell.number_format = "#,##0"
    for cell in sheet["C"][table_row:total_row]:
        cell.number_format = "#,##0"
    for cell in sheet["D"][table_row:total_row]:
        cell.number_format = "0.0%"

    direction_table_row = total_row + 2
    direction_headers = [
        "وضعیت",
        "تعداد واریز",
        "مبلغ واریز (ریال)",
        "تعداد برداشت",
        "مبلغ برداشت (ریال)",
    ]
    for column, header in enumerate(direction_headers, start=1):
        sheet.cell(direction_table_row, column, header)
    _style_header(sheet, direction_table_row, len(direction_headers))

    business_direction = summary.get("business_direction_summary") or {}
    if not business_direction:
        business_direction = {
            status: {
                direction: {
                    "count": sum(
                        1
                        for row in report.get("rows", [])
                        if row.get("business_status") == status
                        and row.get("direction") == direction
                    ),
                    "amount": sum(
                        float(row.get("amount") or 0)
                        for row in report.get("rows", [])
                        if row.get("business_status") == status
                        and row.get("direction") == direction
                    ),
                }
                for direction in ("deposit", "withdrawal")
            }
            for status in statuses
        }

    for offset, status in enumerate(statuses, start=1):
        row_index = direction_table_row + offset
        deposit = business_direction.get(status, {}).get("deposit", {})
        withdrawal = business_direction.get(status, {}).get("withdrawal", {})
        sheet.cell(row_index, 1, _STATUS_LABELS[status])
        sheet.cell(row_index, 2, deposit.get("count") or 0)
        sheet.cell(row_index, 3, deposit.get("amount") or 0)
        sheet.cell(row_index, 4, withdrawal.get("count") or 0)
        sheet.cell(row_index, 5, withdrawal.get("amount") or 0)
        sheet.cell(row_index, 1).fill = PatternFill(
            "solid",
            fgColor=_STATUS_COLORS[status],
        )

    direction_total_row = direction_table_row + len(statuses) + 1
    overall_direction = summary.get("direction_summary") or {}
    deposit_total = overall_direction.get("deposit", {})
    withdrawal_total = overall_direction.get("withdrawal", {})
    sheet.cell(direction_total_row, 1, "جمع")
    sheet.cell(direction_total_row, 2, deposit_total.get("count") or 0)
    sheet.cell(direction_total_row, 3, deposit_total.get("amount") or 0)
    sheet.cell(direction_total_row, 4, withdrawal_total.get("count") or 0)
    sheet.cell(direction_total_row, 5, withdrawal_total.get("amount") or 0)
    for cell in sheet[direction_total_row][:5]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")

    _apply_table_style(
        sheet,
        direction_table_row + 1,
        direction_total_row,
        5,
    )
    for column in (2, 3, 4, 5):
        for row_index in range(direction_table_row + 1, direction_total_row + 1):
            sheet.cell(row_index, column).number_format = "#,##0"

    metric_row = direction_total_row + 2
    definite_rate = summary.get("definite_match_rate")
    confirmed_rate = summary.get("confirmed_coverage_rate")
    tolerant_rate = summary.get("tolerant_coverage_rate")
    resolution_rate = summary.get("definite_resolution_rate")
    if definite_rate is None:
        documentable_count = int(summary.get("documentable_bank_row_count") or 0)
        if not documentable_count:
            documentable_count = max(
                0,
                total_count
                - int(counts.get("reversed") or 0)
                - int(counts.get("internal_transfer") or 0),
            )
        definite_rate = (
            int(counts.get("posted") or 0) / documentable_count
            if documentable_count
            else 0
        )
    if confirmed_rate is None:
        confirmed_rate = (
            (
                int(counts.get("posted") or 0)
                + int(counts.get("reversed") or 0)
            )
            / total_count
            if total_count
            else 0
        )
    if tolerant_rate is None:
        tolerant_rate = (
            (
                int(counts.get("posted") or 0)
                + int(counts.get("reversed") or 0)
                + int(counts.get("needs_review") or 0)
            )
            / total_count
            if total_count
            else 0
        )
    if resolution_rate is None:
        policy_count = int(counts.get("internal_transfer") or 0)
        resolution_rate = (
            (
                int(counts.get("posted") or 0)
                + int(counts.get("reversed") or 0)
                + policy_count
            )
            / total_count
            if total_count
            else 0
        )
    metrics = [
        ("نرخ تطبیق قطعی", definite_rate),
        ("پوشش قطعی با احتساب برگشت", confirmed_rate),
        ("پوشش با تلورانس و موارد احتمالی", tolerant_rate),
        ("نرخ تعیین تکلیف قطعی", resolution_rate),
    ]
    for offset, (label, value) in enumerate(metrics):
        row_index = metric_row + offset
        sheet.cell(row_index, 1, label).font = Font(bold=True)
        sheet.cell(row_index, 1).fill = PatternFill("solid", fgColor="EAF2F8")
        sheet.cell(row_index, 2, value).number_format = "0.0%"
        sheet.merge_cells(
            start_row=row_index,
            start_column=2,
            end_row=row_index,
            end_column=3,
        )
        sheet.merge_cells(
            start_row=row_index,
            start_column=4,
            end_row=row_index,
            end_column=5,
        )
        sheet.cell(row_index, 4, {
            0: "فقط سندهای قطعی",
            1: "سند قطعی و تراکنش خنثی‌شده",
            2: "قطعی، خنثی‌شده و نیازمند بررسی",
            3: "سند قطعی، تراکنش خنثی‌شده و انتقال داخلی شناسایی‌شده",
        }[offset])
        sheet.cell(row_index, 4).alignment = Alignment(wrap_text=True)

    notes_row = metric_row + len(metrics) + 1
    sheet.cell(notes_row, 1, "توضیحات مهم").font = Font(bold=True)
    sheet.cell(
        notes_row + 1,
        1,
        "بدون سند احتمالی یعنی با قواعد فعلی تطبیق خودکار پیدا نشده است و باید توسط خزانه بررسی شود.",
    )
    sheet.merge_cells(
        start_row=notes_row + 1,
        start_column=1,
        end_row=notes_row + 1,
        end_column=5,
    )
    sheet.cell(
        notes_row + 2,
        1,
        "انتقال داخلی شرکت از بدون سند جداست و در مبلغ و درصد بدون سند محاسبه نمی‌شود.",
    )
    sheet.merge_cells(
        start_row=notes_row + 2,
        start_column=1,
        end_row=notes_row + 2,
        end_column=5,
    )
    for row_index in (notes_row + 1, notes_row + 2):
        sheet.cell(row_index, 1).alignment = Alignment(wrap_text=True)
        sheet.row_dimensions[row_index].height = 34

    _set_widths(sheet, [22, 16, 22, 16, 22])
    sheet.freeze_panes = "A12"
    sheet.print_area = f"A1:E{notes_row + 2}"


def build_reconciliation_workbook(report: dict[str, Any]) -> BytesIO:
    """Build a Persian, review-ready Excel report without storing bank files."""

    workbook = Workbook()
    workbook.properties.title = "گزارش مغایرت‌گیری خزانه"
    workbook.properties.subject = "مقایسه صورتحساب بانک با اسناد راهکاران"
    workbook.properties.creator = "Exir Kadous AI Platform"

    _add_summary_sheet(workbook, report)
    for status in (
        "posted",
        "reversed",
        "internal_transfer",
        "needs_review",
        "unposted",
    ):
        _add_bank_status_sheet(workbook, report, status)
    _add_erp_only_sheet(workbook, report)
    _add_daily_balance_sheet(workbook, report)
    _add_fee_sheet(workbook, report)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
