"""اسناد خرید انبار راهکاران — نسخه زنده دو گزارش واحد مالی (فقط خواندنی).

۱) «عملیات گروهی اسناد»: اسناد خرید داخلی/خارجی و برگشت از خرید سال جاری که «ثبت حسابداری نشده»اند،
   با وضعیت قیمت‌گذاری، حسابداری و عامل‌گذاری — همان ستون‌های فهرست اسناد انبار راهکاران.
   مبنا: LGS3.vwInventoryVoucherWithAllStates (AccountingState: 1 نشده، 2 شده؛ PricingState: 1 نشده، 2 شده؛
   SLResolvingState: 1 نشده، 2 شده). اسناد باطل‌شده (State = 3) حساب نمی‌شوند.
۲) «مرور تدارکات»: ردیف‌های رسید خرید دائم و «رسید موقت خرید». ردیف رسید دائم با ReferenceType = 1 به ردیف
   رسید موقتش اشاره می‌کند (ReferenceRef = InventoryVoucherItemID)؛ رسید موقتی که هیچ رسید دائمی به آن اشاره
   نکرده «تبدیل‌نشده» است.
"""

from __future__ import annotations

import threading
import time
from datetime import date, datetime
from io import BytesIO
from typing import Any

import jdatetime
from sqlalchemy import text

from app.utils.jalali import format_jalali_date

# انواع سند (LGS3.InventoryVoucherSpecification)
PURCHASE_SPECS = (2, 3, 55, 64)          # خرید داخلی، خرید خارجی، برگشت به خرید داخلی/خارجی
RECEIPT_SPECS = (2, 3)                   # رسید خرید دائم
TEMP_RECEIPT_SPEC = 10001                # رسید موقت خرید
VOIDED = 3
CACHE_SECONDS = 300

VOUCHER_STATES = {1: "ثبت شده", 2: "تایید شده", 3: "باطل شده"}
PRICING_STATES = {1: "قیمت گذاری نشده", 2: "قیمت گذاری شده"}
ACCOUNTING_STATES = {1: "ثبت حسابداری نشده", 2: "ثبت حسابداری شده", 3: "ثبت حسابداری ناقص"}
SL_STATES = {1: "عامل گذاری نشده", 2: "عامل گذاری شده"}
JALALI_MONTHS = ("فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور", "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند")

_cache: dict[Any, tuple[float, Any]] = {}
_lock = threading.Lock()


def year_start() -> date:
    return jdatetime.date(jdatetime.date.today().year, 1, 1).togregorian()


def _as_date(value: Any) -> date | None:
    if value is None:
        return None
    return value.date() if isinstance(value, datetime) else value


def _t(value: Any) -> str:
    return " ".join(str(value or "").split())


def _ids(values) -> str:
    return ",".join(str(int(v)) for v in values)


class PurchaseDocumentsService:
    def __init__(self, engine=None):
        self._engine = engine

    @property
    def engine(self):
        if self._engine is None:
            from app.database.sqlserver import get_sqlserver_engine
            self._engine = get_sqlserver_engine()
        return self._engine

    def _cached(self, key, build, refresh: bool = False):
        with _lock:
            hit = _cache.get(key)
            if hit and not refresh and time.time() - hit[0] < CACHE_SECONDS:
                return hit[1]
        value = build()
        with _lock:
            _cache[key] = (time.time(), value)
        return value

    # ------------------------------------------------------- عملیات گروهی اسناد
    def documents(self, accounting: str = "unposted", from_date: date | None = None, refresh: bool = False) -> dict[str, Any]:
        start = from_date or year_start()
        return self._cached(("docs", accounting, start), lambda: self._documents(accounting, start), refresh)

    def _documents(self, accounting: str, start: date) -> dict[str, Any]:
        where = " AND v.AccountingState = 1" if accounting == "unposted" else ""
        today = date.today()
        with self.engine.connect() as c:
            rows = c.execute(text(f"""
                SELECT v.InventoryVoucherID, v.Number, v.Date, v.State, v.PricingState, v.AccountingState, v.SLResolvingState,
                       sp.Name AS Title, sp.Direction, s.Name AS StoreName, cs.Name AS CounterpartStore,
                       v.CounterpartEntityText, v.CounterpartDLCode, p.FullName AS Deliverer, b.Title AS BranchName,
                       v.Description, v.Additional1, v.Additional2, v.Additional3, v.Additional4, v.Additional5,
                       u.Name AS CreatorName, it.ItemCount, it.PriceRial
                FROM LGS3.vwInventoryVoucherWithAllStates v
                JOIN LGS3.InventoryVoucherSpecification sp ON sp.InventoryVoucherSpecificationID = v.InventoryVoucherSpecificationRef
                JOIN LGS3.Store s ON s.StoreID = v.StoreRef
                LEFT JOIN LGS3.Store cs ON cs.StoreID = v.CounterpartStoreRef
                LEFT JOIN GNR3.Party p ON p.PartyID = v.DelivererOrReceiverPartyRef
                LEFT JOIN GNR3.Branch b ON b.BranchID = v.BranchRef
                LEFT JOIN SYS3.[User] u ON u.UserID = v.Creator
                OUTER APPLY (
                    SELECT COUNT(*) AS ItemCount, SUM(ip.Price) AS PriceRial
                    FROM LGS3.InventoryVoucherItem i
                    LEFT JOIN LGS3.ItemPrice ip ON ip.InventoryVoucherItemRef = i.InventoryVoucherItemID
                    WHERE i.InventoryVoucherRef = v.InventoryVoucherID
                ) it
                WHERE v.InventoryVoucherSpecificationRef IN ({_ids(PURCHASE_SPECS)}) AND v.State <> {VOIDED}
                  AND v.Date >= :from_date {where}
                ORDER BY v.Date DESC, v.Number DESC
            """), {"from_date": start}).mappings().all()
            # اسناد ثبت‌حسابداری‌نشده سال‌های قبل (فقط شمارش، برای هشدار)
            older = c.execute(text(f"""
                SELECT COUNT(*) FROM LGS3.vwInventoryVoucherWithAllStates v
                WHERE v.InventoryVoucherSpecificationRef IN ({_ids(PURCHASE_SPECS)}) AND v.State <> {VOIDED}
                  AND v.Date < :from_date AND v.AccountingState = 1
            """), {"from_date": start}).scalar() or 0

        docs = []
        for r in rows:
            d = _as_date(r["Date"])
            docs.append({
                "voucher_id": int(r["InventoryVoucherID"]),
                "number": _t(r["Number"]),
                "title": _t(r["Title"]),
                "is_return": r["Direction"] == 2,
                "store_name": _t(r["StoreName"]),
                "date_jalali": format_jalali_date(d) or "",
                "days_waiting": (today - d).days if d else None,
                "counterpart": _t(r["CounterpartEntityText"]),
                "counterpart_code": _t(r["CounterpartDLCode"]),
                "state": VOUCHER_STATES.get(r["State"], str(r["State"])),
                "pricing_state": PRICING_STATES.get(r["PricingState"], str(r["PricingState"])),
                "priced": r["PricingState"] == 2,
                "accounting_state": ACCOUNTING_STATES.get(r["AccountingState"], str(r["AccountingState"])),
                "posted": r["AccountingState"] == 2,
                "sl_state": SL_STATES.get(r["SLResolvingState"], str(r["SLResolvingState"])),
                "deliverer": _t(r["Deliverer"]),
                "branch_name": _t(r["BranchName"]),
                "counterpart_store": _t(r["CounterpartStore"]),
                "description": _t(r["Description"]),
                "additional": [_t(r[f"Additional{i}"]) for i in range(1, 6)],
                "creator": _t(r["CreatorName"]),
                "item_count": int(r["ItemCount"] or 0),
                "price_rial": float(r["PriceRial"] or 0),
            })

        def group(key):
            out: dict[str, dict[str, Any]] = {}
            for x in docs:
                g = out.setdefault(x[key] or "—", {"name": x[key] or "—", "count": 0, "unpriced": 0, "price_rial": 0.0, "oldest_days": 0})
                g["count"] += 1
                g["unpriced"] += 0 if x["priced"] else 1
                g["price_rial"] += x["price_rial"]
                g["oldest_days"] = max(g["oldest_days"], x["days_waiting"] or 0)
            return sorted(out.values(), key=lambda g: -g["count"])

        months: dict[tuple[int, int], dict[str, int]] = {}
        for x in docs:
            y, m = int(x["date_jalali"][:4]), int(x["date_jalali"][5:7])
            g = months.setdefault((y, m), {"count": 0, "unpriced": 0})
            g["count"] += 1
            g["unpriced"] += 0 if x["priced"] else 1
        priced = [x for x in docs if x["priced"]]
        return {
            "status": "success",
            "from_date_jalali": format_jalali_date(start),
            "generated_at_jalali": format_jalali_date(today),
            "accounting": accounting,
            "totals": {
                "count": len(docs),
                "priced_count": len(priced),
                "unpriced_count": len(docs) - len(priced),
                "return_count": sum(1 for x in docs if x["is_return"]),
                "priced_rial": sum(x["price_rial"] for x in priced),
                "over_30_days": sum(1 for x in docs if (x["days_waiting"] or 0) > 30),
                "older_years_unposted": int(older),
            },
            "by_store": group("store_name"),
            "by_creator": group("creator"),
            "by_month": [{"label": f"{JALALI_MONTHS[m - 1]} {y}", **v} for (y, m), v in sorted(months.items())],
            "documents": docs,
            "rule": "اسناد خرید داخلی/خارجی و برگشت از خرید از ابتدای سال که در راهکاران «ثبت حسابداری نشده»اند؛ اسناد باطل‌شده حساب نمی‌شوند.",
        }

    # -------------------------------------------------------------- مرور تدارکات
    def items(self, from_date: date | None = None, refresh: bool = False) -> dict[str, Any]:
        start = from_date or year_start()
        return self._cached(("items", start), lambda: self._items(start), refresh)

    def _items(self, start: date) -> dict[str, Any]:
        today = date.today()
        specs = _ids(RECEIPT_SPECS + (TEMP_RECEIPT_SPEC,))
        with self.engine.connect() as c:
            rows = c.execute(text(f"""
                SELECT i.InventoryVoucherItemID, v.Number, v.Date, v.InventoryVoucherSpecificationRef AS Spec, sp.Name AS Pattern,
                       pt.Code AS PartCode, pt.Name AS PartName, i.Quantity, s.Name AS StoreName, un.Name AS UnitName,
                       v.CounterpartEntityText, v.CounterpartDLCode,
                       tv.Number AS TempNumber, tv.Date AS TempDate
                FROM LGS3.InventoryVoucherItem i
                JOIN LGS3.InventoryVoucher v ON v.InventoryVoucherID = i.InventoryVoucherRef
                JOIN LGS3.InventoryVoucherSpecification sp ON sp.InventoryVoucherSpecificationID = v.InventoryVoucherSpecificationRef
                JOIN LGS3.Store s ON s.StoreID = v.StoreRef
                LEFT JOIN LGS3.Part pt ON pt.PartID = i.PartRef
                LEFT JOIN GNR3.Unit un ON un.UnitID = i.UnitRef
                LEFT JOIN LGS3.InventoryVoucherItem ti ON i.ReferenceType = 1 AND ti.InventoryVoucherItemID = i.ReferenceRef
                LEFT JOIN LGS3.InventoryVoucher tv ON tv.InventoryVoucherID = ti.InventoryVoucherRef AND tv.InventoryVoucherSpecificationRef = {TEMP_RECEIPT_SPEC}
                WHERE v.InventoryVoucherSpecificationRef IN ({specs}) AND v.State <> {VOIDED} AND v.Date >= :from_date
                ORDER BY v.Date DESC, v.Number DESC, i.RowNumber
            """), {"from_date": start}).mappings().all()
            # رسید دائمِ هر ردیف رسید موقت (هر تاریخی؛ ممکن است رسید دائم بعد از بازه ثبت شده باشد)
            links = c.execute(text(f"""
                SELECT i.ReferenceRef AS TempItem, v.Number, v.Date
                FROM LGS3.InventoryVoucherItem i
                JOIN LGS3.InventoryVoucher v ON v.InventoryVoucherID = i.InventoryVoucherRef
                JOIN LGS3.InventoryVoucherItem ti ON ti.InventoryVoucherItemID = i.ReferenceRef
                JOIN LGS3.InventoryVoucher tv ON tv.InventoryVoucherID = ti.InventoryVoucherRef
                WHERE i.ReferenceType = 1 AND v.InventoryVoucherSpecificationRef IN ({_ids(RECEIPT_SPECS)}) AND v.State <> {VOIDED}
                  AND tv.InventoryVoucherSpecificationRef = {TEMP_RECEIPT_SPEC} AND tv.Date >= :from_date
            """), {"from_date": start}).mappings().all()
        permanent_of: dict[int, tuple[str, date | None]] = {}
        for l in links:
            permanent_of[int(l["TempItem"])] = (_t(l["Number"]), _as_date(l["Date"]))

        items = []
        for r in rows:
            d = _as_date(r["Date"])
            temp = r["Spec"] == TEMP_RECEIPT_SPEC
            link = permanent_of.get(int(r["InventoryVoucherItemID"])) if temp else None
            temp_date = _as_date(r["TempDate"])
            items.append({
                "item_id": int(r["InventoryVoucherItemID"]),
                "number": _t(r["Number"]),
                "date_jalali": format_jalali_date(d) or "",
                "part_code": _t(r["PartCode"]),
                "part_name": _t(r["PartName"]),
                "quantity": float(r["Quantity"] or 0),
                "store_name": _t(r["StoreName"]),
                "unit": _t(r["UnitName"]),
                "counterpart": _t(r["CounterpartEntityText"]),
                "counterpart_code": _t(r["CounterpartDLCode"]),
                "pattern": _t(r["Pattern"]),
                "kind": "موقت" if temp else "دائم",
                # موقت ← شماره رسید دائم؛ دائم ← شماره رسید موقت مبدأ
                "linked_number": (link[0] if link else "") if temp else _t(r["TempNumber"]),
                "linked_date_jalali": (format_jalali_date(link[1]) if link and link[1] else "") if temp
                else (format_jalali_date(temp_date) if temp_date else ""),
                "converted": bool(link) if temp else None,
                "days_open": (today - d).days if temp and not link and d else None,
            })
        temps = [x for x in items if x["kind"] == "موقت"]
        pending = [x for x in temps if not x["converted"]]
        perm = [x for x in items if x["kind"] == "دائم"]
        return {
            "status": "success",
            "from_date_jalali": format_jalali_date(start),
            "generated_at_jalali": format_jalali_date(today),
            "totals": {
                "item_count": len(items),
                "voucher_count": len({(x["kind"], x["number"]) for x in items}),
                "permanent_items": len(perm),
                "temporary_items": len(temps),
                "pending_items": len(pending),
                "pending_vouchers": len({x["number"] for x in pending}),
                "permanent_without_temp": sum(1 for x in perm if not x["linked_number"]),
                "parts": len({x["part_code"] for x in items}),
                "suppliers": len({x["counterpart_code"] or x["counterpart"] for x in items}),
            },
            "items": items,
            "rule": "رسید موقت «تبدیل‌نشده» یعنی هیچ ردیف رسید خرید دائمی (باطل‌نشده) به آن ارجاع نداده است.",
        }

    # ------------------------------------------------------ یک سند و سند مرتبطش
    def voucher_pair(self, item_id: int) -> dict[str, Any]:
        """سند ردیف item_id و سند(های) مرتبط موقت↔دائم، هر کدام با سربرگ و همه ردیف‌ها."""
        with self.engine.connect() as c:
            row = c.execute(text("""
                SELECT i.InventoryVoucherRef, i.ReferenceType, i.ReferenceRef, v.InventoryVoucherSpecificationRef AS Spec
                FROM LGS3.InventoryVoucherItem i JOIN LGS3.InventoryVoucher v ON v.InventoryVoucherID = i.InventoryVoucherRef
                WHERE i.InventoryVoucherItemID = :id
            """), {"id": int(item_id)}).mappings().first()
            if not row:
                raise ValueError("ردیف پیدا نشد")
            if row["Spec"] == TEMP_RECEIPT_SPEC:
                # موقت ← رسیدهای دائمی که به این ردیف ارجاع داده‌اند
                linked_ids = c.execute(text(f"""
                    SELECT DISTINCT i.InventoryVoucherRef FROM LGS3.InventoryVoucherItem i
                    JOIN LGS3.InventoryVoucher v ON v.InventoryVoucherID = i.InventoryVoucherRef
                    WHERE i.ReferenceType = 1 AND i.ReferenceRef = :id AND v.State <> {VOIDED}
                      AND v.InventoryVoucherSpecificationRef IN ({_ids(RECEIPT_SPECS)})
                """), {"id": int(item_id)}).scalars().all()
                linked_item = None
            else:
                linked_ids = c.execute(text("""
                    SELECT ti.InventoryVoucherRef FROM LGS3.InventoryVoucherItem ti WHERE ti.InventoryVoucherItemID = :ref
                """), {"ref": row["ReferenceRef"]}).scalars().all() if row["ReferenceType"] == 1 and row["ReferenceRef"] else []
                linked_item = int(row["ReferenceRef"]) if linked_ids else None
            return {
                "status": "success",
                "selected_item_id": int(item_id),
                "linked_item_id": linked_item,
                "document": self._voucher(c, int(row["InventoryVoucherRef"])),
                "linked": [self._voucher(c, int(v)) for v in linked_ids],
            }

    def _voucher(self, c, voucher_id: int) -> dict[str, Any]:
        h = c.execute(text("""
            SELECT v.InventoryVoucherID, v.Number, v.Date, v.State, v.CreationDate, v.Description, v.InventoryVoucherSpecificationRef AS Spec,
                   sp.Name AS Pattern, s.Name AS StoreName, v.CounterpartEntityText, v.CounterpartDLCode, u.Name AS CreatorName,
                   vs.PricingState, vs.AccountingState
            FROM LGS3.InventoryVoucher v
            JOIN LGS3.InventoryVoucherSpecification sp ON sp.InventoryVoucherSpecificationID = v.InventoryVoucherSpecificationRef
            JOIN LGS3.Store s ON s.StoreID = v.StoreRef
            LEFT JOIN SYS3.[User] u ON u.UserID = v.Creator
            LEFT JOIN LGS3.vwInventoryVoucherWithAllStates vs ON vs.InventoryVoucherID = v.InventoryVoucherID
            WHERE v.InventoryVoucherID = :id
        """), {"id": voucher_id}).mappings().first()
        items = c.execute(text("""
            SELECT i.InventoryVoucherItemID, i.RowNumber, i.ReferenceType, i.ReferenceRef, pt.Code AS PartCode, pt.Name AS PartName,
                   i.Quantity, un.Name AS UnitName
            FROM LGS3.InventoryVoucherItem i
            LEFT JOIN LGS3.Part pt ON pt.PartID = i.PartRef
            LEFT JOIN GNR3.Unit un ON un.UnitID = i.UnitRef
            WHERE i.InventoryVoucherRef = :id ORDER BY i.RowNumber
        """), {"id": voucher_id}).mappings().all()
        created = h["CreationDate"]
        return {
            "voucher_id": voucher_id,
            "number": _t(h["Number"]),
            "date_jalali": format_jalali_date(_as_date(h["Date"])) or "",
            "kind": "موقت" if h["Spec"] == TEMP_RECEIPT_SPEC else "دائم",
            "pattern": _t(h["Pattern"]),
            "store_name": _t(h["StoreName"]),
            "counterpart": _t(h["CounterpartEntityText"]),
            "counterpart_code": _t(h["CounterpartDLCode"]),
            "state": VOUCHER_STATES.get(h["State"], str(h["State"])),
            "pricing_state": PRICING_STATES.get(h["PricingState"], "") if h["PricingState"] else "",
            "accounting_state": ACCOUNTING_STATES.get(h["AccountingState"], "") if h["AccountingState"] else "",
            "creator": _t(h["CreatorName"]),
            "created_at_jalali": (f"{format_jalali_date(created.date())} {created:%H:%M}" if isinstance(created, datetime) else ""),
            "description": _t(h["Description"]),
            "items": [{
                "item_id": int(i["InventoryVoucherItemID"]),
                "row": int(i["RowNumber"] or 0),
                "part_code": _t(i["PartCode"]),
                "part_name": _t(i["PartName"]),
                "quantity": float(i["Quantity"] or 0),
                "unit": _t(i["UnitName"]),
                "ref_item_id": int(i["ReferenceRef"]) if i["ReferenceType"] == 1 and i["ReferenceRef"] else None,
            } for i in items],
        }

    # ------------------------------------------------------------------ excel
    def workbook(self, report: str, from_date: date | None = None, year: int | None = None, month: int | None = None) -> BytesIO:
        """year/month: فقط اسناد همان سال (و ماه) شمسی؛ هم‌خوان با فیلتر صفحه."""
        prefix = f"{year}/{month:02d}" if year and month else (f"{year}/" if year else "")

        def in_period(x):
            return not prefix or str(x.get("date_jalali") or "").startswith(prefix)

        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        head = PatternFill("solid", fgColor="1F4E78")
        wb = Workbook()

        def sheet(ws, headers, rows, widths):
            ws.sheet_view.rightToLeft = True
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = head
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for row in rows:
                ws.append(row)
            for i, w in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(i)].width = w
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

        if report == "documents":
            data = self.documents(from_date=from_date)
            ws = wb.active
            ws.title = "عملیات گروهی اسناد"
            # همان ستون‌های خروجی راهکاران + سه ستون کمکی در انتها
            sheet(ws, ["شماره سند", "عنوان سند", "عنوان انبار", "تاریخ سند", "طرف مقابل", "کد تفصیلی طرف مقابل", "وضعیت",
                       "وضعیت قیمت گذاری", "وضعیت حسابداری", "وضعیت عامل گذاری", "تحویل گیرنده/دهنده", "شعبه", "انبار مقابل",
                       "توضیحات", "فیلد اضافی 1", "فیلد اضافی 2", "فیلد اضافی 3", "فیلد اضافی 4", "فیلد اضافی 5", "صادر کننده",
                       "تعداد ردیف", "مبلغ قیمت‌گذاری (ریال)", "روز از تاریخ سند"],
                  [[d["number"], d["title"], d["store_name"], d["date_jalali"], d["counterpart"], d["counterpart_code"], d["state"],
                    d["pricing_state"], d["accounting_state"], d["sl_state"], d["deliverer"], d["branch_name"], d["counterpart_store"],
                    d["description"], *d["additional"], d["creator"], d["item_count"], d["price_rial"] or None, d["days_waiting"]]
                   for d in data["documents"] if in_period(d)],
                  [10, 18, 30, 12, 34, 14, 11, 16, 18, 16, 22, 12, 18, 60, 10, 10, 10, 10, 10, 20, 10, 18, 12])
            for row in ws.iter_rows(min_row=2, min_col=22, max_col=22):
                for cell in row:
                    cell.number_format = "#,##0"
            sheet(wb.create_sheet("خلاصه انبار"), ["انبار", "تعداد سند", "قیمت گذاری نشده", "مبلغ (ریال)", "قدیمی‌ترین (روز)"],
                  [[g["name"], g["count"], g["unpriced"], g["price_rial"], g["oldest_days"]] for g in data["by_store"]], [34, 12, 16, 20, 16])
            sheet(wb.create_sheet("خلاصه صادرکننده"), ["صادر کننده", "تعداد سند", "قیمت گذاری نشده", "مبلغ (ریال)", "قدیمی‌ترین (روز)"],
                  [[g["name"], g["count"], g["unpriced"], g["price_rial"], g["oldest_days"]] for g in data["by_creator"]], [24, 12, 16, 20, 16])
        else:
            data = self.items(from_date=from_date)
            ws = wb.active
            ws.title = "مرور تدارکات"
            sheet(ws, ["شماره سند انبار", "تاریخ سند انبار", "کد قلم", "عنوان قلم", "مقدار", "نام انبار", "واحد", "طرف مقابل",
                       "نام الگوی سند انبار", "دائم / موقت", "سند مرتبط (موقت↔دائم)", "تاریخ سند مرتبط", "وضعیت تبدیل"],
                  [[x["number"], x["date_jalali"], x["part_code"], x["part_name"], x["quantity"], x["store_name"], x["unit"],
                    x["counterpart"], x["pattern"], x["kind"], x["linked_number"], x["linked_date_jalali"],
                    "" if x["converted"] is None else ("تبدیل شده" if x["converted"] else "تبدیل نشده")] for x in data["items"] if in_period(x)],
                  [12, 12, 14, 44, 12, 30, 10, 34, 18, 10, 18, 14, 14])
            pending = [x for x in data["items"] if x["converted"] is False and in_period(x)]
            sheet(wb.create_sheet("رسید موقت تبدیل‌نشده"),
                  ["شماره رسید موقت", "تاریخ", "روز باز", "کد قلم", "عنوان قلم", "مقدار", "واحد", "نام انبار", "طرف مقابل"],
                  [[x["number"], x["date_jalali"], x["days_open"], x["part_code"], x["part_name"], x["quantity"], x["unit"],
                    x["store_name"], x["counterpart"]] for x in pending],
                  [14, 12, 10, 14, 44, 12, 10, 30, 34])
        info = wb.create_sheet("توضیحات")
        info.sheet_view.rightToLeft = True
        period = f"دوره {prefix.rstrip('/')}" if prefix else f"از {data['from_date_jalali']}"
        info.append([f"راهکاران — زنده؛ {period} (تا {data['generated_at_jalali']})"])
        info.append([data["rule"]])
        info.column_dimensions["A"].width = 110
        out = BytesIO()
        wb.save(out)
        out.seek(0)
        return out
