"""Excel: per hybrid → visitor, unsettled returned cheques and remaining customer debt (Karamad).

Read-only: SELECT queries only, through the finance backend's own Karamad connection and
debt logic (same numbers as the finance dashboard's hybrid / visitor files).

Usage:  python tools/export_returned_debt.py [output.xlsx]
"""
from __future__ import annotations

import os
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

FINANCE_BACKEND = Path(os.environ.get("FINANCE_BACKEND", Path(__file__).resolve().parents[2] / "backend"))
sys.path.insert(0, str(FINANCE_BACKEND))
os.chdir(FINANCE_BACKEND)  # the backend reads its .env from here

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from sqlalchemy import bindparam, text  # noqa: E402

from app.database.karamad_sqlserver import get_karamad_sqlserver_engine  # noqa: E402
from app.services.customer_file_service import KARAMAD_HYBRID_DL_CLASS_ID  # noqa: E402
from app.services.karamad_sales_network_service import (  # noqa: E402
    OPEN_RETURNED_CHEQUE_STATUSES, _allocate_debt, _group_by_customer, ledger_positions,
)
from app.utils.jalali import format_jalali_date  # noqa: E402

TOMAN = 10  # amounts are stored in rial
NO_VISITOR = "بدون ویزیتور (مشتری فاکتوری در کارآمد ندارد)"


def t(v) -> str:
    return str(v or "").strip()


def load():
    eng = get_karamad_sqlserver_engine()
    expand = lambda sql, *names: text(sql).bindparams(*(bindparam(n, expanding=True) for n in names))  # noqa: E731
    returned_statuses = list(OPEN_RETURNED_CHEQUE_STATUSES)
    with eng.connect() as c:
        hybrids = defaultdict(list)
        for r in c.execute(text("SELECT [BranchRef], [Code], [Name] FROM dbo.[tblDL] WHERE [ClassRef] = :cls AND [BranchRef] IS NOT NULL ORDER BY [Name]"),
                           {"cls": KARAMAD_HYBRID_DL_CLASS_ID}).mappings():
            hybrids[int(r["BranchRef"])].append((t(r["Code"]), t(r["Name"])))
        branch_ids = sorted(hybrids)
        p = {"branches": branch_ids or [-1], "st": returned_statuses}
        branches = {int(r["ID"]): {"code": r["Code"], "name": t(r["Name"])} for r in c.execute(
            expand("SELECT [ID], [Code], [Name] FROM dbo.[tblBranch] WHERE [ID] IN :branches", "branches"), p).mappings()}
        visitors = {int(r["ID"]): {
            "branch": int(r["BranchRef"]) if r["BranchRef"] is not None else None, "code": r["Code"], "name": t(r["Name"]),
            "active": bool(r["Status"]), "supervisor": t(r["Supervisor"]), "leave": r["LeaveDate"],
        } for r in c.execute(text(
            "SELECT v.[ID], v.[BranchRef], v.[Code], v.[Name], v.[Status], v.[LeaveDate], s.[Name] AS Supervisor "
            "FROM dbo.[tblVisitor] v LEFT JOIN dbo.[tblSupervisor] s ON s.[ID] = v.[SupervisorRef]")).mappings()}

        scope = """SELECT cu.[ID] FROM dbo.[tblCustomer] cu WHERE cu.[DLRef] IS NOT NULL AND (cu.[BranchRef] IN :branches
                   OR cu.[ID] IN (SELECT f.[CustomerRef] FROM dbo.[tblFactorF] f INNER JOIN dbo.[tblVisitor] v ON v.[ID] = f.[VisitorRef]
                                  WHERE v.[BranchRef] IN :branches))"""
        customers = {}
        for r in c.execute(expand(f"SELECT [ID], [DLRef], [Code], [Name], [BranchRef] FROM dbo.[tblCustomer] WHERE [ID] IN ({scope}) ORDER BY [ID]", "branches"), p).mappings():
            customers.setdefault(int(r["DLRef"]), {"code": t(r["Code"]), "name": t(r["Name"]),
                                                   "branch": int(r["BranchRef"]) if r["BranchRef"] is not None else None})
        invoices = c.execute(expand(f"""
            SELECT cu.[DLRef], f.[VisitorRef], f.[FactorPriceP], f.[DateE]
            FROM dbo.[tblFactorF] f INNER JOIN dbo.[tblCustomer] cu ON cu.[ID] = f.[CustomerRef]
            WHERE cu.[DLRef] IS NOT NULL AND f.[CustomerRef] IN ({scope})
            ORDER BY cu.[DLRef], f.[DateE] DESC, f.[ID] DESC""", "branches"), p).mappings().all()
        # every unsettled returned cheque: for the listing (by hybrid) and for the debt allocation (by customer)
        cheques = c.execute(expand("""
            SELECT d.[inx], d.[Serial], d.[InquiryCode], d.[Price], d.[DueDate], d.[BookDate], d.[DLRef], d.[BranchRef],
                   d.[StatusRef], st.[Name] AS StatusName, COALESCE(bk.[Name], d.[Bank]) AS BankName,
                   fx.[VisitorRef], fx.[Code] AS FactorCode
            FROM dbo.[tblChequeD] d
            LEFT JOIN dbo.[tblFactorF] fx ON fx.[ID] = d.[FactorRef]
            OUTER APPLY (SELECT TOP (1) s.[Name] FROM dbo.[tblChequeDStatus] s WHERE s.[Code] = d.[StatusRef]) st
            OUTER APPLY (SELECT TOP (1) b.[Name] FROM dbo.[tblBankList] b WHERE b.[Code] = d.[BankIDRef]) bk
            WHERE d.[StatusRef] IN :st
            ORDER BY d.[DueDate]""", "st"), p).mappings().all()
        names_by_dl = {int(r["DLRef"]): {"code": t(r["Code"]), "name": t(r["Name"])} for r in c.execute(expand(
            "SELECT [DLRef], MIN([Code]) AS Code, MIN([Name]) AS Name FROM dbo.[tblCustomer] WHERE [DLRef] IN :dls GROUP BY [DLRef]", "dls"),
            {"dls": sorted({int(x["DLRef"]) for x in cheques if x["DLRef"] is not None}) or [-1]}).mappings()}
    return branch_ids, hybrids, branches, visitors, customers, invoices, cheques, names_by_dl


def build(path: Path):
    branch_ids, hybrids, branches, visitors, customers, invoices, cheques, names_by_dl = load()
    positions = ledger_positions()
    in_scope = set(branch_ids)

    def branch_of_visitor(vid, fallback):
        b = visitors.get(vid, {}).get("branch") if vid is not None else None
        return b if b in in_scope else (fallback if fallback in in_scope else None)

    # ---- unsettled returned cheques, owned by the visitor of the invoice they paid
    ret_rows = []
    for x in cheques:
        vid = int(x["VisitorRef"]) if x["VisitorRef"] is not None else None
        b = branch_of_visitor(vid, int(x["BranchRef"]) if x["BranchRef"] is not None else None)
        if b is None:
            continue
        cust = names_by_dl.get(int(x["DLRef"]), {}) if x["DLRef"] is not None else {}
        ret_rows.append({"branch": b, "visitor": vid if vid in visitors and visitors[vid]["branch"] == b else None,
                         "amount": float(x["Price"] or 0), "status": t(x["StatusName"]), "serial": t(x["Serial"]),
                         "sayad": t(x["InquiryCode"]), "bank": t(x["BankName"]), "due": format_jalali_date(x["DueDate"]),
                         "booked": format_jalali_date(x["BookDate"]), "factor": t(x["FactorCode"]),
                         "cust_code": cust.get("code", ""), "cust_name": cust.get("name", ""),
                         "days": (date.today() - (x["DueDate"].date() if hasattr(x["DueDate"], "date") else x["DueDate"])).days if x["DueDate"] else None})
    status_names = sorted({r["status"] for r in ret_rows})

    # ---- remaining debt: each customer's ledger balance split into visitor slices (finance dashboard logic)
    by_invoice, by_returned = _group_by_customer(invoices, [
        {"DLRef": x["DLRef"], "VisitorRef": x["VisitorRef"], "Price": x["Price"], "DueDate": x["DueDate"]}
        for x in cheques if x["DLRef"] is not None and int(x["DLRef"]) in customers])
    debt_rows = []
    for dl, cu in customers.items():
        balance = float(positions.get(dl, {}).get("balance_rial") or 0)
        if balance <= 0:
            continue
        for s in _allocate_debt(balance, by_invoice.get(dl, []), by_returned.get(dl, [])):
            vid = s["visitor_id"]
            b = branch_of_visitor(vid, cu["branch"])
            if b is None:
                continue
            debt_rows.append({"branch": b, "visitor": vid if vid in visitors and visitors[vid]["branch"] == b else None,
                              "dl": dl, "amount": s["amount"], "kind": s["kind"], "cust_code": cu["code"], "cust_name": cu["name"]})

    # ---- aggregate per (branch, visitor)
    agg = defaultdict(lambda: {"ret_n": 0, "ret_amt": 0.0, "by_status": defaultdict(float), "debt": 0.0,
                               "debt_from_ret": 0.0, "debtors": set()})
    for r in ret_rows:
        a = agg[(r["branch"], r["visitor"])]
        a["ret_n"] += 1
        a["ret_amt"] += r["amount"]
        a["by_status"][r["status"]] += r["amount"]
    for r in debt_rows:
        a = agg[(r["branch"], r["visitor"])]
        a["debt"] += r["amount"]
        if r["kind"] == "returned_cheque":
            a["debt_from_ret"] += r["amount"]
        a["debtors"].add(r["dl"])
    # visitors of the branch with nothing open still appear (active ones), so the list is complete
    for vid, v in visitors.items():
        if v["branch"] in in_scope and v["active"]:
            agg[(v["branch"], vid)]

    write(path, branch_ids, hybrids, branches, visitors, agg, ret_rows, debt_rows, status_names)
    return len(ret_rows), sum(r["amount"] for r in ret_rows), sum(r["amount"] for r in debt_rows)


# ------------------------------------------------------------------ excel
HEAD_FILL = PatternFill("solid", fgColor="0F3A37")
SUB_FILL = PatternFill("solid", fgColor="DDF3F0")
TOTAL_FILL = PatternFill("solid", fgColor="FCE7C8")
HEAD_FONT = Font(name="Tahoma", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Tahoma", size=10)
BOLD = Font(name="Tahoma", size=10, bold=True)
THIN = Border(bottom=Side(style="thin", color="D5DEE2"))
MONEY = "#,##0"


def sheet(wb, title, headers, widths):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        cell = ws.cell(1, i)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = "A2"
    return ws


def style_row(ws, row, money_cols, font=BODY_FONT, fill=None):
    for cell in ws[row]:
        cell.font = font
        cell.border = THIN
        if fill:
            cell.fill = fill
        if cell.column in money_cols:
            cell.number_format = MONEY


def write(path, branch_ids, hybrids, branches, visitors, agg, ret_rows, debt_rows, status_names):
    wb = Workbook()
    wb.remove(wb.active)
    today = format_jalali_date(date.today())
    toman = lambda rial: round(rial / TOMAN)  # noqa: E731
    vname = lambda vid: visitors[vid]["name"] if vid in visitors else NO_VISITOR  # noqa: E731

    # 1) per hybrid
    ws = sheet(wb, "خلاصه هیبریدها",
               ["کد شعبه", "هیبرید (شعبه)", "هیبرید من", "کد هیبرید من", "تعداد ویزیتور",
                "تعداد چک برگشتی", "مبلغ چک برگشتی (تومان)", "مانده بدهی مشتریان (تومان)",
                "از این بدهی: بابت چک برگشتی (تومان)", "مشتری بدهکار", "سهم چک برگشتی از بدهی"],
               [9, 24, 34, 22, 11, 12, 20, 22, 22, 12, 14])
    totals = defaultdict(float)
    order = sorted(branch_ids, key=lambda b: -sum(a["debt"] for (bb, _), a in agg.items() if bb == b))
    for b in order:
        items = [a for (bb, _), a in agg.items() if bb == b]
        row = {
            "visitors": sum(1 for (bb, v) in agg if bb == b and v is not None),
            "ret_n": sum(a["ret_n"] for a in items), "ret_amt": sum(a["ret_amt"] for a in items),
            "debt": sum(a["debt"] for a in items), "dfr": sum(a["debt_from_ret"] for a in items),
            "debtors": len(set().union(*(a["debtors"] for a in items))) if items else 0,
        }
        for k, v in row.items():
            totals[k] += v
        ws.append([branches.get(b, {}).get("code", b), branches.get(b, {}).get("name", str(b)),
                   "، ".join(n for _, n in hybrids[b]), "، ".join(c for c, _ in hybrids[b]), row["visitors"],
                   row["ret_n"], toman(row["ret_amt"]), toman(row["debt"]), toman(row["dfr"]), row["debtors"],
                   row["ret_amt"] / row["debt"] if row["debt"] else None])
        style_row(ws, ws.max_row, {7, 8, 9})
        ws.cell(ws.max_row, 11).number_format = "0%"
    ws.append(["", "جمع کل", "", "", int(totals["visitors"]), int(totals["ret_n"]), toman(totals["ret_amt"]),
               toman(totals["debt"]), toman(totals["dfr"]), int(totals["debtors"]),
               totals["ret_amt"] / totals["debt"] if totals["debt"] else None])
    style_row(ws, ws.max_row, {7, 8, 9}, BOLD, TOTAL_FILL)
    ws.cell(ws.max_row, 11).number_format = "0%"
    ws.auto_filter.ref = f"A1:K{ws.max_row - 1}"

    # 2) per hybrid × visitor
    ws = sheet(wb, "هیبرید و ویزیتور",
               ["کد شعبه", "هیبرید (شعبه)", "کد ویزیتور", "ویزیتور", "سرپرست", "وضعیت ویزیتور",
                "تعداد چک برگشتی", "مبلغ چک برگشتی (تومان)", *[f"{s} (تومان)" for s in status_names],
                "مانده بدهی مشتریان (تومان)", "از این بدهی: بابت چک برگشتی (تومان)", "مشتری بدهکار"],
               [9, 24, 11, 34, 22, 13, 12, 20, *[20] * len(status_names), 22, 22, 12])
    n_status = len(status_names)
    money_cols = {8, *range(9, 9 + n_status), 9 + n_status, 10 + n_status}
    for b in order:
        keys = sorted([k for k in agg if k[0] == b], key=lambda k: (-agg[k]["debt"], -agg[k]["ret_amt"]))
        sub = defaultdict(float)
        for _, vid in keys:
            a = agg[(b, vid)]
            v = visitors.get(vid, {})
            state = "—" if vid is None else ("فعال" if v.get("active") else "غیرفعال")
            ws.append([branches.get(b, {}).get("code", b), branches.get(b, {}).get("name", str(b)),
                       v.get("code", "") if vid is not None else "", vname(vid), v.get("supervisor", ""), state,
                       a["ret_n"], toman(a["ret_amt"]), *[toman(a["by_status"].get(s, 0)) for s in status_names],
                       toman(a["debt"]), toman(a["debt_from_ret"]), len(a["debtors"])])
            style_row(ws, ws.max_row, money_cols)
            sub["ret_n"] += a["ret_n"]; sub["ret_amt"] += a["ret_amt"]; sub["debt"] += a["debt"]  # noqa: E702
            sub["dfr"] += a["debt_from_ret"]; sub["debtors"] += len(a["debtors"])  # noqa: E702
            for s in status_names:
                sub[s] += a["by_status"].get(s, 0)
        ws.append(["", f"جمع {branches.get(b, {}).get('name', b)}", "", "", "", "", int(sub["ret_n"]), toman(sub["ret_amt"]),
                   *[toman(sub[s]) for s in status_names], toman(sub["debt"]), toman(sub["dfr"]), ""])
        style_row(ws, ws.max_row, money_cols, BOLD, SUB_FILL)
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

    # 3) every unsettled returned cheque
    ws = sheet(wb, "ریز چک‌های برگشتی",
               ["هیبرید (شعبه)", "کد ویزیتور", "ویزیتور", "کد مشتری", "مشتری", "شماره چک", "شناسه صیادی", "بانک",
                "مبلغ (تومان)", "تاریخ سررسید", "روز از سررسید", "تاریخ دریافت", "شماره فاکتور", "وضعیت"],
               [24, 11, 30, 13, 34, 13, 20, 18, 16, 13, 12, 13, 12, 18])
    for r in sorted(ret_rows, key=lambda r: (branches.get(r["branch"], {}).get("name", ""), vname(r["visitor"]), -r["amount"])):
        v = visitors.get(r["visitor"], {})
        ws.append([branches.get(r["branch"], {}).get("name", r["branch"]), v.get("code", ""), vname(r["visitor"]),
                   r["cust_code"], r["cust_name"], r["serial"], r["sayad"], r["bank"], toman(r["amount"]), r["due"],
                   r["days"], r["booked"], r["factor"], r["status"]])
        style_row(ws, ws.max_row, {9})
    ws.auto_filter.ref = f"A1:N{ws.max_row}"

    # 4) debtor customers per visitor
    ws = sheet(wb, "ریز بدهی مشتریان",
               ["هیبرید (شعبه)", "کد ویزیتور", "ویزیتور", "کد مشتری", "مشتری", "سهم این ویزیتور از بدهی (تومان)",
                "از آن بابت چک برگشتی (تومان)"],
               [24, 11, 30, 13, 34, 24, 22])
    per = defaultdict(lambda: [0.0, 0.0])
    meta = {}
    for r in debt_rows:
        k = (r["branch"], r["visitor"], r["dl"])
        per[k][0] += r["amount"]
        if r["kind"] == "returned_cheque":
            per[k][1] += r["amount"]
        meta[k] = (r["cust_code"], r["cust_name"])
    for k in sorted(per, key=lambda k: (branches.get(k[0], {}).get("name", ""), vname(k[1]), -per[k][0])):
        b, vid, _ = k
        ws.append([branches.get(b, {}).get("name", b), visitors.get(vid, {}).get("code", ""), vname(vid),
                   meta[k][0], meta[k][1], toman(per[k][0]), toman(per[k][1])])
        style_row(ws, ws.max_row, {6, 7})
    ws.auto_filter.ref = f"A1:G{ws.max_row}"

    # 5) notes
    ws = wb.create_sheet("توضیحات")
    ws.sheet_view.rightToLeft = True
    ws.column_dimensions["A"].width = 130
    notes = [
        f"گزارش چک برگشتی و مانده بدهی هیبریدها — تاریخ تهیه: {today} — منبع: کارآمد (KDB)، فقط خواندنی",
        "مبالغ به تومان است.",
        "",
        "چک برگشتی: کارآمد وضعیت «واخواست شده» ندارد (آن وضعیت مال راهکاران/B2B است). معادل آن در کارآمد چک‌های برگشتیِ هنوز تسویه‌نشده است:",
        f"   وضعیت‌های {', '.join(status_names) or '—'} (کدهای {', '.join(str(s) for s in OPEN_RETURNED_CHEQUE_STATUSES)}). «برگشتی وصول شده» تسویه شده و در گزارش نیست.",
        "هر چک برگشتی به ویزیتورِ فاکتوری می‌رسد که چک بابت آن دریافت شده؛ اگر چک به فاکتوری وصل نباشد، در ردیف «بدون ویزیتور» شعبه‌اش می‌آید.",
        "",
        "مانده بدهی: مانده دفتر کل هر مشتری (اسناد قطعی و صادرنشده)، مثل داشبورد مالی. سهم هر ویزیتور:",
        "   ۱) اول چک برگشتی تسویه‌نشدهٔ مشتری (به ویزیتور فاکتور همان چک)، ۲) بقیه به روش FIFO: پرداخت‌ها اول قدیمی‌ترین فاکتورها را تسویه می‌کنند،",
        "   پس مانده مال جدیدترین فاکتورهاست و به ویزیتورِ همان فاکتورها می‌رسد. هر ریال بدهی فقط یک بار شمرده می‌شود.",
        "   ستون «از این بدهی: بابت چک برگشتی» بخشی از مانده است که از چک برگشتی آمده؛ می‌تواند از مبلغ چک برگشتی کمتر باشد اگر مشتری بخشی را پرداخته.",
        "هیبرید هر ردیف = شعبهٔ ویزیتور. مشتری بدون هیچ فاکتوری به شعبهٔ اصلی خودش می‌رسد (ردیف «بدون ویزیتور»).",
        "جمع بدهی یک هیبرید در این فایل جمع ویزیتورهایش است و ممکن است با عدد «مطالبات» داشبورد (مشتریانی که شعبه اصلی‌شان این شعبه است) کمی فرق کند.",
        "ویزیتورهای فعالِ بدون بدهی و چک برگشتی هم برای کامل بودن فهرست آمده‌اند.",
    ]
    for line in notes:
        ws.append([line])
        ws.cell(ws.max_row, 1).font = BOLD if line == notes[0] else BODY_FONT
    wb.move_sheet("توضیحات", offset=-(len(wb.sheetnames) - 1))
    wb.active = 1
    wb.save(path)


if __name__ == "__main__":
    today = format_jalali_date(date.today()).replace("/", "-")
    default = Path(os.path.expanduser("~")) / "Desktop" / f"چک-برگشتی-و-بدهی-هیبریدها-{today}.xlsx"
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else default
    n, amt, debt = build(out)
    print(f"OK {out}\nreturned cheques: {n}  amount_toman: {amt / TOMAN:,.0f}  debt_toman: {debt / TOMAN:,.0f}")
